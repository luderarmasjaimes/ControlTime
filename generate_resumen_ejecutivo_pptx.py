from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUTPUT = r'c:\InformeCliente\docs\Resumen_Ejecutivo_Cronograma_VPS_Lima_26_Semanas.pptx'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

COLORS = {
    'navy': RGBColor(17, 34, 64),
    'green': RGBColor(34, 94, 78),
    'sand': RGBColor(245, 239, 224),
    'gold': RGBColor(184, 134, 11),
    'stone': RGBColor(219, 211, 189),
    'slate': RGBColor(74, 85, 104),
    'white': RGBColor(255, 255, 255),
    'ink': RGBColor(28, 37, 44),
}

slides = [
    {
        'title': 'Cronograma Maestro de Implementacion',
        'subtitle': 'Plataforma Minera sobre VPS Linux en Lima | 26 semanas | 2 etapas',
        'tag': 'AURIXA | Programa Ejecutivo',
        'bullets': [
            'Nueva plataforma sobre VPS Linux en Lima, Peru.',
            'Consumo controlado de datos desde la plataforma actual en AWS.',
            'Ejecucion en 26 semanas con lectura ejecutiva 4+2 meses.',
            'Plan construido para maximizar control, continuidad y ocupacion efectiva del equipo.',
        ],
        'highlight': 'Decisiones ya cerradas: VPS Linux en Lima, 26 semanas, 2 etapas y dependencia controlada con la base externa actual.',
    },
    {
        'title': 'Vision Ejecutiva del Proyecto',
        'tag': 'Direccion y Alcance',
        'bullets': [
            'Implementar una nueva plataforma minera de reportes, monitoreo y control documental.',
            'Evitar una nueva dependencia de arquitectura cloud publica para esta solucion.',
            'Mantener continuidad del negocio integrando datos de la plataforma actual en produccion.',
            'Entregar a gerencia una solucion comprensible, auditable y operable desde Lima.',
        ],
        'highlight': 'La arquitectura propuesta reduce complejidad, mejora control operativo y mantiene continuidad del negocio.',
    },
    {
        'title': 'Estructura del Cronograma Oficial',
        'tag': 'Secuencia de Entrega',
        'bullets': [
            'Etapa 1: Semanas 1 a 16. Implementacion funcional core.',
            'Etapa 2: Semanas 17 a 26. Hardening, estabilizacion y go-live.',
            'La referencia ejecutiva de 4 meses + 2 meses se mantiene sin perder rigor calendario.',
            'Cada etapa cierra con entregables, riesgos, responsables y criterios de validacion.',
        ],
        'highlight': 'La narrativa ejecutiva 4+2 se conserva, pero el control real del programa se ejecuta sobre 26 semanas calendario.',
    },
    {
        'title': 'Alcance Ejecutivo de la Etapa 1',
        'tag': 'Etapa 1 | Semanas 1-16',
        'bullets': [
            'Arquitectura, seguridad, conectividad e integracion con la fuente externa.',
            'Motor central de negocio y editor tipo Word para informes mineros.',
            'Tablero gerencial, monitoreo de sensores y visualizacion GIS.',
            'Exportacion documental, trazabilidad y QA funcional de cierre.',
        ],
        'highlight': 'La Etapa 1 entrega el nucleo funcional y deja habilitada la base para pasar a hardening sin reabrir arquitectura.',
    },
    {
        'title': 'Alcance Ejecutivo de la Etapa 2',
        'tag': 'Etapa 2 | Semanas 17-26',
        'bullets': [
            'Hardening de infraestructura, accesos y continuidad operativa.',
            'Estabilizacion del consumo de datos desde la plataforma externa actual.',
            'Pruebas de rendimiento, seguridad, UAT y recuperacion ante incidentes.',
            'Marcha blanca, go-live y transferencia formal al cliente.',
        ],
        'highlight': 'La Etapa 2 concentra cierre de riesgo, adopcion, continuidad y salida controlada a produccion.',
    },
    {
        'title': 'Estrategia de Ocupacion del Equipo',
        'tag': 'Uso de Recursos',
        'bullets': [
            'Arquitectura y sistemas con alta participacion desde el inicio.',
            'Backend sostenido sobre el camino critico desde el mes 2.',
            'Frontend reforzado para adopcion y calidad visual desde el mes 2.',
            'QA e infraestructura intensificados hacia los cierres de etapa y salida a produccion.',
        ],
        'highlight': 'La ocupacion ya no es lineal: fue recalibrada para proteger camino critico, adopcion y go-live.',
    },
    {
        'title': 'Riesgos de Direccion Mas Relevantes',
        'tag': 'Gobierno de Riesgo',
        'bullets': [
            'Dependencia de la base externa actual en AWS.',
            'Capacidad insuficiente de VPS Linux ante picos de carga.',
            'Retraso en hallazgos de seguridad y continuidad.',
            'Baja adopcion operativa en marcha blanca.',
        ],
        'highlight': 'El riesgo dominante no es la infraestructura objetivo, sino la gestion disciplinada de dependencias y cierre de hallazgos.',
    },
    {
        'title': 'Hitos Ejecutivos y Recomendacion Final',
        'tag': 'Cierre Ejecutivo',
        'bullets': [
            'Semana 4: arquitectura y gobierno aprobados.',
            'Semana 16: etapa 1 cerrada con integracion y QA funcional.',
            'Semana 20: hardening y estabilizacion principal completados.',
            'Semana 26: go-live aprobado y transferencia formal al cliente.',
        ],
        'footer': 'Recomendacion: aprobar el cronograma de 26 semanas y mantener gobernanza quincenal de riesgos y dependencias.',
        'highlight': 'El programa ya esta estructurado para aprobacion ejecutiva, seguimiento quincenal y salida a produccion controlada.',
    },
]


def add_background(slide, accent=False, title_slide=False):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = COLORS['sand'] if not title_slide else COLORS['navy']

    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.65))
    band.fill.solid()
    band.fill.fore_color.rgb = COLORS['green'] if accent else COLORS['navy']
    band.line.fill.background()

    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.35), Inches(6.85), Inches(1.7), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS['gold']
    bar.line.fill.background()

    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(12.4), Inches(0), Inches(0.28), prs.slide_height)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = COLORS['gold'] if title_slide else COLORS['stone']
    stripe.line.fill.background()

    if title_slide:
        panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(1.45), Inches(12.0), Inches(4.5))
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(245, 239, 224)
        panel.fill.transparency = 0.06
        panel.line.fill.background()


def add_branding(slide, tag, title_slide=False):
    tag_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(4.8), Inches(0.3))
    tf = tag_box.text_frame
    p = tf.paragraphs[0]
    p.text = tag
    p.font.name = 'Aptos'
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLORS['white']

    logo = slide.shapes.add_textbox(Inches(10.55), Inches(0.13), Inches(1.4), Inches(0.32))
    tf = logo.text_frame
    p = tf.paragraphs[0]
    p.text = 'AURIXA'
    p.font.name = 'Aptos Display'
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLORS['gold'] if title_slide else COLORS['white']
    p.alignment = PP_ALIGN.RIGHT

    footer = slide.shapes.add_textbox(Inches(9.0), Inches(6.68), Inches(3.2), Inches(0.25))
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = 'Programa VPS Lima | Mineria | 2026'
    p.font.name = 'Aptos'
    p.font.size = Pt(9)
    p.font.color.rgb = COLORS['slate'] if not title_slide else COLORS['white']
    p.alignment = PP_ALIGN.RIGHT


def add_title(slide, title, subtitle=None, title_slide=False):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0 if title_slide else 0.8), Inches(11.3), Inches(0.9))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = 'Aptos Display'
    p.font.size = Pt(28 if title_slide else 24)
    p.font.bold = True
    p.font.color.rgb = COLORS['navy'] if not title_slide else COLORS['navy']

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.85 if title_slide else 1.45), Inches(11.0), Inches(0.7))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.name = 'Aptos'
        p.font.size = Pt(13 if title_slide else 12)
        p.font.color.rgb = COLORS['slate']


def add_bullets(slide, bullets, footer=None, title_slide=False):
    body = slide.shapes.add_textbox(Inches(0.95), Inches(2.75 if title_slide else 2.0), Inches(7.7 if not title_slide else 10.8), Inches(3.6))
    tf = body.text_frame
    tf.word_wrap = True
    for index, item in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = 'Aptos'
        p.font.size = Pt(20 if len(bullets) <= 4 else 18)
        p.font.color.rgb = COLORS['ink']
        p.space_after = Pt(10)

    if footer:
        foot = slide.shapes.add_textbox(Inches(0.85), Inches(6.1), Inches(11.6), Inches(0.55))
        tf = foot.text_frame
        p = tf.paragraphs[0]
        p.text = footer
        p.font.name = 'Aptos'
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLORS['green']
        p.alignment = PP_ALIGN.LEFT


def add_highlight(slide, text, title_slide=False):
    if title_slide:
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(5.65), Inches(10.4), Inches(0.7))
    else:
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.95), Inches(1.9), Inches(2.9), Inches(3.9))
    box.fill.solid()
    box.fill.fore_color.rgb = COLORS['green'] if title_slide else COLORS['navy']
    box.fill.transparency = 0.06 if title_slide else 0.02
    box.line.fill.background()

    inner = slide.shapes.add_textbox(
        Inches(1.18) if title_slide else Inches(9.2),
        Inches(5.86) if title_slide else Inches(2.15),
        Inches(9.9) if title_slide else Inches(2.4),
        Inches(0.35) if title_slide else Inches(3.4),
    )
    tf = inner.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = 'Aptos'
    p.font.size = Pt(11 if title_slide else 14)
    p.font.bold = True if title_slide else False
    p.font.color.rgb = COLORS['white']
    p.alignment = PP_ALIGN.LEFT


for idx, data in enumerate(slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    is_title = idx == 0
    add_background(slide, accent=(idx % 2 == 1), title_slide=is_title)
    add_branding(slide, data.get('tag', 'Programa Ejecutivo'), title_slide=is_title)
    add_title(slide, data['title'], data.get('subtitle'), title_slide=is_title)
    add_bullets(slide, data['bullets'], data.get('footer'), title_slide=is_title)
    add_highlight(slide, data.get('highlight', ''), title_slide=is_title)

prs.save(OUTPUT)
print(OUTPUT)
