#!/usr/bin/env python3
"""
Genera 3 convocatorias laborales en formato LinkedIn (modelo trabajoLINKEDIN.docx).
Lenguaje industrial genérico — sin stack propietario, C++, ni detalles de arquitectura confidenciales.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
CONTACT = "LUDER.EDER.ARMAS.JAIMES.LEAJ@GMAIL.COM"

TITLE_COLOR = RGBColor(14, 61, 87)
ACCENT = RGBColor(12, 104, 102)


def _p(doc: Document, text: str, *, bold: bool = False, size: float = 11) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(34, 34, 34)
    run.bold = bold
    para.paragraph_format.space_after = Pt(6)


def _section(doc: Document, title: str) -> None:
    _p(doc, title, bold=True, size=12)


def _bullet(doc: Document, text: str) -> None:
    para = doc.add_paragraph(style="List Bullet")
    run = para.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    para.paragraph_format.space_after = Pt(3)


def _about_job(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run("Acerca del empleo")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    p.paragraph_format.space_after = Pt(8)


def _cta(doc: Document, role_tag: str) -> None:
    doc.add_paragraph()
    _p(
        doc,
        f"¿Listo para asumir este desafío? Envía tu CV en PDF y una carta breve (máx. 1 página) "
        f"a {CONTACT} con asunto: {role_tag} — [Tu nombre].",
    )
    _p(doc, "Modalidad: Remoto / híbrido · Proyecto: 6 meses · Dedicación: Full Time · LATAM", bold=False)


def build_backend_docx(path: Path) -> None:
    doc = Document()
    _about_job(doc)

    _p(
        doc,
        "Organización del sector industrial busca un profesional SENIOR para liderar la capa de datos, "
        "persistencia y servicios backend de una plataforma web enterprise de monitoreo, reportabilidad "
        "operativa y gestión documental avanzada.",
        bold=True,
    )
    _p(
        doc,
        "¿Te apasiona la arquitectura de datos, la continuidad operativa y la trazabilidad en entornos "
        "con conectividad limitada? Esta es tu oportunidad de liderar un proyecto de transformación digital "
        "de alto impacto en la región.",
    )

    _section(doc, "¿Cuál será tu misión?")
    _p(
        doc,
        "Garantizar que la arquitectura de datos y los servicios de núcleo soporten de forma confiable "
        "informes técnicos, telemetría operativa, auditoría de acciones, operación offline con sincronización "
        "controlada e integración segura con sistemas existentes de planta.",
    )

    _section(doc, "¿Qué harás en tu día a día?")
    for item in [
        "Diseñar e implementar modelos de datos transaccionales, documentales y de series temporales (multi-empresa, auditoría y retención).",
        "Desarrollar y mantener servicios backend de alto rendimiento (APIs REST, canales en tiempo real, autoservicio documental).",
        "Construir pipelines ETL/ELT e integraciones con fuentes legacy e IoT/industrial, con idempotencia y monitoreo de rezagos.",
        "Diseñar estrategias de persistencia local, colas de cambios y reconciliación al restablecer conectividad.",
        "Administrar entornos Linux y contenedores: backups probados, recuperación ante desastres, migraciones versionadas.",
        "Aplicar controles de seguridad en capa de datos: segregación por tenant, cifrado, bitácoras inmutables y cumplimiento normativo.",
        "Colaborar con frontend, QA y PMO en contratos de integración, pruebas de carga y documentación operativa.",
        "Participar en metodología ágil con Git y herramienta de gestión de proyectos (ClickUp o equivalente).",
    ]:
        _bullet(doc, item)

    _section(doc, "Requisitos para postular:")
    for item in [
        "8+ años de experiencia en backend y arquitectura de datos; 5+ años con motor relacional en producción (excluyente).",
        "Dominio avanzado de PostgreSQL: modelado, JSON/documentos, partición, índices, replicación, backup/restore y tuning (excluyente).",
        "Experiencia en servicios backend de alto rendimiento, APIs REST y comunicación en tiempo real (excluyente).",
        "Diseño de trazabilidad, auditoría forense y gobierno de datos alineado a Industria 4.0/5.0 (excluyente).",
        "Experiencia en ETL/ELT, integración con sistemas externos y sincronización incremental (excluyente).",
        "Conocimientos de seguridad aplicada a datos: RBAC, cifrado en tránsito/reposo, OWASP, hardening de accesos (excluyente).",
        "Operación en Linux, Docker y entornos productivos con CI/CD (excluyente).",
        "Experiencia en proyectos industriales: manufactura, energía, utilities o plantas automatizadas (excluyente).",
        "Git + gestión ágil en equipos multidisciplinarios (excluyente).",
        "Español fluido; inglés técnico de lectura (excluyente).",
    ]:
        _bullet(doc, item)

    _section(doc, "Conocimientos deseables:")
    for item in [
        "Motores de series temporales y telemetría de alta frecuencia.",
        "Mensajería/streaming para ingesta ordenada de eventos.",
        "Multi-tenant con Row-Level Security y políticas de retención legal.",
        "Validación biométrica integrada a flujos de acceso (modelado de eventos).",
        "Certificaciones DBA, seguridad de la información o IEC 62443.",
        "Experiencia en LATAM con operaciones remotas y conectividad intermitente.",
        "Marcos ISO aplicables a calidad de software y evidencias de trazabilidad.",
    ]:
        _bullet(doc, item)

    _section(doc, "¿Qué te ofrecemos?")
    for item in [
        "Proyecto de alto impacto en transformación digital industrial.",
        "Rol con autonomía técnica en datos, persistencia y núcleo de servicios.",
        "Equipo multidisciplinario con metodología ágil definida.",
        "Compensación acorde a perfil Senior, a coordinar en selección.",
    ]:
        _bullet(doc, item)

    _cta(doc, "Senior Backend / Arquitecto de Datos")
    doc.save(path)


def build_frontend_docx(path: Path) -> None:
    doc = Document()
    _about_job(doc)

    _p(
        doc,
        "Organización del sector industrial busca un Frontend Senior para liderar la experiencia de usuario "
        "de una plataforma enterprise de informes técnicos, tableros operativos y visualización geoespacial.",
        bold=True,
    )
    _p(
        doc,
        "¿Dominas React en producción, editores de texto enriquecido y UX para entornos de campo? "
        "Únete desde el primer mes del proyecto como referente frontend del equipo.",
    )

    _section(doc, "¿Cuál será tu misión?")
    _p(
        doc,
        "Construir una interfaz de clase enterprise: editor documental avanzado, operación offline con sincronización "
        "transparente, navegación optimizada, accesibilidad industrial y integración fluida con servicios backend "
        "y datos en tiempo real.",
    )

    _section(doc, "¿Qué harás en tu día a día?")
    for item in [
        "Levantar y evolucionar arquitectura frontend moderna (React, TypeScript, bundler ágil, design system corporativo).",
        "Implementar editor de informes técnicos con estilos, tablas, imágenes, plantillas y flujos de aprobación.",
        "Desarrollar modo offline: almacenamiento local, indicadores de estado y reconciliación al reconectar.",
        "Optimizar UX: atajos de teclado, menús eficientes, paneles colapsables y maximización del área de trabajo.",
        "Integrar mapas interactivos, dashboards KPI y visualización de telemetría en tiempo real.",
        "Consumir APIs REST/WebSocket con manejo robusto de errores, reconexión y autoservicio de guardado.",
        "Validar usabilidad en tablets de campo, alto contraste y condiciones operativas exigentes.",
        "Trabajar con Git, pruebas E2E y ceremonias ágiles junto a backend, QA y PMO.",
    ]:
        _bullet(doc, item)

    _section(doc, "Requisitos para postular:")
    for item in [
        "6+ años con React/TypeScript en producción; 8+ años en desarrollo frontend (excluyente).",
        "Experiencia con editores rich-text o componentes documentales complejos (excluyente).",
        "Modo offline-first: Service Worker, almacenamiento local y estrategias de sincronización (excluyente).",
        "Integración con APIs en tiempo real y gestión de estado reactiva (excluyente).",
        "UX para entornos industriales: responsive, accesibilidad, tablets de campo (excluyente).",
        "Mapas interactivos y visualización de datos en vivo (excluyente).",
        "Calidad de código: linting, tests E2E y revisión en Git (excluyente).",
        "Git + ClickUp o herramienta ágil equivalente (excluyente).",
        "Disponibilidad de inicio en mes 1 del proyecto (excluyente).",
        "Español fluido; inglés técnico de lectura (excluyente).",
    ]:
        _bullet(doc, item)

    _section(doc, "Conocimientos deseables:")
    for item in [
        "Patrones de toolbar empresarial y experiencia tipo procesador ofimático.",
        "PWA, exportación documental client-side y accesibilidad WCAG 2.1.",
        "Gráficos interactivos conectados a telemetría operativa.",
        "Experiencia en sector industrial, energía o manufactura en LATAM.",
    ]:
        _bullet(doc, item)

    _section(doc, "¿Qué te ofrecemos?")
    for item in [
        "Rol líder frontend desde el arranque del proyecto (6 meses).",
        "Autonomía para definir estándares UX/UI de la plataforma.",
        "Equipo técnico de alto nivel y metodología ágil.",
        "Compensación acorde a perfil Senior.",
    ]:
        _bullet(doc, item)

    _cta(doc, "Senior Frontend — Plataforma Enterprise")
    doc.save(path)


def build_paf_docx(path: Path) -> None:
    doc = Document()
    _about_job(doc)

    _p(
        doc,
        "Organización del sector industrial busca un Analista Funcional y Profesional de Apoyo PMO "
        "para un proyecto de software enterprise de 6 meses con equipo multidisciplinario de 10 especialistas.",
        bold=True,
    )
    _p(
        doc,
        "¿Destacas en documentación clara, coordinación de equipos y capacitación a usuarios? "
        "Serás el eje operativo entre el equipo técnico, el Project Manager y la gerencia de negocio.",
    )

    _section(doc, "¿Cuál será tu misión?")
    _p(
        doc,
        "Asegurar la gestión documental, el seguimiento del cronograma, la comunicación con stakeholders "
        "y la transferencia de conocimiento funcional, manteniendo trazabilidad requisito–entregable y "
        "material de capacitación actualizado.",
    )

    _section(doc, "¿Qué harás en tu día a día?")
    for item in [
        "Actualizar semanalmente la herramienta de gestión (ClickUp): estados, fechas, dependencias y bloqueos.",
        "Coordinar reuniones: convocatorias, agendas, actas, action items y seguimiento hasta cierre.",
        "Redactar manuales funcionales, quick-start y material de capacitación para usuarios de operación.",
        "Mantener matriz de trazabilidad requisito–tarea–módulo y correspondencias funcionales.",
        "Apoyar reportes de avance, gestión de riesgos y comunicación con gerencia.",
        "Preparar escenarios y log de incidencias para pruebas de aceptación (UAT).",
        "Diseñar y ejecutar talleres de capacitación en fases finales del proyecto.",
        "Compilar kit de cierre: índice documental, lecciones aprendidas y acta de transferencia.",
    ]:
        _bullet(doc, item)

    _section(doc, "Requisitos para postular:")
    for item in [
        "5+ años en análisis funcional, documentación o PMO en proyectos TI (excluyente).",
        "Dominio de ClickUp, Jira o Azure DevOps (excluyente).",
        "Redacción ejecutiva: actas, manuales de usuario, procedimientos y matrices RACI (excluyente).",
        "Coordinación de múltiples stakeholders técnicos y de negocio (excluyente).",
        "Metodología ágil: Scrum/Kanban, ceremonias y seguimiento de compromisos (excluyente).",
        "Excel y Word/Markdown para documentación formal (excluyente).",
        "Español nativo o fluido; inglés de lectura (excluyente).",
        "Disponibilidad full time durante 6 meses (excluyente).",
    ]:
        _bullet(doc, item)

    _section(doc, "Conocimientos deseables:")
    for item in [
        "Experiencia como capacitador de usuarios no TI en operaciones de campo.",
        "Conocimiento básico de arquitectura web para dialogar con equipos de desarrollo.",
        "Certificación PMP, Scrum Master o ITIL Foundation.",
        "Proyectos enterprise en sector industrial, utilities o manufactura.",
        "Documentación alineada a ISO 9001 y auditorías.",
    ]:
        _bullet(doc, item)

    _section(doc, "¿Qué te ofrecemos?")
    for item in [
        "Rol transversal con visibilidad directa ante gerencia y PM.",
        "Proyecto enterprise de 6 meses con equipo de élite.",
        "Desarrollo profesional en gestión de proyectos TI complejos.",
        "Compensación acorde a experiencia.",
    ]:
        _bullet(doc, item)

    _cta(doc, "Analista Funcional / Soporte PMO")
    doc.save(path)


def main() -> int:
    outputs = [
        (build_backend_docx, DOCS / "Convocatoria_Laboral_01_Backend_LinkedIn.docx"),
        (build_frontend_docx, DOCS / "Convocatoria_Laboral_02_Frontend_LinkedIn.docx"),
        (build_paf_docx, DOCS / "Convocatoria_Laboral_03_PAF_LinkedIn.docx"),
    ]
    for builder, path in outputs:
        builder(path)
        print(f"Generado: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
