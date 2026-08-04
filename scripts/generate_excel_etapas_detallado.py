#!/usr/bin/env python3
"""
Excel detallado en 2 pestañas (Etapa 1 y Etapa 2) para revision gerencial:
recursos, carga mensual, tareas secuenciadas, entregables por sprint, riesgos.
"""
from __future__ import annotations

import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LINK_FONT = Font(name="Calibri", color="0563C1", underline="single", size=11)

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

import generate_clickup_import as gci  # noqa: E402
from generate_plan_gerencia_etapas_xlsx import (  # noqa: E402
    ETAPA1_RIESGOS,
    ETAPA2_RIESGOS,
    FILLS,
    SPRINT_ENTREGABLES,
)

OUT = ROOT / "docs" / "Plan_Detallado_Etapas_1_y_2_v36.xlsx"

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
SECTION_FONT = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="1F4E79")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)
THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

TRANSVERSAL_FOLDERS = {"Gestion PMO", "Riesgos", "Control de calidad", "DevOps"}

ETAPA1_FOCUS = {
    "BE1": "Nucleo C++ tiempo real, sockets telemetria, APIs editor/mapas/export/offline",
    "BE2": "Seguridad, usuarios/perfiles/accesos, biometria, integracion IA aplicada",
    "BE3": "Diseno BD PostgreSQL, sync AWS, ETL, backups",
    "FE1": "Shell UI, ReportStudio, mapas, dashboards, exportacion",
    "FE2": "UX tablet, accesibilidad, pruebas con operadores (desde Mes 2)",
    "ARQ": "Gobierno tecnico, gates R1-R4, comite ejecutivo",
    "SYS": "Docker, CI/CD, monitoreo base (desde Mes 2)",
    "QA": "Plan de pruebas, regresion, validacion funcional (desde Mes 2)",
    "IA": "Entrenamiento modelos ONNX/NLP/STT (desde Mes 2)",
    "PAF": "ClickUp, actas, trazabilidad FE-BE, documentacion funcional",
}

ETAPA2_FOCUS = {
    "BE1": "Optimizacion tiempo real, pruebas carga sockets, estabilizacion core",
    "BE2": "Pentest OWASP, hardening, IA produccion, certificacion seguridad",
    "BE3": "Simulacro DR, replicacion, retencion datos, tuning consultas",
    "FE1": "Refinamiento UI, performance cliente, ajustes UAT",
    "FE2": "UAT campo, marcha blanca, NPS post-lanzamiento",
    "ARQ": "Sign-off Go-Live, UAT ejecutivo, cierre contractual",
    "SYS": "Hardening infra, HA, blue/green, monitoreo 24/7",
    "QA": "Pentest funcional, load test 10K sensores, certificacion UAT",
    "IA": "Afinamiento modelos, precision inferencia produccion",
    "PAF": "Manuales finales, kit cierre, capacitacion, lecciones aprendidas",
}


def set_col_widths(ws, widths: list) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_section(
    ws, row: int, title: str, ncol: int, fill: PatternFill,
    anchors: list | None = None, anchor_label: str | None = None,
) -> int:
    if anchors is not None and anchor_label:
        anchors.append({"label": anchor_label, "row": row, "section": title.strip()})
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = SECTION_FONT
    cell.fill = fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if ncol > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    return row + 1


def write_table(ws, headers: list, rows: list, start_row: int, fill: PatternFill) -> int:
    ncol = len(headers)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = fill
        cell.alignment = HEADER_ALIGN
        cell.border = THIN
    r = start_row + 1
    for row in rows:
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = THIN
            cell.alignment = WRAP
        r += 1
    return r + 1


def build_successor_map(tasks) -> dict[str, list[str]]:
    succ: dict[str, list[str]] = defaultdict(list)
    for t in tasks:
        tid = t[2]
        for dep in (t[11] or "").split(","):
            dep = dep.strip()
            if dep:
                succ[dep].append(tid)
    return succ


def task_hours_in_month(tid: str, month: int) -> float:
    alloc = gci.TASK_ALLOCS.get(tid, {})
    return round(sum(m.get(month, 0) for m in alloc.values()), 1)


def task_in_etapa_month(task, month: int, etapa_months: list[int], paso: str) -> bool:
    tid = task[2]
    if str(tid).startswith(("CAP-", "FLOOR-")):
        return False
    if task_hours_in_month(tid, month) < 0.5:
        return False
    folder = task[0]
    if folder == paso:
        return True
    if folder in TRANSVERSAL_FOLDERS:
        return True
    return False


def collect_tasks_for_month(month: int, etapa_months: list[int], paso: str):
    items = []
    for task in gci.TASKS:
        if not task_in_etapa_month(task, month, etapa_months, paso):
            continue
        tid = task[2]
        h_m = task_hours_in_month(tid, month)
        start, due = gci.TASK_DATES.get(tid, (None, None))
        items.append((task, h_m, start, due))
    items.sort(key=lambda x: (x[2] or date.max, x[0][2]))
    return items


def pct_status(pct: float) -> str:
    if pct > 100:
        return "SOBRECARGA"
    if pct >= 90:
        return "OPTIMO"
    if pct >= 70:
        return "ACEPTABLE"
    if pct == 0:
        return "NO ACTIVO"
    return "BAJO"


def build_etapa_sheet(ws, cfg: dict) -> list[dict]:
    months = cfg["months"]
    paso = cfg["paso"]
    fill = cfg["fill"]
    risks = cfg["risks"]
    sprint_label = cfg["sprint_etapa"]
    focus_map = cfg["focus"]
    sheet_name = ws.title
    anchors: list[dict] = []
    ncol = 16

    row = 1
    ws.cell(row=row, column=1, value=cfg["banner"]).font = TITLE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    anchors.append({"label": f"Inicio — {sheet_name}", "row": 1, "section": cfg["banner"][:80]})
    row += 2

    # ── 1. Listado de recursos ──
    row = write_section(
        ws, row, "1. LISTADO DE RECURSOS DEL PROYECTO (equipo Etapa)", ncol, fill,
        anchors, "1. Listado de recursos",
    )
    res_headers = [
        "Cod.", "Rol completo", "Funcion principal en esta etapa",
        "Meses activos en etapa", "Tareas asignadas en etapa", "Horas totales etapa",
    ]
    res_rows = []
    for code in gci.RESOURCES:
        n_tasks = 0
        h_etapa = 0.0
        active_in_etapa = []
        for m in months:
            h = gci.MONTHLY_LOAD[code][m]
            if m in gci.active_months_for_resource(code):
                active_in_etapa.append(gci.MONTH_LABELS[m][:3])
                h_etapa += h
        for task in gci.TASKS:
            if code not in gci.parse_assignees(task[5]):
                continue
            tid = task[2]
            if any(task_hours_in_month(tid, m) >= 0.5 for m in months):
                if task[0] == paso or task[0] in TRANSVERSAL_FOLDERS:
                    n_tasks += 1
        res_rows.append([
            code,
            gci.RESOURCE_DISPLAY[code],
            focus_map.get(code, gci.RESOURCES[code][:80]),
            ", ".join(active_in_etapa) if active_in_etapa else "—",
            n_tasks,
            f"{round(h_etapa, 0):.0f} h",
        ])
    row = write_table(ws, res_headers, res_rows, row, fill)
    row += 1

    # ── 2. Asignacion recursos por mes ──
    row = write_section(
        ws, row, "2. ASIGNACION DE RECURSOS POR MES (% DE USO Y HORAS)", ncol, fill,
        anchors, "2. Asignacion de recursos por mes",
    )

    summary_headers = ["Recurso"] + [f"{gci.MONTH_LABELS[m]} (h)" for m in months] + [
        f"{gci.MONTH_LABELS[m]} (%)" for m in months
    ] + ["Prom. etapa %", "Total h etapa"]
    summary_rows = []
    for code in gci.RESOURCES:
        cells_h = []
        cells_p = []
        pcts = []
        total = 0.0
        for m in months:
            h = round(gci.MONTHLY_LOAD[code][m], 1)
            cap = gci.MONTH_CAPACITY[m]
            if m in gci.active_months_for_resource(code):
                pct = round(h / cap * 100, 1)
                pcts.append(pct)
                total += h
                cells_h.append(h)
                cells_p.append(f"{pct}%")
            else:
                cells_h.append("—")
                cells_p.append("—")
        prom = round(sum(pcts) / len(pcts), 1) if pcts else 0
        summary_rows.append([code, *cells_h, *cells_p, f"{prom}%", round(total, 0)])
    row = write_table(ws, summary_headers, summary_rows, row, fill)

    for m in months:
        mlabel = gci.MONTH_LABELS[m]
        row = write_section(
            ws, row,
            f"   Detalle {mlabel} — capacidad max {gci.MONTH_CAPACITY[m]} h/recurso",
            ncol, FILLS["sub"],
            anchors, f"2.{mlabel} — carga recursos",
        )
        det_headers = ["Recurso", "Rol", "Horas asignadas", "Capacidad mes", "% uso", "Estado", "Observacion"]
        det_rows = []
        for code in gci.RESOURCES:
            if m not in gci.active_months_for_resource(code):
                det_rows.append([
                    code, gci.RESOURCE_DISPLAY[code][:35], 0, gci.MONTH_CAPACITY[m],
                    "—", "NO ACTIVO", "Recurso no planificado este mes",
                ])
                continue
            h = round(gci.MONTHLY_LOAD[code][m], 1)
            cap = gci.MONTH_CAPACITY[m]
            pct = round(h / cap * 100, 1)
            obs = "Meta 90-100%" if 90 <= pct <= 100 else (
                "Por debajo del piso 90%" if pct < 90 else "Sobre tope 100%"
            )
            det_rows.append([
                code, gci.RESOURCE_DISPLAY[code][:35], h, cap, f"{pct}%",
                pct_status(pct), obs,
            ])
        row = write_table(ws, det_headers, det_rows, row, FILLS["sub"])
        row += 1

    # ── 3. Tareas detalladas mes a mes ──
    row = write_section(
        ws, row,
        "3. TAREAS DETALLADAS — ORGANIZADAS MES A MES (SECUENCIA, DEPENDENCIAS Y RECURSOS)",
        ncol, fill,
        anchors, "3. Tareas detalladas (todas)",
    )
    successors = build_successor_map(gci.TASKS)
    task_headers = [
        "Seq.", "ID", "Nombre tarea", "Descripcion (resumen)",
        "Recursos", "Roles asignados (detalle)",
        "Horas tot.", "Horas mes", "Sprint", "Inicio", "Fin",
        "Depende de", "Continua en (tareas siguientes)",
        "Release", "Prioridad", "Riesgo",
    ]

    global_seq = 1
    for m in months:
        mlabel = gci.MONTH_LABELS[m]
        items = collect_tasks_for_month(m, months, paso)
        row = write_section(
            ws, row,
            f"   MES: {mlabel} 2026 — {len(items)} tareas programadas",
            ncol, FILLS["sub"],
            anchors, f"3. Tareas — {mlabel}",
        )
        trows = []
        for task, h_m, start, due in items:
            tid = task[2]
            assignees = gci.parse_assignees(task[5])
            roles = gci.format_resource_field(task[5])
            cont = ", ".join(successors.get(tid, [])[:6])
            if len(successors.get(tid, [])) > 6:
                cont += " …"
            sprint = gci.sprint_label_for_range(start, due) if start and due else task[10]
            trows.append([
                global_seq,
                tid,
                task[3],
                (task[4] or "")[:220],
                ", ".join(assignees),
                roles,
                task[15],
                h_m,
                sprint,
                gci.format_clickup_date(start) if start else "",
                gci.format_clickup_date(due) if due else "",
                (task[11] or "").strip() or "—",
                cont or "—",
                task[10] or "",
                task[7],
                task[13],
            ])
            global_seq += 1
        if not trows:
            trows = [["—", "—", "Sin tareas", "", "", "", "", "", "", "", "", "", "", "", "", ""]]
        row = write_table(ws, task_headers, trows, row, FILLS["sub"])
        row += 1

    # ── 4. Entregables por Sprint ──
    row = write_section(
        ws, row, "4. ENTREGABLES A PRESENTAR EN SPRINT (SPRINT REVIEW + GATE PM)", ncol, fill,
        anchors, "4. Entregables por Sprint",
    )
    sprint_headers = [
        "Sprint", "Release", "Inicio", "Fin", "Objetivo del sprint",
        "Entregables concretos (documentos / funciones listas)",
        "Demo ante Gerencia", "Gate / aprobacion PM",
    ]
    sprint_rows = []
    for s in SPRINT_ENTREGABLES:
        if s[1] != sprint_label:
            continue
        sprint_rows.append([s[0], s[2], s[3], s[4], s[5], s[6], s[7], s[8]])
    row = write_table(ws, sprint_headers, sprint_rows, row, fill)
    row += 1

    # ── 5. Riesgos y mitigacion ──
    row = write_section(
        ws, row,
        "5. RIESGOS IDENTIFICADOS Y PLAN DE MITIGACION (SI SE MATERIALIZAN)",
        ncol, fill,
        anchors, "5. Riesgos y mitigacion",
    )
    risk_headers = [
        "ID", "Impacto", "Riesgo identificado", "Consecuencia si ocurre",
        "Acciones de mitigacion / contingencia", "Responsable", "Vigilancia (sprints)",
    ]
    row = write_table(ws, risk_headers, risks, row, FILLS["risk"])

    ws.freeze_panes = "A4"
    set_col_widths(ws, [6, 14, 32, 38, 14, 36, 8, 8, 8, 11, 11, 14, 18, 8, 10, 8])
    for a in anchors:
        a["sheet"] = sheet_name
    return anchors


def cell_link(sheet: str, row: int) -> str:
    """Enlace interno Excel a celda A{row} de la hoja."""
    safe = sheet.replace("'", "''")
    return f"#'{safe}'!A{row}"


def set_hyperlink(cell, location: str, display: str | None = None) -> None:
    cell.value = display or "Ir a seccion"
    cell.hyperlink = location
    cell.font = LINK_FONT
    cell.alignment = WRAP


def build_index_sheet(ws, etapa1_anchors: list, etapa2_anchors: list) -> None:
    ws.title = "Indice"
    ncol = 5
    row = 1
    ws.cell(row=row, column=1, value="INDICE — PLAN DETALLADO ETAPAS 1 Y 2 (AURIXA v36)").font = Font(
        bold=True, size=16, color="1F4E79",
    )
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    row = 3
    ws.cell(row=row, column=1, value=(
        "Use la columna ENLACE para saltar a cada seccion. Tambien puede usar las pestañas inferiores del libro."
    )).alignment = WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    row += 2

    headers = ["Etapa", "Seccion", "Descripcion breve", "Celda", "Enlace"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = FILLS["title"]
        cell.alignment = HEADER_ALIGN
        cell.border = THIN
    row += 1

    def add_group(etapa_name: str, anchors: list[dict], desc_default: str):
        nonlocal row
        for a in anchors:
            loc = cell_link(a["sheet"], a["row"])
            ws.cell(row=row, column=1, value=etapa_name).border = THIN
            ws.cell(row=row, column=2, value=a["label"]).border = THIN
            ws.cell(row=row, column=3, value=a.get("section", desc_default)[:100]).border = THIN
            ws.cell(row=row, column=3).alignment = WRAP
            ws.cell(row=row, column=4, value=f"A{a['row']}").border = THIN
            link_cell = ws.cell(row=row, column=5)
            set_hyperlink(link_cell, loc, "Abrir >>")
            link_cell.border = THIN
            row += 1

    add_group("ETAPA 1", etapa1_anchors, "Jun–Sep 2026 · S1–S8")
    row += 1
    add_group("ETAPA 2", etapa2_anchors, "Oct–Nov 2026 · S9–S13")

    row += 1
    row = write_section(ws, row, "DOCUMENTOS RELACIONADOS (mismo proyecto)", ncol, FILLS["ok"])
    ext_headers = ["Tipo", "Archivo", "Uso", "", "Nota"]
    for c, h in enumerate(ext_headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = FILLS["ok"]
        cell.border = THIN
    row += 1
    related = [
        ("ClickUp", "Plataforma_Minera_ClickUp_v19.xlsx", "Importar cronograma — hoja Import_ClickUp", "482 tareas con Sprint"),
        ("Gerencia", "docs/Plan_Proyecto_Gerencia_Etapas_Sprints_v36.xlsx", "Resumen gates y 11 hojas PMO", "Revision ejecutiva"),
        ("Recursos", "docs/Matriz_Tareas_Recursos_v36_Nueva_Distribucion.md", "Matriz carga 10 recursos", "Validacion PMO"),
    ]
    for tipo, archivo, uso, nota in related:
        ws.cell(row=row, column=1, value=tipo).border = THIN
        ws.cell(row=row, column=2, value=archivo).border = THIN
        ws.cell(row=row, column=3, value=uso).border = THIN
        ws.cell(row=row, column=4, value="").border = THIN
        ws.cell(row=row, column=5, value=nota).border = THIN
        row += 1

    ws.freeze_panes = "A6"
    set_col_widths(ws, [12, 32, 48, 8, 14])


def main() -> int:
    print("Cargando cronograma maestro...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet("Etapa_1_Implementacion")
    anchors1 = build_etapa_sheet(ws1, {
        "banner": "ETAPA 1 — IMPLEMENTACION FUNCIONAL CORE | 4 MESES (JUNIO–SEPTIEMBRE 2026) | SPRINTS S1–S8 | RELEASES R1–R4",
        "months": [6, 7, 8, 9],
        "paso": "PASO 1",
        "fill": FILLS["etapa1"],
        "risks": ETAPA1_RIESGOS,
        "sprint_etapa": "Etapa 1",
        "focus": ETAPA1_FOCUS,
    })

    ws2 = wb.create_sheet("Etapa_2_GoLive")
    anchors2 = build_etapa_sheet(ws2, {
        "banner": "ETAPA 2 — HARDENING, ESTABILIZACION Y GO-LIVE | 2 MESES (OCTUBRE–NOVIEMBRE 2026) | SPRINTS S9–S13 | RELEASES R5–R6",
        "months": [10, 11],
        "paso": "PASO 2",
        "fill": FILLS["etapa2"],
        "risks": ETAPA2_RIESGOS,
        "sprint_etapa": "Etapa 2",
        "focus": ETAPA2_FOCUS,
    })

    ws_idx = wb.create_sheet("Indice", 0)
    build_index_sheet(ws_idx, anchors1, anchors2)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Excel generado: {OUT}")
    print(f"  Pestana Indice: {len(anchors1) + len(anchors2)} enlaces internos")
    print("  Pestana Etapa_1_Implementacion (Jun-Sep, S1-S8)")
    print("  Pestana Etapa_2_GoLive (Oct-Nov, S9-S13)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
