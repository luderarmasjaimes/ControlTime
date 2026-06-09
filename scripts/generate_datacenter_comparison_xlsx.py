from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
CSV_PATH = ROOT / "docs" / "Plantilla_Comparacion_Economica_Datacenter_AURIXA.csv"
XLSX_PATH = ROOT / "docs" / "Plantilla_Comparacion_Economica_Datacenter_AURIXA.xlsx"


GREEN_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFEB9C")
RED_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")


def load_rows() -> list[list[str]]:
    with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def autosize_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is not None:
                length = max(length, len(str(cell.value)))
        worksheet.column_dimensions[column_letter].width = min(max(length + 2, 12), 28)


def build_input_sheet(workbook: Workbook, rows: list[list[str]]) -> None:
    ws = workbook.active
    ws.title = "Comparacion"

    for row in rows:
        ws.append(row)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    autosize_columns(ws)

    currency_cols = [22, 23, 24]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for idx in currency_cols:
            row[idx - 1].number_format = '"USD" #,##0.00'


def build_summary_sheet(workbook: Workbook) -> None:
    ws = workbook.create_sheet("Resumen")
    headers = [
        "Proveedor",
        "Costo Mensual USD",
        "Costo Setup USD",
        "Costo Anual USD",
        "SLA",
        "RTO",
        "Semaforo Costo",
        "Semaforo SLA",
        "Semaforo RTO",
        "Puntaje Costo",
        "Puntaje SLA",
        "Puntaje RTO",
        "Puntaje Total",
        "Ranking",
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    providers = ["Cirion", "GTD", "Claro Empresas", "WIN Empresas", "Equinix / Partner"]
    for index, provider in enumerate(providers, start=2):
        comp_row = index
        ws.cell(index, 1, provider)
        ws.cell(index, 2, f"=Comparacion!V{comp_row}")
        ws.cell(index, 3, f"=Comparacion!W{comp_row}")
        ws.cell(index, 4, f"=Comparacion!X{comp_row}")
        ws.cell(index, 5, f"=Comparacion!J{comp_row}")
        ws.cell(index, 6, f"=Comparacion!K{comp_row}")
        ws.cell(index, 7, f'=IF(B{index}="","PENDIENTE",IF(B{index}<=2500,"VERDE",IF(B{index}<=4000,"AMARILLO","ROJO")))')
        ws.cell(index, 8, f'=IF(E{index}="","PENDIENTE",IF(E{index}>=0.9995,"VERDE",IF(E{index}>=0.999,"AMARILLO","ROJO")))')
        ws.cell(index, 9, f'=IF(F{index}="","PENDIENTE",IF(F{index}<=5,"VERDE",IF(F{index}<=15,"AMARILLO","ROJO")))')
        ws.cell(index, 10, f'=IF(G{index}="VERDE",100,IF(G{index}="AMARILLO",60,IF(G{index}="ROJO",20,"")))')
        ws.cell(index, 11, f'=IF(H{index}="VERDE",100,IF(H{index}="AMARILLO",60,IF(H{index}="ROJO",20,"")))')
        ws.cell(index, 12, f'=IF(I{index}="VERDE",100,IF(I{index}="AMARILLO",60,IF(I{index}="ROJO",20,"")))')
        ws.cell(index, 13, f'=IF(COUNT(J{index}:L{index})<3,"",ROUND(J{index}*0.35+K{index}*0.40+L{index}*0.25,2))')
        ws.cell(index, 14, f'=IF(M{index}="","",RANK(M{index},$M$2:$M$6,0))')

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = '"USD" #,##0.00'

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=10, max_col=13):
        for cell in row:
            cell.number_format = '0.00'

    autosize_columns(ws)
    ws.freeze_panes = "A2"

    for col in (7, 8, 9):
        letter = get_column_letter(col)
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{ws.max_row}",
            CellIsRule(operator="equal", formula=['"VERDE"'], fill=GREEN_FILL),
        )
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{ws.max_row}",
            CellIsRule(operator="equal", formula=['"AMARILLO"'], fill=YELLOW_FILL),
        )
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{ws.max_row}",
            CellIsRule(operator="equal", formula=['"ROJO"'], fill=RED_FILL),
        )

    letter = get_column_letter(13)
    ws.conditional_formatting.add(
        f"{letter}2:{letter}{ws.max_row}",
        CellIsRule(operator="greaterThanOrEqual", formula=["85"], fill=GREEN_FILL),
    )
    ws.conditional_formatting.add(
        f"{letter}2:{letter}{ws.max_row}",
        CellIsRule(operator="between", formula=["60", "84.99"], fill=YELLOW_FILL),
    )
    ws.conditional_formatting.add(
        f"{letter}2:{letter}{ws.max_row}",
        CellIsRule(operator="lessThan", formula=["60"], fill=RED_FILL),
    )


def build_criteria_sheet(workbook: Workbook) -> None:
    ws = workbook.create_sheet("Criterios")
    rows = [
        ["Criterio", "Verde", "Amarillo", "Rojo"],
        ["Costo mensual USD", "<= 2500", "<= 4000", "> 4000"],
        ["SLA disponibilidad", ">= 99.95%", ">= 99.90%", "< 99.90%"],
        ["RTO minutos", "<= 5", "<= 15", "> 15"],
        ["Peso costo", "35%", "", ""],
        ["Peso SLA", "40%", "", ""],
        ["Peso RTO", "25%", "", ""],
        ["Nota", "Ajusta estos umbrales si la cotizacion real usa otro nivel de servicio o alcance.", "", ""],
    ]
    for row in rows:
        ws.append(row)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for cell in ws[8]:
        cell.fill = SECTION_FILL
        cell.font = Font(bold=True)

    autosize_columns(ws)


def main() -> None:
    rows = load_rows()
    workbook = Workbook()
    build_input_sheet(workbook, rows)
    build_summary_sheet(workbook)
    build_criteria_sheet(workbook)
    workbook.save(XLSX_PATH)
    print(f"Generated {XLSX_PATH}")


if __name__ == "__main__":
    main()
