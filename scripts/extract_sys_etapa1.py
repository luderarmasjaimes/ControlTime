#!/usr/bin/env python3
"""Extrae tareas SYS de Etapa 1 del cronograma ClickUp."""
import csv
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(r"c:\InformeCliente")
XLSX = ROOT / "Nueva_Plataforma_Minera_IA_ClickUp.xlsx"
OUT_MD = ROOT / "docs" / "SYS_Etapa1_Tareas_Revision.md"
OUT_CSV = ROOT / "docs" / "SYS_Etapa1_Tareas_Revision.csv"


def parse_codes(recurso: str) -> list[str]:
    r = str(recurso or "").upper()
    codes: list[str] = []
    mapping = [
        ("BACKEND 1", "BE1"),
        ("BACKEND 2", "BE2"),
        ("BACKEND 3", "BE3"),
        ("FRONTEND 1", "FE1"),
        ("FRONTEND 2", "FE2"),
        ("IA-ML", "IA"),
        ("ARQ", "ARQ"),
        ("SYS", "SYS"),
        ("QA", "QA"),
    ]
    for key, code in mapping:
        if key in r and code not in codes:
            codes.append(code)
    return codes


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["Import_ClickUp"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}

    rows_sys = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        tid = row[idx["Task ID"]]
        if not tid or str(tid).startswith("CAP-"):
            continue
        folder = str(row[idx["Folder Name"]] or "")
        if "Etapa 1" not in folder:
            continue
        codes = parse_codes(row[idx["Recurso"]])
        if "SYS" not in codes:
            continue
        rows_sys.append(
            {
                "id": str(tid),
                "name": str(row[idx["Task Name"]] or ""),
                "list": str(row[idx["List Name"]] or ""),
                "start": str(row[idx["Start Date"]] or ""),
                "due": str(row[idx["Due Date"]] or ""),
                "hours": row[idx["Time Estimate"]] or 0,
                "priority": str(row[idx["Priority"]] or ""),
                "recurso": str(row[idx["Recurso"]] or ""),
                "codes": ",".join(codes),
                "solo_sys": codes == ["SYS"],
                "stream": str(row[idx["Stream"]] or ""),
                "depends": str(row[idx["Depends On"]] or ""),
                "tags": str(row[idx["Tags"]] or ""),
            }
        )

    rows_sys.sort(key=lambda x: (x["start"], x["list"], x["id"]))
    solo = [r for r in rows_sys if r["solo_sys"]]
    shared = [r for r in rows_sys if not r["solo_sys"]]
    total_h = sum(r["hours"] for r in rows_sys)
    solo_h = sum(r["hours"] for r in solo)

    OUT_MD.parent.mkdir(exist_ok=True)
    lines = [
        "# Tareas SYS — Etapa 1 (Plataforma Reportabilidad)",
        "",
        f"**Fuente:** `{XLSX.name}` (hoja Import_ClickUp)",
        "**Periodo etapa:** Jun–Nov 2026",
        f"**Total tareas con SYS:** {len(rows_sys)} ({len(solo)} solo SYS, {len(shared)} compartidas)",
        f"**Horas estimadas:** {total_h} h (solo SYS: {solo_h} h)",
        "",
        "---",
        "",
    ]

    by_list = defaultdict(list)
    for r in rows_sys:
        by_list[r["list"]].append(r)

    for lst in sorted(by_list.keys()):
        items = by_list[lst]
        h = sum(x["hours"] for x in items)
        lines.append(f"## {lst} ({len(items)} tareas, {h} h)")
        lines.append("")
        lines.append("| ID | Tarea | Inicio | Fin | h | Asignacion | Prioridad |")
        lines.append("|---|---|---|---|---:|---|---|")
        for r in items:
            name = r["name"].replace("|", "/")
            lines.append(
                f"| {r['id']} | {name} | {r['start']} | {r['due']} | {r['hours']} | {r['recurso']} | {r['priority']} |"
            )
        lines.append("")

    if shared:
        lines.extend(["---", "", "## Tareas compartidas (SYS + otros recursos)", ""])
        for r in shared:
            lines.append(f"- **{r['id']}** — {r['name']} ({r['codes']}) — {r['hours']} h")
        lines.append("")

    OUT_MD.write_text("\n".join(lines), encoding="utf-8")

    # Archivos separados: solo SYS vs compartidas
    solo_md = OUT_MD.parent / "SYS_Etapa1_Solo_SYS.md"
    shared_md = OUT_MD.parent / "SYS_Etapa1_Compartidas.md"

    def write_split_md(path, title, items):
        h_total = sum(x["hours"] for x in items)
        out = [
            f"# {title}",
            "",
            f"**Total:** {len(items)} tareas | **Horas:** {h_total} h",
            "",
            "| ID | Tarea | Lista | Inicio | Fin | h | Con quien | Prioridad |",
            "|---|---|---|---|---|---:|---|---|",
        ]
        for r in items:
            otros = r["recurso"] if r["solo_sys"] else r["recurso"].replace("SYS", "").replace(", ,", ",").strip(" ,")
            if r["solo_sys"]:
                comp = "— (solo SYS)"
            else:
                comp = ", ".join(c for c in r["codes"].split(",") if c != "SYS") or r["recurso"]
            name = r["name"].replace("|", "/")
            out.append(
                f"| {r['id']} | {name} | {r['list']} | {r['start']} | {r['due']} | {r['hours']} | {comp} | {r['priority']} |"
            )
        path.write_text("\n".join(out), encoding="utf-8")

    write_split_md(solo_md, "SYS Etapa 1 — Tareas SOLO SYS", solo)
    write_split_md(shared_md, "SYS Etapa 1 — Tareas COMPARTIDAS con otros recursos", shared)

    solo_csv = OUT_MD.parent / "SYS_Etapa1_Solo_SYS.csv"
    shared_csv = OUT_MD.parent / "SYS_Etapa1_Compartidas.csv"
    fields = [
        "id", "name", "list", "start", "due", "hours", "priority",
        "recurso", "codes", "stream", "depends", "tags",
    ]
    for path, items in ((solo_csv, solo), (shared_csv, shared)):
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
            w.writeheader()
            w.writerows(items)

    with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(
            f,
            fieldnames=[
                "id", "name", "list", "start", "due", "hours", "priority",
                "recurso", "codes", "solo_sys", "stream", "depends", "tags",
            ],
        )
        w.writeheader()
        w.writerows(rows_sys)

    print(f"Total: {len(rows_sys)} | Solo SYS: {len(solo)} | Compartidas: {len(shared)} | Horas: {total_h}")
    print(f"MD:  {OUT_MD}")
    print(f"Solo SYS:    {solo_md} + {solo_csv}")
    print(f"Compartidas: {shared_md} + {shared_csv}")
    print(f"CSV completo: {OUT_CSV}")
    for lst in sorted(by_list.keys()):
        print(f"  {lst}: {len(by_list[lst])} tareas, {sum(x['hours'] for x in by_list[lst])} h")


if __name__ == "__main__":
    main()
