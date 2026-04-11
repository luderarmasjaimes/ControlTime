# -*- coding: utf-8 -*-
"""Genera DOCX de presentación para gerencia TI y usuario empresa (InformeCliente + referencia ThingsBoard)."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, fill_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, headers, rows, header_fill="1F4E79"):
    ncols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(hdr_cells[i], header_fill)
    for ri, row in enumerate(rows):
        row_cells = table.rows[ri + 1].cells
        for ci, text in enumerate(row):
            row_cells[ci].text = str(text)
            for p in row_cells[ci].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
    doc.add_paragraph()


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out = root / "docs" / "MATRIZ_FUNCIONALIDADES_PRESENTACION_GERENCIA_2026-04-10.docx"

    doc = Document()
    for s in doc.sections:
        s.top_margin = Pt(56.7 * 4)
        s.bottom_margin = Pt(56.7 * 4)
        s.left_margin = Pt(56.7 * 5)
        s.right_margin = Pt(56.7 * 5)

    title = doc.add_heading(
        "Plataforma InformeCliente — Matriz funcional para presentación",
        level=0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Estado implementado, brechas y alineación con referencia IoT (ThingsBoard)\n"
        "Fecha de documento: abril 2026"
    )
    r.font.size = Pt(11)
    r.italic = True

    doc.add_paragraph()

    doc.add_heading("1. Contexto ejecutivo", level=1)
    doc.add_paragraph(
        "InformeCliente es la plataforma vertical de la operación minera: autenticación "
        "(incluida verificación facial asistida por IA), monitoreo, cartografía, motor de "
        "fórmulas, informes técnicos y despliegue en contenedores Docker."
    )
    doc.add_paragraph(
        "ThingsBoard (código de referencia en C:\\thingsboard-master) es una plataforma IoT "
        "de mercado: dispositivos y activos, telemetría, dashboards, motor de reglas, alarmas, "
        "notificaciones y modelo multi-tenant. No sustituye los informes corporativos ni el "
        "flujo biométrico propio; sirve como referencia para decidir qué capacidades IoT/operación "
        "se desean igualar, integrar o mantener en desarrollo propio."
    )

    doc.add_heading("2. Funcionalidades actuales (por módulo)", level=1)
    h1 = ["#", "Qué ofrece al usuario / negocio", "Módulo / componente", "Descripción", "Estado"]
    r1 = [
        [
            "1",
            "Acceso con usuario y contraseña",
            "Frontend AuthGateway · Backend /api/auth/login/password",
            "Inicio de sesión estándar vinculado a empresa/usuario.",
            "Implementado",
        ],
        [
            "2",
            "Acceso con reconocimiento facial",
            "Frontend auth · Backend login/registro facial · Servicio ai_engine",
            "Registro y login con plantilla facial y validaciones asistidas por IA.",
            "Implementado",
        ],
        [
            "3",
            "Alta de usuario y validación de empresa",
            "Backend register/validate-company · BD auth",
            "Registro y comprobación de datos de compañía.",
            "Implementado",
        ],
        [
            "4",
            "Catálogo de empresas para auth",
            "Backend /api/auth/companies",
            "Soporte a flujos de login por compañía.",
            "Implementado",
        ],
        [
            "5",
            "Auditoría de accesos",
            "Frontend AuditCenter · Backend /api/auth/audit (+ export CSV)",
            "Traza de eventos de autenticación para cumplimiento y TI.",
            "Implementado",
        ],
        [
            "6",
            "Panel ejecutivo y KPIs",
            "MiningDashboard · /api/dashboard/metrics",
            "Resumen operativo y métricas agregadas.",
            "Implementado (revisar datos demo/fallback por entorno)",
        ],
        [
            "7",
            "Telemetría y sensores",
            "AdvancedSensors · /api/sensors/data · Timescale",
            "Visualización de series y estado de instrumentación.",
            "Implementado / parcial unificación multi-sitio",
        ],
        [
            "8",
            "Mapa base y mapa HD",
            "MapViewer, DetailedMap · tileserver · /api/map/markers",
            "Mapas y capas MBTiles de alta resolución.",
            "Implementado",
        ],
        [
            "9",
            "Captura desde mapa hacia informe",
            "Mapas + ReportStudio",
            "Inserción de imagen de mapa en informe.",
            "Implementado (evidencia como activo auditable: mejorable)",
        ],
        [
            "10",
            "Conversión geoespacial (p. ej. ECW → MBTiles)",
            "Backend /api/convert, /api/jobs/{id}",
            "Jobs de conversión para alimentar el visor.",
            "Implementado",
        ],
        [
            "11",
            "Vigilancia / cámaras",
            "VideoDiagram · /api/surveillance/cameras",
            "Referencia a fuentes de video en el tablero.",
            "Implementado",
        ],
        [
            "12",
            "Vista 3D / gemelo",
            "Viewer3D",
            "Visualización 3D del contexto minero.",
            "Implementado",
        ],
        [
            "13",
            "Gráficos geotécnicos",
            "Inclinómetros, desplazamientos, azimuth",
            "Análisis de estabilidad y movimiento.",
            "Implementado",
        ],
        [
            "14",
            "Motor de fórmulas mineras",
            "FormulaEngineEmbed · formula_engine · /api/formula y /api/analysis",
            "Catálogos, análisis (p. ej. temperaturas), flujo de ingeniería.",
            "Implementado",
        ],
        [
            "15",
            "Informes técnicos y corporativos",
            "Report, ReportStudioV2 · CRUD /api/reports",
            "Editor multipágina, plantillas, bloques; ciclo de vida en evolución.",
            "Implementado con brechas multiempresa y metadatos",
        ],
        [
            "16",
            "Mantenimiento de usuarios en UI",
            "UserMaintenanceModal (y almacenamiento local en parte)",
            "Administración de usuarios de empresa desde la aplicación.",
            "Parcial — cerrar con API y auditoría central",
        ],
        [
            "17",
            "Despliegue Docker",
            "docker-compose: db, web, frontend, ai_engine, formula_*, tileserver",
            "Entorno reproducible para TI.",
            "Implementado",
        ],
    ]
    add_table(doc, h1, r1)

    doc.add_heading("3. Necesidades y brechas (prioridad)", level=1)
    h2 = ["#", "Objetivo", "Módulo", "Descripción", "Prioridad"]
    r2 = [
        [
            "1",
            "Gobierno multiempresa en informes y datos",
            "Backend web + PostgreSQL",
            "Filtrado estricto por tenant/empresa; metadatos created_by, revisión, versión.",
            "P0",
        ],
        [
            "2",
            "Usuarios administrables de forma central",
            "Auth API + ReportStudio",
            "Listado, mantenimiento y auditoría sin depender de localStorage para lo crítico.",
            "P0",
        ],
        [
            "3",
            "Ciclo de vida documental",
            "Backend + ReportStudioV2",
            "Estados (borrador, revisión, aprobado), roles e historial.",
            "P1",
        ],
        [
            "4",
            "Compartir informes con registro",
            "Backend + report_shares / notificaciones",
            "Compartición persistente y trazable, no solo simulada en UI.",
            "P1",
        ],
        [
            "5",
            "Evidencias de mapa auditables",
            "Backend + almacenamiento",
            "Captura con usuario, tiempo, contexto geográfico y vínculo al informe.",
            "P2",
        ],
        [
            "6",
            "API unificada de telemetría",
            "Backend + Timescale",
            "Filtros por sitio/tenant, paginación y rangos (patrón tipo Device/Asset).",
            "P1–P2",
        ],
        [
            "7",
            "Hardening seguridad",
            "Infra, secretos, TLS",
            "Sin secretos en repositorio; gestión formal de credenciales y certificados.",
            "P0",
        ],
        [
            "8",
            "Contrato API y pruebas E2E",
            "Transversal",
            "OpenAPI y pruebas end-to-end sobre backend real.",
            "P1",
        ],
    ]
    add_table(doc, h2, r2, header_fill="2E75B6")

    doc.add_heading("4. Alineación con ThingsBoard (referencia)", level=1)
    doc.add_paragraph(
        "ThingsBoard aporta patrones de producto IoT. La columna derecha indica la orientación "
        "en InformeCliente (desarrollo propio o integración futura)."
    )
    h3 = [
        "Capacidad (ThingsBoard)",
        "Equivalente hoy en InformeCliente",
        "Línea de acción propuesta",
    ]
    r3 = [
        [
            "Tenant / Customer / jerarquía",
            "Empresa y sesión en auth; esquema telemetría en BD",
            "Modelo explícito tenant → sitio → activo/sensor en API y UI; alinear con informes",
        ],
        [
            "Device / Asset y relaciones",
            "Sensores y proyectos vía API y SQL",
            "Catálogo gobernado y relaciones (sensor ↔ frente ↔ informe)",
        ],
        [
            "Telemetría y almacenamiento",
            "/api/sensors/data, Timescale",
            "Normalización, retención, calidad; opción puente MQTT/HTTP o integración TB",
        ],
        [
            "Dashboards tiempo real",
            "MiningDashboard y gráficos propios",
            "Widgets avanzados o embed de dashboards TB si se adopta como capa IoT",
        ],
        [
            "Rule engine y alarmas",
            "Sin equivalente completo",
            "Motor de umbrales/escalamiento propio o TB dedicado a alarmística",
        ],
        [
            "Notificaciones",
            "Esquema en evolución para informes",
            "Política única: TB para IoT + backend propio para documentos, o bus común",
        ],
        [
            "SCADA / control industrial",
            "No como producto genérico",
            "Solo si negocio lo exige; TB como referencia de UI/alarma",
        ],
        [
            "Provisioning de dispositivos",
            "Limitado frente a TB",
            "Definir fuente de verdad: InformeCliente, TB o modelo híbrido",
        ],
    ]
    add_table(doc, h3, r3, header_fill="375623")

    doc.add_heading("5. Resumen: usuario de empresa vs gerencia TI", level=1)
    h4 = ["Usuario de la empresa", "Gerencia TI"]
    r4 = [
        [
            "Una sola entrada: login seguro y, si aplica, reconocimiento facial.",
            "Rutas de auth auditables; revisar secretos y certificados en despliegue.",
        ],
        [
            "Tableros, sensores, mapas e informes en un mismo entorno.",
            "Stack Docker modular; plan de escala, backups y monitoreo.",
        ],
        [
            "Informes técnicos y corporativos desde el navegador.",
            "Cerrar multiempresa y ciclo de vida documental en backend antes de producción fuerte.",
        ],
        [
            "Motor de fórmulas para análisis de ingeniería.",
            "formula_engine + PostgreSQL dedicado; continuidad operativa y respaldos.",
        ],
        [
            "Mapas de alta resolución y conversiones geoespaciales.",
            "tileserver, volúmenes de datos y jobs de conversión bajo supervisión.",
        ],
    ]
    add_table(doc, h4, r4, header_fill="7030A0")

    doc.add_heading("6. Decisión estratégica sugerida para la reunión", level=1)
    doc.add_paragraph(
        "Integrar ThingsBoard como capa IoT (telemetría, alarmas, dashboards de dispositivos) "
        "manteniendo InformeCliente como capa de negocio, informes y biométrica — frente a — "
        "replicar solo patrones de ThingsBoard sin desplegarlo, ampliando reglas y "
        "telemetría en el backend propio. La tabla 4 ayuda a alinear expectativas entre "
        "negocio y TI."
    )

    doc.save(out)
    print(f"Escrito: {out}")


if __name__ == "__main__":
    main()
