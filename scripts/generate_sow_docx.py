from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE_COLOR = RGBColor(14, 61, 87)
ACCENT_COLOR = RGBColor(12, 104, 102)
BODY_COLOR = RGBColor(34, 34, 34)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def clean_inline_markdown(text: str) -> str:
    text = text.replace("**", "")
    text = text.replace("__", "")
    return text.strip()


def configure_styles(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY_COLOR

    for style_name, size, color in [
        ("Title", 22, TITLE_COLOR),
        ("Heading 1", 15, TITLE_COLOR),
        ("Heading 2", 12.5, ACCENT_COLOR),
        ("Heading 3", 11.5, TITLE_COLOR),
    ]:
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    if "Quote" in document.styles:
        quote = document.styles["Quote"]
        quote.font.name = "Calibri"
        quote.font.size = Pt(10)
        quote.font.italic = True

    if "Table Grid" in document.styles:
        table_style = document.styles["Table Grid"]
        table_style.font.name = "Calibri"
        table_style.font.size = Pt(9.5)

    if "SOW Bullet" not in document.styles:
        style = document.styles.add_style("SOW Bullet", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)


def add_cover(document: Document, title: str, subtitle: str, source_name: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = TITLE_COLOR

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(subtitle)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = ACCENT_COLOR

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(14)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(f"Documento generado para presentacion ejecutiva\nFuente: {source_name}")
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = BODY_COLOR

    document.add_paragraph()
    document.add_paragraph()
    document.add_section(WD_SECTION.NEW_PAGE)


def parse_table(lines: list[str], start_index: int):
    rows = []
    index = start_index
    while index < len(lines):
        line = lines[index].strip()
        if not line.startswith("|"):
            break
        parts = [cell.strip() for cell in line.strip("|").split("|")]
        rows.append(parts)
        index += 1

    if len(rows) >= 2 and all(set(cell) <= {"-", ":"} for cell in rows[1]):
        rows.pop(1)

    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=column_count)
    table.style = "Table Grid"

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index in range(column_count):
            value = clean_inline_markdown(row[col_index]) if col_index < len(row) else ""
            cells[col_index].text = value
            cells[col_index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cells[col_index].paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9.5)
        if row_index == 0:
            for cell in cells:
                set_cell_shading(cell, "DCE6F1")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

    document.add_paragraph()


def add_numbered_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    run = paragraph.add_run(clean_inline_markdown(re.sub(r"^\d+\.\s+", "", text)))
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)


def add_bullet_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    run = paragraph.add_run(clean_inline_markdown(re.sub(r"^-\s+", "", text)))
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(clean_inline_markdown(text))
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)


def render_markdown(document: Document, markdown_text: str) -> None:
    lines = markdown_text.splitlines()
    index = 0
    title = ""
    subtitle = ""

    while index < len(lines):
        stripped = lines[index].strip()
        if stripped.startswith("# ") and not title:
            title = clean_inline_markdown(stripped[2:])
        elif stripped.startswith("## ") and not subtitle:
            subtitle = clean_inline_markdown(stripped[3:])
            break
        index += 1

    add_cover(document, title or "Documento Ejecutivo", subtitle or "Presentacion", "markdown fuente")

    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            index += 1
            continue

        if stripped.startswith("# "):
            document.add_paragraph(clean_inline_markdown(stripped[2:]), style="Title")
            index += 1
            continue
        if stripped.startswith("## "):
            document.add_paragraph(clean_inline_markdown(stripped[3:]), style="Heading 1")
            index += 1
            continue
        if stripped.startswith("### "):
            document.add_paragraph(clean_inline_markdown(stripped[4:]), style="Heading 2")
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        if re.match(r"^\d+\.\s+", stripped):
            add_numbered_paragraph(document, stripped)
            index += 1
            continue
        if stripped.startswith("- "):
            add_bullet_paragraph(document, stripped)
            index += 1
            continue
        add_paragraph(document, stripped)
        index += 1


def build_document(source_path: Path, output_path: Path) -> None:
    document = Document()
    configure_styles(document)
    markdown_text = source_path.read_text(encoding="utf-8")
    render_markdown(document, markdown_text)

    footer = document.sections[-1].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("AURIXA 2026 - Documento de presentacion ejecutiva")
    footer_run.font.size = Pt(8.5)
    footer_run.font.color.rgb = ACCENT_COLOR

    document.save(output_path)


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print("Uso: python scripts/generate_sow_docx.py <input.md> <output.docx>")
        return 1

    source_path = Path(argv[1])
    output_path = Path(argv[2])
    build_document(source_path, output_path)
    print(f"DOCX generado: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))