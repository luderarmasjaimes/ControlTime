#!/usr/bin/env python3
"""
Genera Excel Etapa 1 — Plataforma Reportabilidad (Jun–Nov 2026)
listo para importacion ClickUp al 100%.

Requisitos previos: ejecuta generate_clickup_import.py (se invoca automaticamente).
Salida:
  Etapa1_Plataforma_Reportabilidad_ClickUp.xlsx
  Etapa1_Plataforma_Reportabilidad_ClickUp_Import.csv
"""
from __future__ import annotations

import csv
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"c:\InformeCliente")
SOURCE = ROOT / "Plataforma_Minera_ClickUp_v19.xlsx"
GENERATOR = ROOT / "scripts" / "generate_clickup_import.py"
OUTPUT_XLSX = ROOT / "Etapa1_Plataforma_Reportabilidad_ClickUp.xlsx"
OUTPUT_CSV = ROOT / "Etapa1_Plataforma_Reportabilidad_ClickUp_Import.csv"

SPACE_NAME = "PLATAFORMA MINERA IA"
ETAPA1_FOLDER = "Etapa 1 - Plataforma Reportabilidad"

ETAPA1_LISTS = [
    "MES1_JUNIO_Analisis_Arquitectura",
    "MES2_JULIO_Core_Seguridad",
    "MES3_AGOSTO_UX_Sensores",
    "MES4_SEPTIEMBRE_Integracion_QA",
    "MES5_OCTUBRE_Optimizacion",
    "MES6_NOVIEMBRE_GoLive",
    "PMO_Seguimiento",
    "QA_Transversal",
    "DevOps_Infra",
    "Gestion_Riesgos",
    "Capacidad_Recursos",
]

ETAPA1_HEADERS = [
    "Task Name",
    "Task Description",
    "Folder Name",
    "List Name",
    "Status",
    "Priority",
    "Start Date",
    "Due Date",
    "Tags",
    "Time Estimate",
    "Task ID",
    "Recurso",
    "Release",
    "Depends On",
    "Complejidad",
    "Riesgo",
    "Stream",
    "Ambiente",
    "Sprint",
    "Mes Calendario",
    "Carga Mes Porcentaje",
    "Tipo Metodologia",
    "Secuencia Import",
]

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def regenerate_source() -> None:
    subprocess.run([sys.executable, str(GENERATOR)], check=True, cwd=str(ROOT))


def load_source_rows():
    wb = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb["Import_ClickUp"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[idx["Task ID"]]:
            continue
        rows.append(dict(zip(headers, row)))
    wb.close()
    return rows


def validate_rows(rows: list[dict]) -> dict:
    report = {"ok": True, "errors": [], "warnings": [], "stats": {}}
    ids = [r["Task ID"] for r in rows]
    dupes = sorted({i for i in ids if ids.count(i) > 1})
    if dupes:
        report["ok"] = False
        report["errors"].append(f"IDs duplicados: {dupes[:8]}")

    idset = set(ids)
    broken = []
    for r in rows:
        dep = (r.get("Depends On") or "").strip()
        if not dep:
            continue
        for d in dep.split(","):
            d = d.strip()
            if d and d not in idset:
                broken.append(f"{r['Task ID']} -> {d}")
    if broken:
        report["ok"] = False
        report["errors"].append(f"Dependencias rotas: {len(broken)}")
        report["errors"].extend(broken[:6])

    missing = [r["Task ID"] for r in rows if not (r.get("Task Name") or "").strip()]
    if missing:
        report["ok"] = False
        report["errors"].append(f"Tareas sin nombre: {missing[:5]}")

    bad_lists = sorted({r["List Name"] for r in rows if r["List Name"] not in ETAPA1_LISTS})
    if bad_lists:
        report["warnings"].append(f"Listas no estandar Etapa 1: {bad_lists}")

    no_dates = [r["Task ID"] for r in rows if not r.get("Start Date") or not r.get("Due Date")]
    if no_dates:
        report["ok"] = False
        report["errors"].append(f"Sin fechas: {no_dates[:5]}")

    report["stats"] = {
        "tasks": len(rows),
        "floor_tasks": sum(1 for r in rows if str(r["Task ID"]).startswith("FLOOR-")),
        "cap_tasks": sum(1 for r in rows if str(r["Task ID"]).startswith("CAP-")),
        "tech_tasks": sum(
            1 for r in rows
            if not str(r["Task ID"]).startswith(("CAP-", "FLOOR-"))
        ),
    }
    return report


def to_etapa1_row(src: dict, seq: int) -> list:
    return [
        src["Task Name"],
        src["Task Description"],
        ETAPA1_FOLDER,
        src["List Name"],
        src.get("Status") or "Backlog",
        src.get("Priority") or "Normal",
        src["Start Date"],
        src["Due Date"],
        src.get("Tags") or "",
        src.get("Time Estimate") or 0,
        src["Task ID"],
        src.get("Recurso") or "",
        src.get("Release") or "",
        (src.get("Depends On") or "").strip(),
        src.get("Complejidad") or "Media",
        src.get("Riesgo") or "Medio",
        src.get("Stream") or "",
        src.get("Ambiente") or "Desarrollo",
        src.get("Sprint") or "",
        src.get("Mes Calendario") or "",
        src.get("Carga Mes Porcentaje") or 0,
        src.get("Tipo Metodologia") or "",
        seq,
    ]


def write_sheet_headers(ws, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def build_workbook(rows: list[dict], validation: dict) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Import_ClickUp"
    write_sheet_headers(ws, ETAPA1_HEADERS)

    for i, src in enumerate(rows, 1):
        vals = to_etapa1_row(src, i)
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=i + 1, column=col, value=val)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [55, 85, 34, 32, 10, 10, 12, 12, 26, 8, 14, 30, 6, 16, 10, 8, 16, 12, 10, 14, 10, 24, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ETAPA1_HEADERS))}{len(rows) + 1}"
    ws.freeze_panes = "A2"

    # Mapeo columnas
    ws_map = wb.create_sheet("Mapeo_Columnas_ClickUp")
    write_sheet_headers(ws_map, ["Columna Excel", "Mapear en ClickUp (ES)", "Tipo", "Notas"])
    mapeo = [
        ("Task Name", "Nombre de la tarea", "Nativo", ""),
        ("Task Description", "Descripcion", "Nativo", ""),
        ("Folder Name", "Carpeta (Folder)", "Nativo", f"Crear folder '{ETAPA1_FOLDER}' antes de importar"),
        ("List Name", "Lista (List)", "Nativo", "Crear las 11 listas del folder Etapa 1"),
        ("Status", "Estado", "Nativo", "Backlog"),
        ("Priority", "Prioridad", "Nativo", ""),
        ("Start Date", "Fecha inicio", "Nativo", "Formato Year-Month-Day"),
        ("Due Date", "Fecha limite", "Nativo", "Formato Year-Month-Day"),
        ("Tags", "Etiquetas", "Nativo", ""),
        ("Time Estimate", "Duracion estimada", "Nativo", "Horas"),
        ("Task ID", "Campo personalizado Task ID", "Texto", ""),
        ("Recurso", "Campo personalizado Recurso", "Texto", "NO usar Persona asignada"),
        ("Release", "Campo personalizado Release", "Desplegable", ""),
        ("Depends On", "Campo personalizado Depends On", "Texto", "IDs separados por coma"),
        ("Complejidad", "Campo personalizado Complejidad", "Desplegable", ""),
        ("Riesgo", "Campo personalizado Riesgo", "Desplegable", ""),
        ("Stream", "Campo personalizado Stream", "Desplegable", ""),
        ("Ambiente", "Campo personalizado Ambiente", "Desplegable", ""),
        ("Sprint", "Campo personalizado Sprint", "Desplegable", ""),
        ("Mes Calendario", "Campo personalizado Mes Calendario", "Desplegable", ""),
        ("Carga Mes Porcentaje", "Campo personalizado Carga Mes Porcentaje", "Numero", ""),
        ("Tipo Metodologia", "Campo personalizado Tipo Metodologia", "Desplegable", ""),
        ("Secuencia Import", "Opcional / referencia", "Numero", "Orden topologico sugerido"),
    ]
    for r, row in enumerate(mapeo, 2):
        for c, val in enumerate(row, 1):
            ws_map.cell(row=r, column=c, value=val).border = BORDER

    # Estructura ClickUp Etapa 1
    ws_struct = wb.create_sheet("Estructura_Etapa1")
    write_sheet_headers(ws_struct, ["Space", "Folder", "List", "Periodo", "Tareas est."])
    list_counts = {}
    for r in rows:
        list_counts[r["List Name"]] = list_counts.get(r["List Name"], 0) + 1
    periods = {
        "MES1_JUNIO_Analisis_Arquitectura": "Jun 2026",
        "MES2_JULIO_Core_Seguridad": "Jul 2026",
        "MES3_AGOSTO_UX_Sensores": "Ago 2026",
        "MES4_SEPTIEMBRE_Integracion_QA": "Sep 2026",
        "MES5_OCTUBRE_Optimizacion": "Oct 2026",
        "MES6_NOVIEMBRE_GoLive": "Nov 2026",
        "PMO_Seguimiento": "Transversal",
        "QA_Transversal": "Transversal",
        "DevOps_Infra": "Transversal",
        "Gestion_Riesgos": "Transversal",
        "Capacidad_Recursos": "Indicadores capacidad",
    }
    for lst in ETAPA1_LISTS:
        ws_struct.append([SPACE_NAME, ETAPA1_FOLDER, lst, periods.get(lst, ""), list_counts.get(lst, 0)])

    # Secuencia / dependencias
    ws_seq = wb.create_sheet("Secuencia_Tareas")
    write_sheet_headers(
        ws_seq,
        ["Secuencia", "Task ID", "Task Name", "List Name", "Recurso", "Start Date", "Due Date", "Depends On"],
    )
    for i, src in enumerate(rows, 2):
        ws_seq.append([
            i - 1,
            src["Task ID"],
            src["Task Name"],
            src["List Name"],
            src.get("Recurso"),
            src["Start Date"],
            src["Due Date"],
            (src.get("Depends On") or "").strip(),
        ])

    # Copiar resumen mensual si existe
    if SOURCE.exists():
        wb_src = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
        if "Resumen_Mensual_Detallado" in wb_src.sheetnames:
            ws_src = wb_src["Resumen_Mensual_Detallado"]
            ws_res = wb.create_sheet("Resumen_Mensual_Detallado")
            for row in ws_src.iter_rows(values_only=True):
                ws_res.append(list(row))
        wb_src.close()

    # Guia importacion
    ws_guide = wb.create_sheet("Guia_Importacion_Etapa1")
    write_sheet_headers(ws_guide, ["Paso", "Accion", "Detalle"])
    stats = validation["stats"]
    guide = [
        ("0", "Archivo", str(OUTPUT_XLSX.name)),
        ("1", f"Crear Space '{SPACE_NAME}'", "Si no existe en ClickUp"),
        ("2", f"Crear Folder '{ETAPA1_FOLDER}'", "Dentro del Space"),
        ("3", "Crear 11 listas", ", ".join(ETAPA1_LISTS)),
        ("4", "Custom Fields", "Task ID, Recurso, Depends On, Release, Complejidad, Riesgo, Stream, Ambiente, Sprint, Mes Calendario, Carga Mes Porcentaje, Tipo Metodologia"),
        ("5", "Importar Excel", "Hoja Import_ClickUp (1ra pestana)"),
        ("6", "Formato fecha", "Year-Month-Day (YYYY-MM-DD) — obligatorio"),
        ("7", "Mapeo", "Ver hoja Mapeo_Columnas_ClickUp — NO mapear Recurso a Persona asignada"),
        ("8", "Depends On", "Campo texto con IDs (ej: BE1-001,ARQ-002). ClickUp no crea links automaticos"),
        ("9", "Secuencia", f"{stats['tasks']} filas en orden topologico (predecesores primero)"),
        ("10", "Recursos", "10 recursos: BE1, BE2, BE3, FE1, FE2, ARQ, SYS, QA, IA, PAF"),
        ("11", "Tareas", f"{stats['tech_tasks']} tecnicas + {stats.get('floor_tasks', 0)} refuerzo FLOOR + {stats['cap_tasks']} indicadores CAP"),
    ]
    for r, row in enumerate(guide, 2):
        for c, val in enumerate(row, 1):
            ws_guide.cell(row=r, column=c, value=val).border = BORDER

    wb.active = 0
    return wb


def main():
    print("Regenerando cronograma base (v19)...")
    regenerate_source()
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    print("Cargando Import_ClickUp...")
    rows = load_source_rows()
    validation = validate_rows(rows)
    wb = build_workbook(rows, validation)
    wb.save(OUTPUT_XLSX)

    with open(OUTPUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(ETAPA1_HEADERS)
        for i, src in enumerate(rows, 1):
            writer.writerow(to_etapa1_row(src, i))

    print(f"\nExcel Etapa 1: {OUTPUT_XLSX}")
    print(f"CSV Etapa 1:   {OUTPUT_CSV}")
    print(f"Tareas import: {validation['stats']['tasks']}")
    print(f"  Tecnicas:    {validation['stats']['tech_tasks']}")
    print(f"  FLOOR:       {validation['stats'].get('floor_tasks', 0)}")
    print(f"  CAP:         {validation['stats']['cap_tasks']}")
    print(f"Validacion:    {'OK' if validation['ok'] else 'REVISAR'}")
    if validation["errors"]:
        for e in validation["errors"]:
            print(f"  ERROR: {e}")
    if validation["warnings"]:
        for w in validation["warnings"]:
            print(f"  AVISO: {w}")
    print(f"\nImportar hoja 'Import_ClickUp' en Space '{SPACE_NAME}' > Folder '{ETAPA1_FOLDER}'")


if __name__ == "__main__":
    main()
