#!/usr/bin/env python3
"""
Genera:
  1. Excel detallado de carga mensual por recurso (horas y %)
  2. DOCX gerencial de carga mensual
  3. DOCX de las 3 convocatorias laborales (Backend, FE1, PAF)
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SCRIPTS = ROOT / "scripts"

sys.path.insert(0, str(SCRIPTS))

# Importa cronograma ya calendarizado (tarda ~25s la primera vez)
import generate_clickup_import as gci  # noqa: E402

from generate_sow_docx import build_document  # noqa: E402

TARGET = gci.TARGET_LOAD_PCT
MONTHS = list(range(6, 12))
MONTH_LABELS = [gci.MONTH_LABELS[m] for m in MONTHS]

TITLE_COLOR = RGBColor(14, 61, 87)
ACCENT_COLOR = RGBColor(12, 104, 102)
OK_FILL = "C8E6C9"
WARN_FILL = "FFF9C4"
OVER_FILL = "FFCDD2"
HEADER_FILL = "1F4E79"
HEADER_FILL_ORANGE = "E65100"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def pct_status(pct: float, active: bool) -> str:
    if not active:
        return "N/A"
    if pct > 100:
        return "SOBRECARGA"
    if pct >= TARGET:
        return "ÓPTIMO"
    if pct >= 70:
        return "MODERADO"
    return "BAJO"


def pct_fill_excel(pct: float, active: bool) -> str | None:
    if not active:
        return "EEEEEE"
    if pct > 100:
        return OVER_FILL
    if pct >= TARGET:
        return OK_FILL
    if pct < 70:
        return WARN_FILL
    return None


def build_resource_rows():
    """Filas consolidadas y detalle tareas por recurso/mes."""
    rows = []
    task_detail: dict[str, list[dict]] = {c: [] for c in gci.RESOURCES}

    for task in gci.TASKS:
        tid, name, hours = task[2], task[3], task[15]
        assignees = gci.parse_assignees(task[5])
        alloc = gci.TASK_ALLOCS.get(tid, {})
        for code in assignees:
            if code not in gci.RESOURCES:
                continue
            code_alloc = alloc.get(code, {})
            for m in MONTHS:
                h = code_alloc.get(m, 0.0)
                if h < 0.05:
                    continue
                cap = gci.MONTH_CAPACITY[m]
                active = m in gci.active_months_for_resource(code)
                task_detail[code].append({
                    "mes": m,
                    "mes_label": gci.MONTH_LABELS[m],
                    "tid": tid,
                    "nombre": name,
                    "horas": round(h, 1),
                    "pct_mes": round(h / cap * 100, 1) if active else 0,
                })

    for code in gci.RESOURCES:
        display = gci.RESOURCE_DISPLAY[code]
        group = gci.RESOURCE_GROUP.get(code, "")
        desc = gci.RESOURCES[code]
        active_months = gci.active_months_for_resource(code)
        month_data = []
        total_h = 0.0
        active_pcts = []

        for m in MONTHS:
            h = round(gci.MONTHLY_LOAD[code][m], 1)
            cap = gci.MONTH_CAPACITY[m]
            active = m in active_months
            pct = round(h / cap * 100, 1) if active and cap else 0.0
            if active:
                total_h += h
                active_pcts.append(pct)
            month_data.append({
                "mes": m,
                "label": gci.MONTH_LABELS[m],
                "horas": h,
                "capacidad": cap,
                "pct": pct,
                "active": active,
                "estado": pct_status(pct, active),
            })

        avg_active = round(sum(active_pcts) / len(active_pcts), 1) if active_pcts else 0.0
        avg_6m = round(sum(d["pct"] for d in month_data if d["active"]) / 6, 1) if active_months else 0.0
        max_cap = gci.max_hours_for_resource(code)

        rows.append({
            "code": code,
            "display": display,
            "group": group,
            "desc": desc,
            "months": month_data,
            "total_horas": round(total_h, 1),
            "max_horas": max_cap,
            "promedio_meses_activos": avg_active,
            "promedio_6m_calendario": avg_6m,
            "tareas": sorted(task_detail[code], key=lambda x: (x["mes"], x["tid"])),
        })
    return rows


def style_header_row(ws, row: int, cols: int, fill=HEADER_FILL) -> None:
    hf = PatternFill(start_color=fill, end_color=fill, fill_type="solid")
    bf = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = hf
        cell.font = bf
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


thin = Side(style="thin")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_excel(rows, out_path: Path) -> None:
    wb = openpyxl.Workbook()

    # ── Hoja 1: Resumen consolidado ──
    ws = wb.active
    ws.title = "Resumen_Consolidado"
    headers = [
        "Código", "Recurso", "Grupo", "Total Horas", "Capacidad Proyecto (h)",
        "Prom. Meses Activos %", "Estado Global",
    ] + [f"{lb} (h)" for lb in MONTH_LABELS] + [f"{lb} (%)" for lb in MONTH_LABELS]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    t = ws.cell(row=1, column=1, value="REPORTE MENSUAL DE RECURSOS — AURIXA v36 (Jun–Nov 2026)")
    t.font = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    t.fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    t.alignment = Alignment(horizontal="center")

    for c, h in enumerate(headers, 1):
        ws.cell(row=2, column=c, value=h)
    style_header_row(ws, 2, len(headers), HEADER_FILL_ORANGE)

    cap_row = ["", "CAPACIDAD MÁXIMA MENSUAL (h)", "", "", "", ""] + [gci.MONTH_CAPACITY[m] for m in MONTHS]
    cap_row += [""] * 6
    for c, v in enumerate(cap_row, 1):
        cell = ws.cell(row=3, column=c, value=v)
        cell.font = Font(bold=True)
        cell.border = BORDER

    for r, row in enumerate(rows, 4):
        md = {d["label"]: d for d in row["months"]}
        estado = "ÓPTIMO" if row["promedio_meses_activos"] >= TARGET else "REVISAR"
        vals = [
            row["code"], row["display"], row["group"],
            row["total_horas"], row["max_horas"],
            f"{row['promedio_meses_activos']}%", estado,
        ]
        for lb in MONTH_LABELS:
            d = md[lb]
            vals.append(d["horas"] if d["active"] else "—")
        for lb in MONTH_LABELS:
            d = md[lb]
            vals.append(f"{d['pct']}%" if d["active"] else "N/A")
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if c >= 14:
                m_idx = c - 14
                d = row["months"][m_idx]
                fill = pct_fill_excel(d["pct"], d["active"])
                if fill:
                    cell.fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")

    for i, w in enumerate([8, 38, 14, 12, 18, 16, 14] + [10] * 12, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"

    # ── Hoja 2: Detalle horas + % (formato gerencial) ──
    ws2 = wb.create_sheet("Detalle_Horas_Porcentaje")
    h2 = ["Recurso", "Código", "Grupo"]
    for lb in MONTH_LABELS:
        h2.extend([f"{lb} (h)", f"{lb} (%)", f"{lb} Estado"])
    h2.extend(["Total (h)", "Prom. activos %", "Estado"])

    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(h2))
    t2 = ws2.cell(row=1, column=1, value="DETALLE MENSUAL — HORAS Y % DE CARGA LABORAL")
    t2.font = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
    t2.fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    t2.alignment = Alignment(horizontal="center")

    for c, h in enumerate(h2, 1):
        ws2.cell(row=2, column=c, value=h)
    style_header_row(ws2, 2, len(h2))

    for r, row in enumerate(rows, 3):
        vals = [row["display"], row["code"], row["group"]]
        for d in row["months"]:
            if d["active"]:
                vals.extend([d["horas"], f"{d['pct']}%", d["estado"]])
            else:
                vals.extend(["—", "N/A", "Sin asignación"])
        vals.extend([
            row["total_horas"],
            f"{row['promedio_meses_activos']}%",
            "ÓPTIMO" if row["promedio_meses_activos"] >= TARGET else "SUBUTILIZADO",
        ])
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws2.freeze_panes = "A3"

    # ── Hoja por mes ──
    for m in MONTHS:
        lb = gci.MONTH_LABELS[m]
        ws_m = wb.create_sheet(f"Mes_{m - 5}_{lb[:3]}")
        ws_m.cell(row=1, column=1, value=f"MES {m - 5} — {lb.upper()} 2026 — Carga por recurso")
        ws_m.cell(row=1, column=1).font = Font(bold=True, size=12, color="FFFFFF")
        ws_m.cell(row=1, column=1).fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
        ws_m.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)

        mh = ["Código", "Recurso", "Horas asignadas", "Capacidad mes (h)", "% Carga", "Estado", "Horas disponibles", "Grupo"]
        for c, h in enumerate(mh, 1):
            ws_m.cell(row=2, column=c, value=h)
        style_header_row(ws_m, 2, len(mh))

        cap = gci.MONTH_CAPACITY[m]
        ws_m.cell(row=3, column=1, value="")
        ws_m.cell(row=3, column=2, value="Capacidad bruta del mes")
        ws_m.cell(row=3, column=4, value=cap)
        ws_m.cell(row=3, column=5, value="100%")
        for c in range(1, 9):
            ws_m.cell(row=3, column=c).font = Font(bold=True, italic=True)

        for r, row in enumerate(rows, 4):
            d = row["months"][m - 6]
            disp_h = d["horas"] if d["active"] else 0
            disp_pct = d["pct"] if d["active"] else 0
            libres = round(cap - disp_h, 1) if d["active"] else cap
            ws_m.cell(row=r, column=1, value=row["code"])
            ws_m.cell(row=r, column=2, value=row["display"])
            ws_m.cell(row=r, column=3, value=d["horas"] if d["active"] else "—")
            ws_m.cell(row=r, column=4, value=cap if d["active"] else "—")
            ws_m.cell(row=r, column=5, value=f"{disp_pct}%" if d["active"] else "N/A")
            ws_m.cell(row=r, column=6, value=d["estado"])
            ws_m.cell(row=r, column=7, value=libres if d["active"] else "—")
            ws_m.cell(row=r, column=8, value=row["group"])
            fill = pct_fill_excel(disp_pct, d["active"])
            if fill:
                for c in range(1, 9):
                    ws_m.cell(row=r, column=c).fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")
            for c in range(1, 9):
                ws_m.cell(row=r, column=c).border = BORDER

    # ── Hoja detalle tareas por recurso ──
    ws_t = wb.create_sheet("Detalle_Tareas_por_Recurso")
    th = ["Código Recurso", "Recurso", "Mes", "Task ID", "Tarea", "Horas asignadas", "% del mes"]
    for c, h in enumerate(th, 1):
        ws_t.cell(row=1, column=c, value=h)
    style_header_row(ws_t, 1, len(th))

    tr = 2
    for row in rows:
        for t in row["tareas"]:
            ws_t.cell(row=tr, column=1, value=row["code"])
            ws_t.cell(row=tr, column=2, value=row["display"])
            ws_t.cell(row=tr, column=3, value=t["mes_label"])
            ws_t.cell(row=tr, column=4, value=t["tid"])
            ws_t.cell(row=tr, column=5, value=t["nombre"])
            ws_t.cell(row=tr, column=6, value=t["horas"])
            ws_t.cell(row=tr, column=7, value=f"{t['pct_mes']}%")
            for c in range(1, 8):
                ws_t.cell(row=tr, column=c).border = BORDER
            tr += 1
    ws_t.freeze_panes = "A2"

    wb.save(out_path)


def add_docx_table(doc: Document, headers: list[str], data_rows: list[list]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "DCE6F1")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)

    for row_vals in data_rows:
        cells = table.add_row().cells
        for i, val in enumerate(row_vals):
            cells[i].text = str(val)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()


def write_docx_report(rows, out_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Reporte Mensual de Recursos — Carga Laboral Detallada")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = TITLE_COLOR

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("Proyecto AURIXA · Plataforma Minera Enterprise · Junio–Noviembre 2026")
    r2.font.size = Pt(12)
    r2.font.color.rgb = ACCENT_COLOR

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        f"Modelo: {gci.HOURS_PER_WORKDAY} h/día laboral · Feriados Perú 2026 excluidos · "
        f"Meta mínima: {TARGET}% en meses activos · Tope máximo: 100%"
    )
    mr.font.size = Pt(10)

    doc.add_paragraph()
    doc.add_heading("1. Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "Este reporte consolida la carga laboral planificada de los 10 recursos del proyecto. "
        "Los porcentajes se calculan sobre la capacidad mensual real (días netos × 8,5 h). "
        "Recursos que inician en mes 2 (FE2, SYS, QA, IA) muestran N/A en junio por diseño del plan de contratación."
    )

    doc.add_heading("2. Matriz consolidada — Horas y % por mes", level=1)
    headers = ["Recurso", "Código"] + MONTH_LABELS + ["Prom. activos %", "Estado"]
    matrix = []
    for row in rows:
        pcts = [f"{d['pct']}%" if d["active"] else "N/A" for d in row["months"]]
        hrs = [str(d["horas"]) if d["active"] else "—" for d in row["months"]]
        combined = [f"{h} / {p}" for h, p in zip(hrs, pcts)]
        estado = "ÓPTIMO" if row["promedio_meses_activos"] >= TARGET else "REVISAR"
        matrix.append([row["display"], row["code"]] + combined + [f"{row['promedio_meses_activos']}%", estado])
    add_docx_table(doc, headers, matrix)

    doc.add_heading("3. Capacidad mensual del proyecto", level=1)
    cap_rows = [[gci.MONTH_LABELS[m], f"{gci.MONTH_CAPACITY[m]} h", f"{gci.month_net_workdays(m)} días netos"] for m in MONTHS]
    add_docx_table(doc, ["Mes", "Capacidad máxima", "Días laborables netos"], cap_rows)

    doc.add_heading("4. Detalle por recurso", level=1)
    for row in rows:
        doc.add_heading(f"{row['code']} — {row['display']}", level=2)
        p = doc.add_paragraph()
        p.add_run(row["desc"]).font.size = Pt(10)

        detail_headers = ["Mes", "Horas", "Capacidad (h)", "% Carga", "Estado", "Horas libres"]
        detail_data = []
        for d in row["months"]:
            libres = round(d["capacidad"] - d["horas"], 1) if d["active"] else "—"
            detail_data.append([
                d["label"],
                d["horas"] if d["active"] else "—",
                d["capacidad"] if d["active"] else "—",
                f"{d['pct']}%" if d["active"] else "N/A",
                d["estado"],
                libres,
            ])
        add_docx_table(doc, detail_headers, detail_data)

        top_tasks = row["tareas"][:12]
        if top_tasks:
            doc.add_paragraph("Principales tareas asignadas (muestra):")
            th = ["Mes", "ID", "Tarea", "Horas"]
            td = [[t["mes_label"], t["tid"], t["nombre"][:80], t["horas"]] for t in top_tasks]
            add_docx_table(doc, th, td)

    doc.add_heading("5. Notas metodológicas", level=1)
    for note in [
        "Fuente: scripts/generate_clickup_import.py — cronograma v36 con alcances A–G.",
        f"Promedio 'meses activos' excluye meses sin contratación (ej. FE2/QA/IA/SYS en junio).",
        "Estado ÓPTIMO: ≥ 90% en meses con asignación; tope 100% sin sobrecarga.",
        "Generado automáticamente para gerencia y seguimiento PMO.",
    ]:
        doc.add_paragraph(note, style="List Bullet")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("AURIXA 2026 — Reporte mensual de recursos v36")
    fr.font.size = Pt(8)
    fr.font.color.rgb = ACCENT_COLOR

    doc.save(out_path)


def split_convocatoria_fe1_paf() -> tuple[Path, Path]:
    src = DOCS / "Convocatoria_Laboral_02_Frontend_PAF_AURIXA.md"
    text = src.read_text(encoding="utf-8")
    fe1_path = DOCS / "Convocatoria_Laboral_02_Frontend_FE1_AURIXA.md"
    paf_path = DOCS / "Convocatoria_Laboral_03_PAF_AURIXA.md"

    fe1 = """# Convocatoria Laboral — Frontend Senior FE1
## ReportStudio e Informes Técnicos Mineros Enterprise

> **Proyecto:** Plataforma AURIXA · **Inicio:** Mes 1 (único frontend activo en junio) · **Duración:** 6 meses

---

""" + text.split("# PARTE A — Frontend Senior FE1")[1].split("# PARTE B — Profesional de Apoyo Funcional (PAF)")[0].strip()

    paf = """# Convocatoria Laboral — Profesional de Apoyo Funcional (PAF)
## Analista Funcional · PMO · Documentación · Capacitación

> **Proyecto:** Plataforma AURIXA · **Duración:** 6 meses (mes 1 al 6)

---

""" + text.split("# PARTE B — Profesional de Apoyo Funcional (PAF)")[1].split("## Versión corta LinkedIn")[0].strip()

    fe1_path.write_text(fe1, encoding="utf-8")
    paf_path.write_text(paf, encoding="utf-8")
    return fe1_path, paf_path


def main() -> int:
    print("Calculando carga mensual desde cronograma...")
    rows = build_resource_rows()

    xlsx_out = DOCS / "Reporte_Mensual_Recursos_v36_Detallado.xlsx"
    docx_out = DOCS / "Reporte_Mensual_Recursos_v36_Detallado.docx"
    write_excel(rows, xlsx_out)
    print(f"Excel: {xlsx_out}")
    write_docx_report(rows, docx_out)
    print(f"DOCX reporte: {docx_out}")

    convocatorias = [
        (DOCS / "Convocatoria_Laboral_01_Backend_Base_Datos_AURIXA.md",
         DOCS / "Convocatoria_Laboral_01_Backend_Base_Datos_AURIXA.docx"),
    ]
    fe1_md, paf_md = split_convocatoria_fe1_paf()
    convocatorias.extend([
        (fe1_md, DOCS / "Convocatoria_Laboral_02_Frontend_FE1_AURIXA.docx"),
        (paf_md, DOCS / "Convocatoria_Laboral_03_PAF_AURIXA.docx"),
    ])

    for md, docx in convocatorias:
        build_document(md, docx)
        print(f"Convocatoria DOCX: {docx}")

    print("\n--- Validación carga (meses activos) ---")
    for row in rows:
        flag = "OK" if row["promedio_meses_activos"] >= TARGET else "BAJO"
        print(f"  {row['code']:4s} prom activos {row['promedio_meses_activos']:5.1f}% [{flag}]")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
