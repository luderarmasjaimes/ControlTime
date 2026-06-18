"""
Genera CONVOCATORIA_LINKEDIN_BE_DATABASE.docx
Formato tipo LinkedIn: una columna, tipografía limpia, emojis, listo para copiar/pegar.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Colores ──────────────────────────────────────────────────────────────────
C_NEGRO       = RGBColor(0x1B, 0x1B, 0x1B)
C_GRIS_TENUE  = RGBColor(0x66, 0x66, 0x66)
C_AZUL_LI     = RGBColor(0x00, 0x66, 0xC2)   # Azul LinkedIn
C_FONDO_TAG   = RGBColor(0xEE, 0xF3, 0xF8)
C_BORDE       = RGBColor(0xCC, 0xCC, 0xCC)
C_WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
C_ROJO        = RGBColor(0xC0, 0x39, 0x2B)
C_VERDE       = RGBColor(0x1A, 0x7A, 0x3C)


def hex_str(c: RGBColor) -> str:
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def set_shading(para, color: RGBColor):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_str(color))
    pPr.append(shd)


def add_bottom_border(para, color: RGBColor = C_BORDE, sz=6):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), str(sz))
    b.set(qn('w:space'), '6')
    b.set(qn('w:color'), hex_str(color))
    pBdr.append(b)
    pPr.append(pBdr)


def parse(text: str):
    """Devuelve lista de (segmento, is_bold) desde **negrita**."""
    parts = text.split('**')
    return [(p, i % 2 == 1) for i, p in enumerate(parts) if p]


# ─────────────────────────────────────────────────────────────────────────────

def build():
    doc = Document()

    # Márgenes generosos — simula el ancho del feed LinkedIn
    for sec in doc.sections:
        sec.top_margin    = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin   = Cm(3.5)
        sec.right_margin  = Cm(3.5)
        sec.page_width    = Cm(21)   # A4

    style = doc.styles['Normal']
    style.font.name = 'Segoe UI'
    style.font.size = Pt(11)
    style.font.color.rgb = C_NEGRO

    # ── Helper: párrafo de texto con espaciado LinkedIn ──────────────────────
    def p_li(before=4, after=4):
        """Párrafo en blanco con espaciado estándar."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after  = Pt(after)
        return p

    def run(para, text, bold=False, size=None, color=C_NEGRO, italic=False):
        r = para.add_run(text)
        r.bold   = bold
        r.italic = italic
        r.font.name = 'Segoe UI'
        r.font.size = Pt(size or 11)
        r.font.color.rgb = color
        return r

    def run_mix(para, segments, size=11):
        for txt, is_bold in segments:
            run(para, txt, bold=is_bold, size=size)

    def section_title(icon, title, color=C_AZUL_LI, size=12):
        """Encabezado de sección estilo LinkedIn: emoji + texto en azul, borde inferior."""
        p = p_li(before=14, after=4)
        run(p, f"{icon} ", bold=False, size=size, color=color)
        run(p, title, bold=True, size=size, color=color)
        add_bottom_border(p, color=C_AZUL_LI, sz=6)
        return p

    def bullet(icon, text, icon_color=C_NEGRO, size=11):
        """Ítem de lista."""
        p = doc.add_paragraph()
        p.paragraph_format.left_indent       = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after  = Pt(3)
        r = p.add_run(f"{icon}  ")
        r.font.name = 'Segoe UI'
        r.font.size = Pt(size)
        r.font.color.rgb = icon_color
        run_mix(p, parse(text), size=size)
        return p

    def spacer():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        run(p, "", size=4)
        return p

    # ══════════════════════════════════════════════════════════════════════════
    # CABECERA — simulando encabezado de post LinkedIn
    # ══════════════════════════════════════════════════════════════════════════
    p_rol = p_li(before=0, after=2)
    p_rol.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run(p_rol, "Se busca:", bold=False, size=11, color=C_GRIS_TENUE)

    p_titulo = p_li(before=0, after=4)
    run(p_titulo, "Arquitecto de Datos & Backend Senior", bold=True, size=20, color=C_NEGRO)

    p_sub = p_li(before=0, after=2)
    run(p_sub, "Plataforma Industrial de Alto Impacto", bold=False, size=13, color=C_GRIS_TENUE)

    # Badges de modalidad
    p_badges = p_li(before=4, after=8)
    for badge, sep in [
        ("📍 Remoto / Híbrido", "   "),
        ("⏱ Full Time", "   "),
        ("📄 6 meses renovable", "   "),
        ("🚀 Incorporación inmediata", ""),
    ]:
        run(p_badges, badge, bold=False, size=10, color=C_AZUL_LI)
        if sep:
            run(p_badges, sep, size=10, color=C_GRIS_TENUE)

    add_bottom_border(p_badges, color=C_BORDE, sz=4)

    # ── Intro ─────────────────────────────────────────────────────────────────
    spacer()
    p_i1 = p_li(before=6, after=4)
    p_i1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run_mix(p_i1, parse(
        "Organización del sector industrial convoca a un profesional **Senior** para liderar "
        "la capa de datos, persistencia y servicios backend de una plataforma web enterprise "
        "de monitoreo operativo, reportabilidad técnica y gestión documental avanzada."
    ), size=11)

    p_i2 = p_li(before=0, after=10)
    p_i2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(p_i2,
        "¿Te apasionan la arquitectura de datos, la continuidad operativa y la trazabilidad "
        "en entornos con conectividad limitada? Esta es tu oportunidad de liderar un proyecto "
        "de transformación digital de alto impacto en la región.",
        italic=True, size=11, color=C_GRIS_TENUE)

    # ══════════════════════════════════════════════════════════════════════════
    # TU MISIÓN
    # ══════════════════════════════════════════════════════════════════════════
    section_title("🎯", "TU MISIÓN")
    p_m = p_li(before=4, after=10)
    p_m.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(p_m,
        "Garantizar que la arquitectura de datos y los servicios de núcleo soporten de forma "
        "confiable: informes técnicos estructurados, telemetría operativa de alta frecuencia, "
        "auditoría completa de acciones, operación offline con sincronización controlada e "
        "integración segura con los sistemas existentes de planta.", size=11)

    # ══════════════════════════════════════════════════════════════════════════
    # LO QUE HARÁS
    # ══════════════════════════════════════════════════════════════════════════
    section_title("🔧", "LO QUE HARÁS EN TU DÍA A DÍA")
    tareas = [
        "Diseñar e implementar modelos de datos transaccionales, documentales y de **series temporales** con soporte multi-empresa, auditoría y políticas de retención.",
        "Desarrollar y mantener servicios backend de **alto rendimiento**: APIs REST, canales de comunicación en tiempo real y autoservicio documental.",
        "Construir pipelines **ETL/ELT** e integraciones con fuentes legacy e IoT/industrial, garantizando idempotencia y monitoreo de rezagos.",
        "Diseñar estrategias de persistencia local, colas de cambios pendientes y mecanismos de reconciliación al restablecer la conectividad.",
        "Administrar entornos **Linux y contenedores**: backups probados, planes de recuperación ante desastres y migraciones de esquema versionadas.",
        "Aplicar controles de seguridad: **Row-Level Security**, cifrado en tránsito y en reposo, bitácoras inmutables y cumplimiento normativo.",
        "Definir estándares de arquitectura de datos; guiar y asesorar a equipos de desarrollo, QA y PMO.",
        "Colaborar en contratos de integración con frontend, revisiones de carga y documentación operativa.",
        "Participar en metodología ágil con **Git** y herramienta de gestión de proyectos (ClickUp o equivalente).",
    ]
    for t in tareas:
        bullet("◆", t, icon_color=C_AZUL_LI)
    spacer()

    # ══════════════════════════════════════════════════════════════════════════
    # REQUISITOS EXCLUYENTES
    # ══════════════════════════════════════════════════════════════════════════
    section_title("✅", "REQUISITOS EXCLUYENTES", color=C_ROJO)
    p_nota_excl = p_li(before=2, after=4)
    set_shading(p_nota_excl, RGBColor(0xFF, 0xF0, 0xEF))
    run(p_nota_excl, "  ⚠  La ausencia de cualquiera de estos puntos descalifica la postulación.",
        size=10, color=C_ROJO, italic=True)

    excl = [
        "**8+ años** de experiencia en backend y arquitectura de datos; mínimo **5 años** con motor relacional en producción.",
        "**Dominio avanzado de PostgreSQL:** modelado relacional y documental (JSON/JSONB), partición, índices avanzados, replicación, backup/restore y tuning.",
        "Servicios backend de **alto rendimiento**, APIs REST y comunicación en tiempo real (WebSockets, SSE o equivalente).",
        "**Trazabilidad forense, auditoría de acciones y gobierno de datos** alineado a estándares de Industria 4.0/5.0.",
        "Pipelines **ETL/ELT**, integración con sistemas externos (legacy e IoT) y sincronización incremental controlada.",
        "Seguridad en datos: **RBAC, cifrado, OWASP** aplicado a APIs y accesos, hardening de bases de datos.",
        "Operación fluida en **Linux, Docker y entornos productivos con CI/CD.**",
        "Experiencia en **proyectos industriales:** manufactura, energía, utilities, minería o plantas automatizadas.",
        "**Git y gestión ágil** de proyectos en equipos multidisciplinarios.",
        "**Español fluido; inglés técnico de lectura** (documentación, RFCs, whitepapers).",
    ]
    for e in excl:
        bullet("▸", e, icon_color=C_ROJO)
    spacer()

    # ══════════════════════════════════════════════════════════════════════════
    # REQUISITOS IMPORTANTES
    # ══════════════════════════════════════════════════════════════════════════
    section_title("⭐", "REQUISITOS IMPORTANTES")
    p_nota_imp = p_li(before=2, after=4)
    set_shading(p_nota_imp, C_FONDO_TAG)
    run(p_nota_imp, "  Alta valoración — marcan la diferencia en la evaluación.", size=10, color=C_GRIS_TENUE, italic=True)

    imp = [
        "Diseño de arquitecturas de datos end-to-end con foco en **escalabilidad, resiliencia y seguridad.**",
        "Patrones de arquitectura: **Data Lake, Lakehouse, Data Warehouse,** procesamiento batch y streaming.",
        "Bases de datos **relacionales** (PostgreSQL, SQL Server, MySQL) **y NoSQL** (MongoDB, Redis), con criterios claros de selección.",
        "Técnicas de **modelado de datos** (ERD, UML, Draw.io) y planificación de capacidad.",
        "**SQL avanzado** y Python para análisis, scripting y decisiones de arquitectura.",
        "Capacidad para **asesorar a perfiles junior y mid** (Data Engineers, BI, DevOps).",
    ]
    for i in imp:
        bullet("▸", i, icon_color=C_AZUL_LI)
    spacer()

    # ══════════════════════════════════════════════════════════════════════════
    # CONOCIMIENTOS DESEABLES
    # ══════════════════════════════════════════════════════════════════════════
    section_title("💡", "CONOCIMIENTOS DESEABLES", color=C_GRIS_TENUE)
    p_nota_des = p_li(before=2, after=4)
    set_shading(p_nota_des, RGBColor(0xF8, 0xF8, 0xF8))
    run(p_nota_des, "  No excluyentes — altamente valorados.", size=10, color=C_GRIS_TENUE, italic=True)

    des = [
        "**Series temporales:** TimescaleDB, InfluxDB y telemetría de alta frecuencia.",
        "**Mensajería/streaming:** Kafka, Redpanda, NATS o equivalentes.",
        "**Multi-tenancy** con Row-Level Security y políticas de retención por regulación.",
        "Modelado de eventos con **validación biométrica** integrada a flujos de acceso.",
        "Operaciones con **conectividad intermitente** y reconciliación de datos offline.",
        "Certificaciones en **DBA, seguridad de la información** o normas IEC 62443.",
        "Marcos **ISO** aplicables a calidad de software y trazabilidad documental.",
        "Infraestructura como código **(IaC)** y orquestación de contenedores.",
    ]
    for d in des:
        bullet("▸", d, icon_color=C_GRIS_TENUE)
    spacer()

    # ══════════════════════════════════════════════════════════════════════════
    # QUÉ TE OFRECEMOS
    # ══════════════════════════════════════════════════════════════════════════
    section_title("🏆", "QUÉ TE OFRECEMOS", color=C_VERDE)
    oferta = [
        ("✔", "Proyecto de **alto impacto real** en transformación digital industrial — tus decisiones de arquitectura se verán en producción desde el primer sprint."),
        ("✔", "**Autonomía técnica plena:** aquí se trabaja con argumentos, no con jerarquías."),
        ("✔", "Equipo multidisciplinario consolidado con metodología ágil operativa y procesos definidos."),
        ("✔", "**Planilla completa desde el primer día.** Contrato 6 meses, renovable según desempeño."),
        ("✔", "Compensación **acorde a perfil Senior** — a coordinar en el proceso de selección."),
    ]
    for icon, txt in oferta:
        bullet(icon, txt, icon_color=C_VERDE)
    spacer()

    # ══════════════════════════════════════════════════════════════════════════
    # CÓMO POSTULAR
    # ══════════════════════════════════════════════════════════════════════════
    section_title("📩", "¿CÓMO POSTULAR?")

    p_env = p_li(before=4, after=2)
    run(p_env, "Envía los siguientes elementos al correo ", size=11)
    run(p_env, "[CORREO DE CONTACTO]", bold=True, size=11, color=C_ROJO)
    run(p_env, " con el asunto:", size=11)

    p_asunto = p_li(before=2, after=8)
    p_asunto.paragraph_format.left_indent = Cm(0.8)
    set_shading(p_asunto, C_FONDO_TAG)
    run(p_asunto, '  "Senior Backend / Arquitecto de Datos — [TU NOMBRE COMPLETO]"',
        bold=True, italic=True, size=11, color=C_AZUL_LI)

    pasos = [
        ("1.", "CV actualizado en PDF", " (máx. 3 páginas)."),
        ("2.", "Portafolio o repositorios de código", " (GitHub, GitLab — opcional pero valorado)."),
        ("3.", "Respuesta breve en el correo", " (máx. 150 palabras) a la pregunta de filtro:"),
    ]
    for num, bold_txt, rest in pasos:
        pb = doc.add_paragraph()
        pb.paragraph_format.left_indent = Cm(0.8)
        pb.paragraph_format.space_before = Pt(2)
        pb.paragraph_format.space_after = Pt(3)
        run(pb, f"{num}  ", bold=True, size=11, color=C_AZUL_LI)
        run(pb, bold_txt, bold=True, size=11)
        run(pb, rest, size=11)

    # Pregunta filtro
    p_q = p_li(before=4, after=4)
    p_q.paragraph_format.left_indent = Cm(1.0)
    p_q.paragraph_format.right_indent = Cm(0.5)
    set_shading(p_q, C_FONDO_TAG)
    run(p_q,
        '  "¿Cuál fue el proyecto de datos de mayor complejidad técnica en el que participaste? '
        'Describe específicamente tu rol, las decisiones de arquitectura que tomaste y el impacto medible que generaste."',
        italic=True, size=10.5, color=C_NEGRO)

    p_prio = p_li(before=6, after=10)
    run(p_prio,
        "Los candidatos que respondan esta pregunta tienen prioridad. "
        "Los que no la respondan no serán considerados en la primera revisión.",
        bold=True, size=10.5, color=C_ROJO)

    # ══════════════════════════════════════════════════════════════════════════
    # PROCESO DE SELECCIÓN — tabla simple estilo LinkedIn
    # ══════════════════════════════════════════════════════════════════════════
    section_title("📋", "PROCESO DE SELECCIÓN")
    spacer()

    tbl = doc.add_table(rows=6, cols=3)
    tbl_xml = tbl._tbl
    tblPr = tbl_xml.find(qn('w:tblPr')) or OxmlElement('w:tblPr')
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '9360')
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)

    def cell_fmt(cell, text, bold=False, bg=None, align=WD_ALIGN_PARAGRAPH.LEFT, color=C_NEGRO):
        if bg:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), hex_str(bg))
            tcPr.append(shd)
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)
        r = p.add_run(text)
        r.bold = bold
        r.font.name = 'Segoe UI'
        r.font.size = Pt(9.5)
        r.font.color.rgb = color

    hdrs = ["FASE", "DURACIÓN", "RESPONSABLE"]
    for i, h in enumerate(hdrs):
        cell_fmt(tbl.rows[0].cells[i], h, bold=True,
                 bg=C_AZUL_LI, color=C_WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

    rows = [
        ("Revisión de CV + respuesta técnica obligatoria", "48 h", "Líder técnico / ARQ"),
        ("Entrevista técnica — Arquitectura de datos y PostgreSQL (45 min)", "1 sesión", "ARQ"),
        ("Challenge técnico asíncrono (máx. 4 h del candidato)", "3–5 días", "Candidato"),
        ("Entrevista final — encaje técnico y cultural (45 min)", "1 sesión", "ARQ + PM"),
        ("Oferta, negociación y firma", "48 h", "PM / RRHH"),
    ]
    alt = [RGBColor(0xF0, 0xF6, 0xFD), C_WHITE]
    for ri, (f, d, r_) in enumerate(rows):
        bg = alt[ri % 2]
        for ci, val in enumerate([(f, WD_ALIGN_PARAGRAPH.LEFT), (d, WD_ALIGN_PARAGRAPH.CENTER), (r_, WD_ALIGN_PARAGRAPH.CENTER)]):
            cell_fmt(tbl.rows[ri+1].cells[ci], val[0], bg=bg, align=val[1])

    spacer()
    spacer()

    # ══════════════════════════════════════════════════════════════════════════
    # HASHTAGS — estilo tag LinkedIn
    # ══════════════════════════════════════════════════════════════════════════
    add_bottom_border(p_li(before=4, after=4), color=C_BORDE, sz=4)

    p_ref = p_li(before=6, after=4)
    run(p_ref, "🔁 Si conoces a alguien con este perfil, comparte esta publicación. "
               "El mejor talento técnico llega por referidos.", size=10.5, color=C_GRIS_TENUE, italic=True)

    tags = [
        "#ArquitectoDeDatos", "#BackendSenior", "#PostgreSQL", "#DataEngineering",
        "#TimescaleDB", "#Docker", "#Linux", "#IndustriaMinera",
        "#TransformacionDigital", "#Industria40", "#TechJobsLATAM",
        "#RemoteWork", "#DataArchitecture", "#ETL",
    ]
    p_tags = p_li(before=6, after=6)
    p_tags.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for tag in tags:
        run(p_tags, tag + "  ", bold=False, size=9.5, color=C_AZUL_LI)

    # ── Pie de página ─────────────────────────────────────────────────────────
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f = fp.add_run("Convocatoria — Arquitecto de Datos & Backend Senior · Mayo 2026 · Para uso y publicación en LinkedIn")
    r_f.font.name = 'Segoe UI'
    r_f.font.size = Pt(8)
    r_f.font.color.rgb = C_GRIS_TENUE

    # ── Guardar ───────────────────────────────────────────────────────────────
    out = r"c:\InformeCliente\CONVOCATORIA_LINKEDIN_BE_DATABASE.docx"
    doc.save(out)
    print(f"✅  Generado: {out}")


if __name__ == "__main__":
    build()
