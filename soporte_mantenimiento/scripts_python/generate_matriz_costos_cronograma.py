import copy

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUTPUT_XLSX = r'c:\InformeCliente\docs\Matriz_Costos_Cronograma_26_Semanas.xlsx'
OUTPUT_MD = r'c:\InformeCliente\docs\Matriz_Costos_Cronograma_26_Semanas.md'

roles = [
    {'rol': 'ARQ', 'descripcion': 'Arquitecto Senior LATAM', 'monthly_rate': 1900.00},
    {'rol': 'BE1', 'descripcion': 'Arquitecto de Datos Senior', 'monthly_rate': 1900.00},
    {'rol': 'BE2', 'descripcion': 'Especialista en Seguridad e IA', 'monthly_rate': 1900.00},
    {'rol': 'BE3', 'descripcion': 'Ingeniero de Integracion', 'monthly_rate': 1900.00},
    {'rol': 'FE1', 'descripcion': 'Frontend Senior Principal', 'monthly_rate': 1800.00},
    {'rol': 'FE2', 'descripcion': 'Frontend/UX Senior', 'monthly_rate': 1800.00},
    {'rol': 'SYS', 'descripcion': 'Especialista Infraestructura VPS/Linux', 'monthly_rate': 1800.00},
    {'rol': 'QA', 'descripcion': 'Auditor de Calidad Senior', 'monthly_rate': 1800.00},
    {'rol': 'IA', 'descripcion': 'Ingeniero IA Aplicada', 'monthly_rate': 1900.00},
]

period_labels = ['Mes 1', 'Mes 2', 'Mes 3', 'Mes 4', 'Mes 5', 'Mes 6']
stage1_periods = 4
stage2_periods = 2
contingency_pct = 0.12
management_reserve_pct = 0.05
currency = 'USD'

# Costos adicionales editables
laptop_unit_cost = 2200.00
laptop_units_per_role = 1
payroll_factors = [
    ('AFP', 0.10, 'Aporte previsional referencial'),
    ('EsSalud', 0.09, 'Seguro de salud sobre remuneracion'),
    ('CTS', 0.0972, 'Compensacion por tiempo de servicios'),
    ('Gratificaciones', 0.1667, 'Dos gratificaciones anuales prorrateadas'),
    ('Vacaciones', 0.0833, 'Provision anual de vacaciones'),
    ('Utilidades', 0.05, 'Participacion de utilidades referencial'),
    ('Escolaridad', 0.02, 'Provision referencial de beneficio escolar'),
]

thin = Side(style='thin', color='C8C8C8')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='center')
header_fill = PatternFill('solid', fgColor='1F3A5F')
section_fill = PatternFill('solid', fgColor='2D6A4F')
sub_fill = PatternFill('solid', fgColor='DCEFE3')
alt_fill = PatternFill('solid', fgColor='F5F7FA')
highlight_fill = PatternFill('solid', fgColor='FFF2CC')


def money(cell):
    cell.number_format = '$#,##0.00'


def pct(cell):
    cell.number_format = '0%'


def write_headers(ws, row, headers):
    for idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=idx, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = section_fill
        cell.alignment = wrap
        cell.border = border


role_cost_rows = []
monthly_totals = [0, 0, 0, 0, 0, 0]
stage1_total = 0
stage2_total = 0
base_total = 0

for item in roles:
    period_costs = [round(item['monthly_rate'], 2) for _ in period_labels]
    cost1 = round(sum(period_costs[:stage1_periods]), 2)
    cost2 = round(sum(period_costs[stage1_periods:]), 2)
    base = round(cost1 + cost2, 2)
    with_reserve = round(base * (1 + contingency_pct + management_reserve_pct), 2)
    percentages = [1.0 for _ in period_labels]

    monthly_totals = [round(total + cost, 2) for total, cost in zip(monthly_totals, period_costs)]
    stage1_total += cost1
    stage2_total += cost2
    base_total += base
    role_cost_rows.append({
        'rol': item['rol'],
        'descripcion': item['descripcion'],
        'monthly_rate': item['monthly_rate'],
        'cost1': cost1,
        'cost2': cost2,
        'base': base,
        'with_reserve': with_reserve,
        'period_costs': period_costs,
        'percentages': percentages,
    })

stage1_total = round(stage1_total, 2)
stage2_total = round(stage2_total, 2)
base_total = round(base_total, 2)
base_contingency = round(base_total * contingency_pct, 2)
base_reserve = round(base_total * management_reserve_pct, 2)
recommended_total = round(base_total + base_contingency + base_reserve, 2)

total_laptops_cost = round(len(roles) * laptop_units_per_role * laptop_unit_cost, 2)
payroll_total = 0
payroll_rows = []
for label, factor, description in payroll_factors:
    subtotal = round(base_total * factor, 2)
    payroll_total += subtotal
    payroll_rows.append((label, factor, subtotal, description))
payroll_total = round(payroll_total, 2)

subtotal_expanded = round(base_total + total_laptops_cost + payroll_total, 2)
expanded_contingency = round(subtotal_expanded * contingency_pct, 2)
expanded_reserve = round(subtotal_expanded * management_reserve_pct, 2)
expanded_total = round(subtotal_expanded + expanded_contingency + expanded_reserve, 2)

wb = openpyxl.Workbook()

# Resumen de Costos
ws = wb.active
ws.title = 'Resumen de Costos'
for col, width in {'A': 14, 'B': 38, 'C': 16, 'D': 16, 'E': 16, 'F': 16, 'G': 16}.items():
    ws.column_dimensions[col].width = width

row = 1
write_headers(ws, row, ['Rol', 'Descripcion', 'Tarifa Mensual', 'Costo Etapa 1', 'Costo Etapa 2', 'Costo Base Total', 'Costo con Reserva'])
row += 1

for idx, data in enumerate(role_cost_rows, start=row):
    values = [data['rol'], data['descripcion'], data['monthly_rate'], data['cost1'], data['cost2'], data['base'], data['with_reserve']]
    for cidx, value in enumerate(values, 1):
        cell = ws.cell(row=idx, column=cidx, value=value)
        cell.font = Font(size=10)
        cell.alignment = wrap
        cell.border = border
        if cidx >= 3:
            money(cell)
        if idx % 2 == 0:
            cell.fill = alt_fill
    row = idx + 1

summary_labels = [
    ('Total Etapa 1', stage1_total),
    ('Total Etapa 2', stage2_total),
    ('Total Base', base_total),
    ('Contingencia 12%', base_contingency),
    ('Reserva de gestion 5%', base_reserve),
    ('Total Recomendado Base', recommended_total),
    ('Costo Laptops', total_laptops_cost),
    ('Cargas Laborales Adicionales', payroll_total),
    ('Subtotal Ampliado', subtotal_expanded),
    ('Contingencia 12% Ampliada', expanded_contingency),
    ('Reserva de gestion 5% Ampliada', expanded_reserve),
    ('Total Proyecto Realista', expanded_total),
]
for label, value in summary_labels:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    left = ws.cell(row=row, column=1, value=label)
    left.font = Font(bold=True)
    left.fill = highlight_fill if ('Realista' in label or 'Ampliad' in label or 'Laptops' in label or 'Cargas' in label) else sub_fill
    left.alignment = Alignment(horizontal='right', vertical='center')
    left.border = border
    right = ws.cell(row=row, column=6, value=value)
    right.font = Font(bold=True)
    right.fill = copy.copy(left.fill)
    right.alignment = Alignment(horizontal='center', vertical='center')
    right.border = border
    money(right)
    row += 1

# Resumen Ejecutivo Costos
we = wb.create_sheet('Resumen Ejecutivo Costos', 1)
for col, width in {'A': 24, 'B': 22, 'C': 14, 'D': 52}.items():
    we.column_dimensions[col].width = width

row = 1
we.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
cell = we.cell(row=row, column=1, value='MATRIZ DE COSTOS REFERENCIAL - CRONOGRAMA 26 SEMANAS')
cell.font = Font(bold=True, color='FFFFFF', size=14)
cell.fill = header_fill
cell.alignment = Alignment(horizontal='center', vertical='center')
row += 2

for note in [
    'Version restaurada con tarifas mensuales por recurso y usabilidad al 100% en todos los periodos.',
    'Se preserva el presupuesto base restaurado y se agrega un presupuesto ampliado realista.',
    'Los costos adicionales incluyen equipamiento y factores de planilla editables.',
]:
    we.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    note_cell = we.cell(row=row, column=1, value=note)
    note_cell.font = Font(italic=True, size=10, color='555555')
    note_cell.alignment = wrap
    row += 1
row += 1

sections = [
    ('PRESUPUESTO BASE RESTAURADO', [
        ('Total Etapa 1', stage1_total, 'USD', 'Costo acumulado de los meses 1 a 4 con 100% de dedicacion.'),
        ('Total Etapa 2', stage2_total, 'USD', 'Costo acumulado de los meses 5 y 6 con 100% de dedicacion.'),
        ('Total Base Laboral', base_total, 'USD', 'Costo laboral base restaurado antes de adicionales.'),
        ('Contingencia 12%', base_contingency, 'USD', 'Cobertura de ajustes y variaciones sobre la base.'),
        ('Reserva de gestion 5%', base_reserve, 'USD', 'Margen ejecutivo sobre la base.'),
        ('Total Recomendado Base', recommended_total, 'USD', 'Presupuesto base restaurado con contingencia y reserva.'),
    ], sub_fill),
    ('PRESUPUESTO AMPLIADO REALISTA', [
        ('Costo Base Laboral', base_total, 'USD', 'Costo laboral base restaurado.'),
        ('Costo Laptops', total_laptops_cost, 'USD', '9 laptops Core i9 con 1TB SSD, una por rol.'),
        ('Cargas Laborales Adicionales', payroll_total, 'USD', 'AFP, EsSalud, CTS, gratificaciones, vacaciones, utilidades y escolaridad.'),
        ('Subtotal Ampliado', subtotal_expanded, 'USD', 'Base + equipamiento + cargas laborales.'),
        ('Contingencia 12% Ampliada', expanded_contingency, 'USD', 'Cobertura sobre subtotal ampliado.'),
        ('Reserva de gestion 5% Ampliada', expanded_reserve, 'USD', 'Margen ejecutivo sobre subtotal ampliado.'),
        ('Total Proyecto Realista', expanded_total, 'USD', 'Presupuesto integral recomendado para aprobacion.'),
    ], highlight_fill),
]
for section_title, items, fill in sections:
    we.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    title_cell = we.cell(row=row, column=1, value=section_title)
    title_cell.font = Font(bold=True, color='FFFFFF', size=11)
    title_cell.fill = section_fill
    title_cell.alignment = Alignment(horizontal='left', vertical='center')
    row += 1
    write_headers(we, row, ['Concepto', 'Monto', 'Moneda', 'Observacion ejecutiva'])
    row += 1
    for index, values in enumerate(items):
        for col, value in enumerate(values, 1):
            cell = we.cell(row=row, column=col, value=value)
            cell.font = Font(size=10, bold=('Total Proyecto Realista' in str(values[0]) or 'Total Recomendado Base' in str(values[0])))
            cell.alignment = wrap
            cell.border = border
            if col == 2:
                money(cell)
            if index % 2 == 0:
                cell.fill = fill
        row += 1
    row += 1

# Detalle por Recurso
wd = wb.create_sheet('Detalle por Recurso')
for col, width in {'A': 12, 'B': 32, 'C': 16, 'D': 10, 'E': 10, 'F': 10, 'G': 10, 'H': 10, 'I': 10, 'J': 16}.items():
    wd.column_dimensions[col].width = width

row = 1
wd.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
cell = wd.cell(row=row, column=1, value='DETALLE DE COSTOS POR RECURSO Y PERIODO')
cell.font = Font(bold=True, color='FFFFFF', size=13)
cell.fill = header_fill
cell.alignment = Alignment(horizontal='center', vertical='center')
row += 2
write_headers(wd, row, ['Rol', 'Descripcion', 'Tarifa Mensual', 'Mes 1', 'Mes 2', 'Mes 3', 'Mes 4', 'Mes 5', 'Mes 6', 'Costo Base Total'])
row += 1

for idx, item in enumerate(role_cost_rows):
    values = [item['rol'], item['descripcion'], item['monthly_rate']] + item['period_costs'] + [item['base']]
    for col, value in enumerate(values, 1):
        cell = wd.cell(row=row, column=col, value=value)
        cell.font = Font(size=10)
        cell.alignment = wrap
        cell.border = border
        if col >= 3:
            money(cell)
        if idx % 2 == 0:
            cell.fill = alt_fill
    row += 1

# Costos por Mes
wm = wb.create_sheet('Costos por Mes')
for col, width in {'A': 12, 'B': 34, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 14, 'I': 16}.items():
    wm.column_dimensions[col].width = width

row = 1
wm.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
cell = wm.cell(row=row, column=1, value='DISTRIBUCION DE COSTOS POR MES Y RECURSO')
cell.font = Font(bold=True, color='FFFFFF', size=13)
cell.fill = header_fill
cell.alignment = Alignment(horizontal='center', vertical='center')
row += 2
write_headers(wm, row, ['Rol', 'Descripcion'] + period_labels + ['Costo Base Total'])
row += 1

for idx, item in enumerate(role_cost_rows):
    values = [item['rol'], item['descripcion']] + item['period_costs'] + [item['base']]
    for col, value in enumerate(values, 1):
        cell = wm.cell(row=row, column=col, value=value)
        cell.font = Font(size=10)
        cell.alignment = wrap
        cell.border = border
        if col >= 3:
            money(cell)
        if idx % 2 == 0:
            cell.fill = alt_fill
    row += 1

for label, value in list(zip(period_labels, monthly_totals)) + [('Costo Base Total', base_total)]:
    merge_to = 2 if label != 'Costo Base Total' else 8
    wm.merge_cells(start_row=row, start_column=1, end_row=row, end_column=merge_to)
    left = wm.cell(row=row, column=1, value=f'Total {label}')
    left.font = Font(bold=True)
    left.fill = sub_fill
    left.alignment = Alignment(horizontal='right', vertical='center')
    left.border = border
    target_col = 3 if label != 'Costo Base Total' else 9
    right = wm.cell(row=row, column=target_col, value=value)
    right.font = Font(bold=True)
    right.fill = sub_fill
    right.alignment = Alignment(horizontal='center', vertical='center')
    right.border = border
    money(right)
    row += 1

# Tarifas Editables
wr = wb.create_sheet('Tarifas Editables')
for col, width in {'A': 12, 'B': 34, 'C': 16, 'D': 32, 'E': 42}.items():
    wr.column_dimensions[col].width = width

row = 1
wr.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
cell = wr.cell(row=row, column=1, value='TARIFAS EDITABLES POR RECURSO')
cell.font = Font(bold=True, color='FFFFFF', size=13)
cell.fill = header_fill
cell.alignment = Alignment(horizontal='center', vertical='center')
row += 2
write_headers(wr, row, ['Rol', 'Descripcion', 'Tarifa Mensual USD', 'Uso sugerido', 'Observacion'])
row += 1

for idx, item in enumerate(roles):
    values = [
        item['rol'],
        item['descripcion'],
        item['monthly_rate'],
        'Editar si existe tarifa contractual validada',
        'Version restaurada con usabilidad al 100% en todos los periodos.',
    ]
    for col, value in enumerate(values, 1):
        cell = wr.cell(row=row, column=col, value=value)
        cell.font = Font(size=10)
        cell.alignment = wrap
        cell.border = border
        if col == 3:
            money(cell)
        if idx % 2 == 0:
            cell.fill = alt_fill
    row += 1

# Costos Adicionales
wc = wb.create_sheet('Costos Adicionales')
for col, width in {'A': 26, 'B': 44, 'C': 16, 'D': 16, 'E': 18, 'F': 46}.items():
    wc.column_dimensions[col].width = width

row = 1
wc.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
cell = wc.cell(row=row, column=1, value='COSTOS ADICIONALES DEL PROYECTO')
cell.font = Font(bold=True, color='FFFFFF', size=13)
cell.fill = header_fill
cell.alignment = Alignment(horizontal='center', vertical='center')
row += 2

wc.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
cell = wc.cell(row=row, column=1, value='Equipamiento')
cell.font = Font(bold=True, color='FFFFFF')
cell.fill = section_fill
cell.alignment = Alignment(horizontal='left', vertical='center')
row += 1
write_headers(wc, row, ['Concepto', 'Descripcion', 'Cantidad', 'Costo Unitario', 'Subtotal', 'Observacion'])
row += 1
for values in [
    ('Laptop Core i9 1TB SSD', 'Equipo asignado a cada rol del proyecto', len(roles) * laptop_units_per_role, laptop_unit_cost, total_laptops_cost, 'Supuesto referencial editable por Finanzas o PMO.'),
]:
    for col, value in enumerate(values, 1):
        cell = wc.cell(row=row, column=col, value=value)
        cell.font = Font(size=10)
        cell.alignment = wrap
        cell.border = border
        if col in (4, 5):
            money(cell)
    row += 1

row += 1
wc.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
cell = wc.cell(row=row, column=1, value='Cargas Laborales y Planilla')
cell.font = Font(bold=True, color='FFFFFF')
cell.fill = section_fill
cell.alignment = Alignment(horizontal='left', vertical='center')
row += 1
write_headers(wc, row, ['Concepto', 'Base de Calculo', 'Factor', 'Monto Base', 'Subtotal', 'Observacion'])
row += 1
for idx, (label, factor, subtotal, description) in enumerate(payroll_rows):
    values = [label, 'Costo Base Laboral', factor, base_total, subtotal, description]
    for col, value in enumerate(values, 1):
        cell = wc.cell(row=row, column=col, value=value)
        cell.font = Font(size=10)
        cell.alignment = wrap
        cell.border = border
        if col == 3:
            cell.number_format = '0.00%'
        if col in (4, 5):
            money(cell)
        if idx % 2 == 0:
            cell.fill = alt_fill
    row += 1

for label, value in [
    ('Total Cargas Laborales', payroll_total),
    ('Subtotal Ampliado', subtotal_expanded),
    ('Contingencia 12% Ampliada', expanded_contingency),
    ('Reserva de gestion 5% Ampliada', expanded_reserve),
    ('Total Proyecto Realista', expanded_total),
]:
    wc.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    left = wc.cell(row=row, column=1, value=label)
    left.font = Font(bold=True)
    left.fill = highlight_fill
    left.alignment = Alignment(horizontal='right', vertical='center')
    left.border = border
    right = wc.cell(row=row, column=5, value=value)
    right.font = Font(bold=True)
    right.fill = highlight_fill
    right.alignment = Alignment(horizontal='center', vertical='center')
    right.border = border
    money(right)
    row += 1

# Usabilidad por Recurso
wu = wb.create_sheet('Usabilidad por Recurso')
for col, width in {'A': 12, 'B': 34, 'C': 12, 'D': 12, 'E': 12, 'F': 12, 'G': 12, 'H': 12}.items():
    wu.column_dimensions[col].width = width

row = 1
wu.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
cell = wu.cell(row=row, column=1, value='USABILIDAD RESTAURADA POR RECURSO')
cell.font = Font(bold=True, color='FFFFFF', size=13)
cell.fill = header_fill
cell.alignment = Alignment(horizontal='center', vertical='center')
row += 2
write_headers(wu, row, ['Rol', 'Descripcion', 'Mes 1', 'Mes 2', 'Mes 3', 'Mes 4', 'Mes 5', 'Mes 6'])
row += 1

for idx, item in enumerate(role_cost_rows):
    values = [item['rol'], item['descripcion']] + item['percentages']
    for col, value in enumerate(values, 1):
        cell = wu.cell(row=row, column=col, value=value)
        cell.font = Font(size=10)
        cell.alignment = wrap
        cell.border = border
        if col >= 3:
            pct(cell)
        if idx % 2 == 0:
            cell.fill = alt_fill
    row += 1

# Supuestos y Reserva
wa = wb.create_sheet('Supuestos y Reserva')
wa.column_dimensions['A'].width = 32
wa.column_dimensions['B'].width = 92
row = 1
for title, text in [
    ('Moneda', currency),
    ('Base de tarifas', 'Tarifas mensuales restauradas segun indicacion del usuario.'),
    ('Usabilidad', 'Todos los recursos quedan al 100% en los seis periodos mensuales.'),
    ('Tarifas restauradas', 'ARQ 1900, BE1 1900, BE2 1900, BE3 1900, FE1 1800, FE2 1800, SYS 1800, QA 1800, IA 1900.'),
    ('Contingencia base', '12% sobre costo base para cubrir variaciones de integracion, pruebas y ajustes.'),
    ('Reserva de gestion base', '5% sobre costo base para decisiones ejecutivas y eventos no planificados de cierre.'),
    ('Laptop por recurso', f'1 equipo Core i9 con 1TB SSD por rol definido en la matriz. Costo unitario referencial: USD {laptop_unit_cost:,.2f}.'),
    ('Factores de planilla', 'AFP 10.00%, EsSalud 9.00%, CTS 9.72%, Gratificaciones 16.67%, Vacaciones 8.33%, Utilidades 5.00%, Escolaridad 2.00%.'),
    ('Uso sugerido', 'Material de soporte para aprobacion gerencial, presupuesto preliminar y afinamiento comercial/tecnico con presupuesto ampliado realista.'),
]:
    c1 = wa.cell(row=row, column=1, value=title)
    c1.font = Font(bold=True)
    c1.fill = sub_fill
    c1.border = border
    c2 = wa.cell(row=row, column=2, value=text)
    c2.alignment = wrap
    c2.border = border
    row += 1

for sheet in wb.worksheets:
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions

wb.save(OUTPUT_XLSX)

with open(OUTPUT_MD, 'w', encoding='utf-8') as fh:
    fh.write('# Matriz de Costos Referencial\n')
    fh.write('## Cronograma de 26 Semanas\n\n')
    fh.write(f'- Moneda base: {currency}\n')
    fh.write('- Tipo de estimacion: referencial editable\n')
    fh.write('- Version restaurada: tarifas mensuales + usabilidad 100%\n')
    fh.write(f'- Total Etapa 1: USD {stage1_total:,.2f}\n')
    fh.write(f'- Total Etapa 2: USD {stage2_total:,.2f}\n')
    fh.write(f'- Total base original restaurado: USD {base_total:,.2f}\n')
    fh.write(f'- Total recomendado base con contingencia y reserva: USD {recommended_total:,.2f}\n')
    fh.write(f'- Costo adicional laptops: USD {total_laptops_cost:,.2f}\n')
    fh.write(f'- Cargas laborales adicionales: USD {payroll_total:,.2f}\n')
    fh.write(f'- Subtotal ampliado: USD {subtotal_expanded:,.2f}\n')
    fh.write(f'- Contingencia 12% sobre subtotal ampliado: USD {expanded_contingency:,.2f}\n')
    fh.write(f'- Reserva de gestion 5% sobre subtotal ampliado: USD {expanded_reserve:,.2f}\n')
    fh.write(f'- Total proyecto realista: USD {expanded_total:,.2f}\n\n')
    fh.write('## Tarifas mensuales restauradas\n')
    for item in roles:
        fh.write(f"- {item['rol']} ({item['descripcion']}): USD {item['monthly_rate']:,.2f}/mes\n")
    fh.write('\n')
    fh.write('## Distribucion referencial por mes\n')
    for label, value in zip(period_labels, monthly_totals):
        fh.write(f'- {label}: USD {value:,.2f}\n')
    fh.write('\n')
    fh.write('## Factores de planilla aplicados\n')
    for label, factor, subtotal, _ in payroll_rows:
        fh.write(f'- {label}: {factor * 100:.2f}% = USD {subtotal:,.2f}\n')
    fh.write('\n')
    fh.write('## Observaciones\n')
    fh.write('- La matriz fue restaurada con tarifas mensuales definidas por el usuario.\n')
    fh.write('- Todos los recursos fueron configurados con usabilidad al 100% en todos los periodos.\n')
    fh.write('- Se mantienen los costos adicionales de equipamiento y planilla para un presupuesto mas realista.\n')

print(OUTPUT_XLSX)
print(OUTPUT_MD)
