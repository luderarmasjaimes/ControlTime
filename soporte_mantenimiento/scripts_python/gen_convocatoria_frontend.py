"""
Genera Convocatoria_Laboral_02_Frontend_LinkedIn.docx
Perfil: Desarrollador Frontend Semi-Senior — Plataforma Industrial de Alto Impacto
Stack real extraído del código fuente del proyecto.
Salida: c:\\Users\\BEEMETRY\\Desktop\\26-05-2026\\Convocatoria_Laboral_02_Frontend_LinkedIn.docx
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Paleta LinkedIn ───────────────────────────────────────────────────────────
C_NEGRO      = RGBColor(0x1B, 0x1B, 0x1B)
C_GRIS       = RGBColor(0x66, 0x66, 0x66)
C_AZUL_LI    = RGBColor(0x00, 0x66, 0xC2)
C_FONDO_TAG  = RGBColor(0xEE, 0xF3, 0xF8)
C_BORDE      = RGBColor(0xCC, 0xCC, 0xCC)
C_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
C_ROJO       = RGBColor(0xC0, 0x39, 0x2B)
C_VERDE      = RGBColor(0x1A, 0x7A, 0x3C)
C_NARANJA    = RGBColor(0xD3, 0x55, 0x00)   # acento Frontend


def h(c: RGBColor) -> str:
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def shading(para, color: RGBColor):
    pPr = para._p.get_or_add_pPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), h(color))
    pPr.append(s)


def border_bottom(para, color=C_BORDE, sz=6):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), str(sz))
    b.set(qn('w:space'), '6')
    b.set(qn('w:color'), h(color))
    pBdr.append(b)
    pPr.append(pBdr)


def parse(text: str):
    parts = text.split('**')
    return [(p, i % 2 == 1) for i, p in enumerate(parts) if p]


# ─────────────────────────────────────────────────────────────────────────────

def build():
    doc = Document()

    for sec in doc.sections:
        sec.top_margin    = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin   = Cm(3.5)
        sec.right_margin  = Cm(3.5)
        sec.page_width    = Cm(21)

    style = doc.styles['Normal']
    style.font.name  = 'Segoe UI'
    style.font.size  = Pt(11)
    style.font.color.rgb = C_NEGRO

    # ── Helpers ──────────────────────────────────────────────────────────────
    def p(before=4, after=4):
        pa = doc.add_paragraph()
        pa.paragraph_format.space_before = Pt(before)
        pa.paragraph_format.space_after  = Pt(after)
        return pa

    def run(para, text, bold=False, size=11, color=C_NEGRO, italic=False):
        r = para.add_run(text)
        r.bold = bold; r.italic = italic
        r.font.name = 'Segoe UI'
        r.font.size = Pt(size)
        r.font.color.rgb = color
        return r

    def runs(para, segs, size=11):
        for txt, bold in segs:
            run(para, txt, bold=bold, size=size)

    def sec_title(icon, title, color=C_AZUL_LI, size=12):
        pa = p(before=14, after=4)
        run(pa, f"{icon} ", size=size, color=color)
        run(pa, title, bold=True, size=size, color=color)
        border_bottom(pa, color=color, sz=6)
        return pa

    def bullet(icon, text, icon_color=C_NEGRO, size=11):
        pa = doc.add_paragraph()
        pa.paragraph_format.left_indent       = Cm(0.6)
        pa.paragraph_format.first_line_indent = Cm(-0.6)
        pa.paragraph_format.space_before = Pt(2)
        pa.paragraph_format.space_after  = Pt(3)
        r = pa.add_run(f"{icon}  ")
        r.font.name = 'Segoe UI'; r.font.size = Pt(size)
        r.font.color.rgb = icon_color
        runs(pa, parse(text), size=size)
        return pa

    def sp():
        pa = doc.add_paragraph()
        pa.paragraph_format.space_before = Pt(0)
        pa.paragraph_format.space_after  = Pt(0)
        run(pa, "", size=4)

    # ══════════════════════════════════════════════════════════════════════════
    # CABECERA
    # ══════════════════════════════════════════════════════════════════════════
    p0 = p(before=0, after=2)
    run(p0, "Se busca:", size=11, color=C_GRIS)

    p1 = p(before=0, after=4)
    run(p1, "Desarrollador Frontend Semi-Senior", bold=True, size=20)

    p2 = p(before=0, after=2)
    run(p2, "Plataforma Industrial de Alto Impacto — React · GIS · Tiempo Real", size=13, color=C_GRIS)

    pb = p(before=4, after=8)
    for badge, sep in [
        ("📍 Remoto / Híbrido", "   "),
        ("⏱ Full Time", "   "),
        ("📄 6 meses renovable", "   "),
        ("🚀 Incorporación inmediata", ""),
    ]:
        run(pb, badge, size=10, color=C_AZUL_LI)
        if sep:
            run(pb, sep, size=10, color=C_GRIS)
    border_bottom(pb, color=C_BORDE, sz=4)

    # ── Intro ─────────────────────────────────────────────────────────────────
    sp()
    pi1 = p(before=6, after=4)
    pi1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(pi1, parse(
        "Organización del sector industrial convoca a un profesional **Frontend Semi-Senior** "
        "para desarrollar la interfaz de usuario de una plataforma web enterprise "
        "de monitoreo operativo minero: dashboards en tiempo real, cartografía GIS interactiva, "
        "editor documental avanzado, videovigilancia y flujos de verificación biométrica."
    ))

    pi2 = p(before=0, after=10)
    pi2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(pi2,
        "¿Disfrutas construir UIs de alto rendimiento con mapas y tiempo real? Este proyecto ya está "
        "en producción y tendrás la guía de un equipo Senior — "
        "tu código llega a operadores en campo desde el primer sprint.",
        italic=True, size=11, color=C_GRIS)

    # ══════════════════════════════════════════════════════════════════════════
    # TU MISIÓN
    # ══════════════════════════════════════════════════════════════════════════
    sec_title("🎯", "TU MISIÓN")
    pm = p(before=4, after=10)
    pm.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(pm,
        "Garantizar que la interfaz de la plataforma sea rápida, confiable y accesible "
        "para operadores mineros en campo: visualización de telemetría en tiempo real, "
        "mapas geoespaciales de alta resolución, editor de informes técnicos multipágina "
        "y flujos de acceso con verificación biométrica — todo funcionando con conectividad limitada.")

    # ══════════════════════════════════════════════════════════════════════════
    # STACK TÉCNICO EN PRODUCCIÓN — bloque destacado
    # ══════════════════════════════════════════════════════════════════════════
    sec_title("⚙️", "STACK TÉCNICO EN PRODUCCIÓN", color=C_NARANJA)
    stack_items = [
        ("React 18", " + Vite 5 · SPA con JSX y hooks modernos, bundling optimizado para producción en nginx."),
        ("MapLibre GL JS", " + Leaflet · Mapas satelitales, capas WMS institucionales (INGEMMET/MINEM/MINAM), MBTiles, GeoJSON y geocercas."),
        ("Tiptap", " · Editor de texto enriquecido para informes técnicos (párrafos, tablas, KPIs, imágenes, sensores en vivo)."),
        ("Three.js", " · Visualización 3D interactiva del frente minero — gemelo digital simplificado con WebSocket."),
        ("WebSockets / SSE", " · Telemetría de sensores industriales en tiempo real (~10 000 puntos), alertas y estado de operaciones."),
        ("Framer Motion", " + Lucide React · Animaciones de UI y sistema de iconografía consistente."),
        ("Axios", " · Capa HTTP hacia API C++ (backend) y formula engine (microservicio independiente)."),
        ("Playwright", " + Vitest · Tests E2E (dashboard, mapas, tabs) y pruebas unitarias de componentes."),
        ("Docker", " + nginx:alpine · Build multi-stage, proxy inverso a backend C++ y formula engine; despliegue en VPS Linux."),
    ]
    for bold_part, rest in stack_items:
        pa = doc.add_paragraph()
        pa.paragraph_format.left_indent = Cm(0.6)
        pa.paragraph_format.first_line_indent = Cm(-0.6)
        pa.paragraph_format.space_before = Pt(2)
        pa.paragraph_format.space_after  = Pt(3)
        run(pa, "▸  ", size=11, color=C_NARANJA)
        run(pa, bold_part, bold=True, size=11)
        run(pa, rest, size=11)
    sp()

    # ══════════════════════════════════════════════════════════════════════════
    # LO QUE HARÁS
    # ══════════════════════════════════════════════════════════════════════════
    sec_title("🔧", "LO QUE HARÁS EN TU DÍA A DÍA")
    tareas = [
        "Desarrollar y mantener componentes React para dashboards de **telemetría en tiempo real**: KPIs, gráficas de series temporales, alertas y estados operativos.",
        "Construir e integrar vistas GIS: mapas satelitales (**MapLibre GL JS / Leaflet**), capas WMS institucionales, GeoJSON, geocercas, rutas de acarreo y zonas de perforación.",
        "Evolucionar el **editor de informes técnicos** (ReportStudio V2): bloques de texto, tablas, KPIs, mapas embebidos, sensores en vivo y exportación a PDF.",
        "Implementar flujos de **verificación biométrica** en UI: captura de rostro, validación ICAO, indicadores de calidad y secuencias de autenticación multi-paso.",
        "Mantener la **escena 3D** (Three.js): gemelo digital del frente minero con actualización en tiempo real vía WebSocket.",
        "Gestionar la capa de **comunicación en tiempo real** (WebSocket / SSE): reconexión automática, buffers de datos de alta frecuencia y rendering sin bloqueo del hilo principal.",
        "Garantizar **rendimiento y accesibilidad**: métricas First Paint < 2 s, 60 fps en mapas, funcionalidad degradada con conectividad intermitente.",
        "Escribir y mantener **tests E2E con Playwright** y pruebas unitarias de componentes con Vitest.",
        "Colaborar en el diseño de **contratos de API** con el equipo backend, documentar componentes reutilizables y mantener el sistema de diseño.",
        "Participar en **revisiones de código, sprints y demos** con metodología ágil (Git + ClickUp).",
    ]
    for t in tareas:
        bullet("◆", t, icon_color=C_AZUL_LI)
    sp()

    # ══════════════════════════════════════════════════════════════════════════
    # REQUISITOS EXCLUYENTES
    # ══════════════════════════════════════════════════════════════════════════
    sec_title("✅", "REQUISITOS EXCLUYENTES", color=C_ROJO)
    pne = p(before=2, after=4)
    shading(pne, RGBColor(0xFF, 0xF0, 0xEF))
    run(pne, "  ⚠  La ausencia de cualquiera de estos puntos descalifica la postulación.",
        size=10, color=C_ROJO, italic=True)

    excl = [
        "**3+ años** de experiencia en desarrollo frontend profesional; mínimo **2 años** con **React** en aplicaciones de producción.",
        "**React 18** con hooks modernos (useCallback, useMemo, useRef, custom hooks), gestión de estado y optimización de re-renders.",
        "Experiencia comprobable con **mapas interactivos en web:** MapLibre GL JS, Leaflet, OpenLayers o equivalente — capas, fuentes, eventos y rendimiento.",
        "Integración de **APIs REST y WebSocket** en tiempo real: manejo de errores, reconexiones y rendering de datos.",
        "**Build tooling moderno:** Vite, Webpack o equivalente; configuración de proxy, variables de entorno y builds para Docker/nginx.",
        "Dominio de **JavaScript ES2022+ / JSX**; comprensión de asincronía, closures, módulos y performance del DOM.",
        "**Testing frontend:** pruebas E2E (Playwright, Cypress o equivalente) y unitarias (Vitest, Jest).",
        "Trabajo colaborativo con **Git** (ramas, PRs, revisiones) y metodología ágil en equipos multidisciplinarios.",
        "**Español fluido; inglés técnico de lectura** (documentación, MDN, specs de W3C).",
    ]
    for e in excl:
        bullet("▸", e, icon_color=C_ROJO)
    sp()

    # ══════════════════════════════════════════════════════════════════════════
    # REQUISITOS IMPORTANTES
    # ══════════════════════════════════════════════════════════════════════════
    sec_title("⭐", "REQUISITOS IMPORTANTES")
    pni = p(before=2, after=4)
    shading(pni, C_FONDO_TAG)
    run(pni, "  Alta valoración — marcan la diferencia en la evaluación.", size=10, color=C_GRIS, italic=True)

    imp = [
        "Experiencia con **editores de contenido enriquecido** (Tiptap, Lexical, ProseMirror, Draft.js o equivalente) en producción.",
        "Conocimiento de **visualización 3D en web:** Three.js, Babylon.js, deck.gl o WebGL — incluso a nivel básico/intermedio.",
        "Comprensión de protocolos de datos geoespaciales: **WMS, GeoJSON, MBTiles, TileJSON** y estándares OGC básicos.",
        "Diseño de **arquitecturas de componentes escalables:** atomic design, compound components, render props y patrones de estado compartido.",
        "**Optimización de rendimiento frontend:** lazy loading, code splitting, virtualización de listas, throttle/debounce y profiling con DevTools.",
        "Experiencia integrando **flujos de cámara web** (MediaStream API, canvas, video) para captura y análisis de imágenes en cliente.",
        "Comprensión básica de **nginx** como servidor de SPA: configuración de proxy, caché de assets y manejo de rutas SPA.",
        "Familiaridad con **accesibilidad (WCAG)** y diseño responsive para pantallas industriales (tablets, monitores de control).",
    ]
    for i in imp:
        bullet("▸", i, icon_color=C_AZUL_LI)
    sp()

    # ══════════════════════════════════════════════════════════════════════════
    # CONOCIMIENTOS DESEABLES
    # ══════════════════════════════════════════════════════════════════════════
    sec_title("💡", "CONOCIMIENTOS DESEABLES", color=C_GRIS)
    pnd = p(before=2, after=4)
    shading(pnd, RGBColor(0xF8, 0xF8, 0xF8))
    run(pnd, "  No excluyentes — altamente valorados.", size=10, color=C_GRIS, italic=True)

    des = [
        "**TypeScript** en proyectos React de producción; tipado estricto en componentes y hooks.",
        "**Framer Motion** u otras librerías de animación declarativa para UX de alto impacto.",
        "Experiencia en **plataformas industriales, SCADA, HMI o dashboards de monitoreo operativo** (minería, energía, utilities).",
        "Procesamiento básico de imágenes en cliente: **Canvas API, WebGL shaders** o integración con librerías de visión por computadora (onnxruntime-web, face-api.js).",
        "Conocimiento de **PWA** y estrategias de caché para operación offline o con conectividad degradada.",
        "Integración con **sistemas de autenticación biométrica** o FIDO2 en flujos de UI.",
        "Diseño de **sistemas de exportación de documentos** (PDF, DOCX) desde el cliente o vía API.",
        "Familiaridad con **Playwright** como herramienta de automatización más allá del testing: demos, screenshots, reportes.",
    ]
    for d in des:
        bullet("▸", d, icon_color=C_GRIS)
    sp()

    # ══════════════════════════════════════════════════════════════════════════
    # QUÉ TE OFRECEMOS
    # ══════════════════════════════════════════════════════════════════════════
    sec_title("🏆", "QUÉ TE OFRECEMOS", color=C_VERDE)
    oferta = [
        ("✔", "Una UI **ya en producción** — no partes de cero: hay componentes, arquitectura y tests que mejorar y escalar."),
        ("✔", "**Mentoría de perfil Senior:** tendrás guía técnica del tech lead y arquitecto del proyecto."),
        ("✔", "Equipo multidisciplinario consolidado (C++ backend, Python AI, GIS, DevOps) con metodología ágil operativa."),
        ("✔", "**Planilla completa desde el primer día.** Contrato inicial de 6 meses, renovable según desempeño."),
        ("✔", "Compensación **S/. 4,500 – S/. 6,500 bruto/mes** (planilla completa, según experiencia demostrada)."),
    ]
    for icon, txt in oferta:
        pa = doc.add_paragraph()
        pa.paragraph_format.left_indent       = Cm(0.6)
        pa.paragraph_format.first_line_indent = Cm(-0.6)
        pa.paragraph_format.space_before = Pt(2)
        pa.paragraph_format.space_after  = Pt(4)
        run(pa, f"{icon}  ", size=11, color=C_VERDE)
        runs(pa, parse(txt))
    sp()

    # ══════════════════════════════════════════════════════════════════════════
    # CÓMO POSTULAR
    # ══════════════════════════════════════════════════════════════════════════
    sec_title("📩", "¿CÓMO POSTULAR?")

    penv = p(before=4, after=2)
    run(penv, "Envía los siguientes elementos al correo ")
    run(penv, "[CORREO DE CONTACTO]", bold=True, size=11, color=C_ROJO)
    run(penv, " con el asunto:")

    pas = p(before=2, after=8)
    pas.paragraph_format.left_indent = Cm(0.8)
    shading(pas, C_FONDO_TAG)
    run(pas, '  "Semi-Senior Frontend — [TU NOMBRE COMPLETO]"',
        bold=True, italic=True, size=11, color=C_AZUL_LI)

    for num, btxt, rest in [
        ("1.", "CV actualizado en PDF", " (máx. 3 páginas)."),
        ("2.", "Portafolio, repositorios o demos en vivo", " (GitHub, GitLab, URL de deploy — muy valorado)."),
        ("3.", "Respuesta breve en el correo", " (máx. 150 palabras) a la pregunta de filtro:"),
    ]:
        pb2 = doc.add_paragraph()
        pb2.paragraph_format.left_indent = Cm(0.8)
        pb2.paragraph_format.space_before = Pt(2)
        pb2.paragraph_format.space_after  = Pt(3)
        run(pb2, f"{num}  ", bold=True, size=11, color=C_AZUL_LI)
        run(pb2, btxt, bold=True, size=11)
        run(pb2, rest, size=11)

    pq = p(before=4, after=4)
    pq.paragraph_format.left_indent  = Cm(1.0)
    pq.paragraph_format.right_indent = Cm(0.5)
    shading(pq, C_FONDO_TAG)
    run(pq,
        '  "Describe una UI de alta complejidad que hayas construido con React: '
        '¿qué problema técnico fue el más difícil (rendimiento, estado, integración de mapas, tiempo real)? '
        '¿Cómo lo resolviste y cuál fue el impacto medible?"',
        italic=True, size=10.5)

    ppr = p(before=6, after=10)
    run(ppr,
        "Los candidatos que respondan esta pregunta tienen prioridad. "
        "Los que no la respondan no serán considerados en la primera revisión.",
        bold=True, size=10.5, color=C_ROJO)

    # ══════════════════════════════════════════════════════════════════════════
    # PROCESO DE SELECCIÓN
    # ══════════════════════════════════════════════════════════════════════════
    sec_title("📋", "PROCESO DE SELECCIÓN")
    sp()

    tbl = doc.add_table(rows=6, cols=3)

    def cell_set(cell, text, bold=False, bg=None,
                 align=WD_ALIGN_PARAGRAPH.LEFT, color=C_NEGRO):
        if bg:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            s = OxmlElement('w:shd')
            s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto')
            s.set(qn('w:fill'), h(bg))
            tcPr.append(s)
        pa = cell.paragraphs[0]
        pa.alignment = align
        pa.paragraph_format.space_before = Pt(4)
        pa.paragraph_format.space_after  = Pt(4)
        r = pa.add_run(text)
        r.bold = bold; r.font.name = 'Segoe UI'
        r.font.size = Pt(9.5); r.font.color.rgb = color

    for i, hdr in enumerate(["FASE", "DURACIÓN", "RESPONSABLE"]):
        cell_set(tbl.rows[0].cells[i], hdr, bold=True,
                 bg=C_AZUL_LI, color=C_WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

    fases = [
        ("Revisión de CV + respuesta técnica + portafolio", "48 h", "Lead Frontend / ARQ"),
        ("Entrevista técnica — React, mapas, rendimiento y arquitectura (45 min)", "1 sesión", "Lead Frontend"),
        ("Challenge técnico asíncrono: componente real del proyecto (máx. 4 h)", "3–5 días", "Candidato"),
        ("Revisión de challenge + entrevista final de encaje (30 min)", "1 sesión", "Lead + PM"),
        ("Oferta, negociación y firma", "48 h", "PM / RRHH"),
    ]
    alt = [RGBColor(0xF0, 0xF6, 0xFD), C_WHITE]
    for ri, (f, d, r_) in enumerate(fases):
        for ci, (val, aln) in enumerate([
            (f,  WD_ALIGN_PARAGRAPH.LEFT),
            (d,  WD_ALIGN_PARAGRAPH.CENTER),
            (r_, WD_ALIGN_PARAGRAPH.CENTER),
        ]):
            cell_set(tbl.rows[ri+1].cells[ci], val, bg=alt[ri % 2], align=aln)

    sp(); sp()

    # ══════════════════════════════════════════════════════════════════════════
    # CHALLENGE DE EJEMPLO
    # ══════════════════════════════════════════════════════════════════════════
    sec_title("🧩", "EJEMPLO DE CHALLENGE TÉCNICO", color=C_NARANJA)
    pch = p(before=4, after=4)
    pch.paragraph_format.left_indent  = Cm(0.8)
    pch.paragraph_format.right_indent = Cm(0.5)
    shading(pch, RGBColor(0xFF, 0xF8, 0xF0))
    run(pch,
        '  "Implementa un componente React que consuma un endpoint WebSocket de telemetría '
        '(recibirá un JSON con sensor_id, value, unit, timestamp cada 500 ms). '
        'Muestra los últimos 60 valores de un sensor seleccionable en un gráfico de línea (sin librerías de charting — SVG o Canvas). '
        'El componente debe manejar reconexión automática, pausar el render cuando la pestaña no está visible '
        'y no superar 50 re-renders por minuto bajo carga continua. '
        'Incluye un test E2E con Playwright que simule 10 mensajes y verifique que el gráfico actualiza. '
        'Tiempo máximo: 4 horas."',
        italic=True, size=10.5, color=C_NEGRO)
    sp(); sp()

    # ── Referidos + hashtags ──────────────────────────────────────────────────
    border_bottom(p(before=4, after=4), color=C_BORDE, sz=4)

    pref = p(before=6, after=4)
    run(pref,
        "🔁 Si conoces a alguien con este perfil, comparte esta publicación. "
        "El mejor talento técnico llega por referidos.",
        size=10.5, color=C_GRIS, italic=True)

    ptags = p(before=6, after=6)
    tags = [
        "#FrontendSemiSenior", "#DesarrolladorFrontend", "#ReactJS", "#MapLibreGL", "#GIS", "#WebSocket",
        "#Vite", "#ThreeJS", "#Tiptap", "#Playwright", "#IndustriaMinera",
        "#TransformacionDigital", "#TechJobsLATAM", "#RemoteWork", "#Hiring",
    ]
    for tag in tags:
        run(ptags, tag + "  ", size=9.5, color=C_AZUL_LI)

    # ── Pie de página ─────────────────────────────────────────────────────────
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(
        "Convocatoria Laboral 02 — Desarrollador Frontend Semi-Senior · Mayo 2026 · Publicación LinkedIn"
    )
    r.font.name = 'Segoe UI'; r.font.size = Pt(8)
    r.font.color.rgb = C_GRIS

    # ── Guardar ───────────────────────────────────────────────────────────────
    import os
    out_dir = r"c:\Users\BEEMETRY\Desktop\26-05-2026"
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "Convocatoria_Laboral_02_Frontend_LinkedIn.docx")
    doc.save(out)
    print(f"✅  Generado: {out}")


if __name__ == "__main__":
    build()
