"""
Genera Convocatoria_Laboral_03_PAF_LinkedIn.docx
Perfil: Analista Funcional & Soporte PMO — Plataforma Industrial Enterprise
Formato LinkedIn: Segoe UI, azul #0066C2, una columna, emojis, tabla proceso.
Salida: c:\\Users\\BEEMETRY\\Desktop\\26-05-2026\\Convocatoria_Laboral_03_PAF_LinkedIn.docx
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

C_NEGRO     = RGBColor(0x1B, 0x1B, 0x1B)
C_GRIS      = RGBColor(0x66, 0x66, 0x66)
C_AZUL_LI   = RGBColor(0x00, 0x66, 0xC2)
C_FONDO_TAG = RGBColor(0xEE, 0xF3, 0xF8)
C_BORDE     = RGBColor(0xCC, 0xCC, 0xCC)
C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
C_ROJO      = RGBColor(0xC0, 0x39, 0x2B)
C_VERDE     = RGBColor(0x1A, 0x7A, 0x3C)
C_PURPURA   = RGBColor(0x6A, 0x0D, 0xAD)   # acento PMO


def h(c): return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"

def shading(para, color):
    pPr = para._p.get_or_add_pPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), h(color)); pPr.append(s)

def border_bottom(para, color=C_BORDE, sz=6):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), str(sz))
    b.set(qn('w:space'), '6'); b.set(qn('w:color'), h(color))
    pBdr.append(b); pPr.append(pBdr)

def parse(text):
    parts = text.split('**')
    return [(p, i % 2 == 1) for i, p in enumerate(parts) if p]


def build():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(3.5); sec.right_margin = Cm(3.5)
        sec.page_width = Cm(21)

    style = doc.styles['Normal']
    style.font.name = 'Segoe UI'
    style.font.size = Pt(11)
    style.font.color.rgb = C_NEGRO

    def p(before=4, after=4):
        pa = doc.add_paragraph()
        pa.paragraph_format.space_before = Pt(before)
        pa.paragraph_format.space_after  = Pt(after)
        return pa

    def run(para, text, bold=False, size=11, color=C_NEGRO, italic=False):
        r = para.add_run(text)
        r.bold = bold; r.italic = italic
        r.font.name = 'Segoe UI'; r.font.size = Pt(size)
        r.font.color.rgb = color; return r

    def runs(para, segs, size=11):
        for txt, bold in segs: run(para, txt, bold=bold, size=size)

    def sec_title(icon, title, color=C_AZUL_LI, size=12):
        pa = p(before=14, after=4)
        run(pa, f"{icon} ", size=size, color=color)
        run(pa, title, bold=True, size=size, color=color)
        border_bottom(pa, color=color, sz=6); return pa

    def bullet(icon, text, icon_color=C_NEGRO, size=11):
        pa = doc.add_paragraph()
        pa.paragraph_format.left_indent       = Cm(0.6)
        pa.paragraph_format.first_line_indent = Cm(-0.6)
        pa.paragraph_format.space_before = Pt(2)
        pa.paragraph_format.space_after  = Pt(3)
        r = pa.add_run(f"{icon}  ")
        r.font.name = 'Segoe UI'; r.font.size = Pt(size)
        r.font.color.rgb = icon_color
        runs(pa, parse(text), size=size); return pa

    def sp():
        pa = doc.add_paragraph()
        pa.paragraph_format.space_before = Pt(0)
        pa.paragraph_format.space_after  = Pt(0)
        run(pa, "", size=4)

    # ── CABECERA ──────────────────────────────────────────────────────────────
    run(p(before=0, after=2), "Se busca:", size=11, color=C_GRIS)
    run(p(before=0, after=4), "Analista Funcional & Soporte PMO", bold=True, size=20)
    run(p(before=0, after=2),
        "Plataforma Industrial Enterprise — ClickUp · Documentación · Gestión de Proyectos TI",
        size=13, color=C_GRIS)

    pb = p(before=4, after=8)
    for badge, sep in [("📍 Remoto / Híbrido","   "),("⏱ Full Time","   "),
                        ("📄 6 meses (mes 1–6)","   "),("🚀 Incorporación inmediata","")]:
        run(pb, badge, size=10, color=C_AZUL_LI)
        if sep: run(pb, sep, size=10, color=C_GRIS)
    border_bottom(pb, color=C_BORDE, sz=4)

    # ── INTRO ─────────────────────────────────────────────────────────────────
    sp()
    pi1 = p(before=6, after=4); pi1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(pi1, parse(
        "Organización del sector industrial convoca a un profesional **Semi-senior a Senior** "
        "para liderar la gestión documental, el seguimiento del cronograma y la comunicación "
        "con stakeholders en un proyecto de software enterprise de alto impacto con equipo "
        "multidisciplinario de 10 especialistas."
    ))
    pi2 = p(before=0, after=10); pi2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(pi2, "¿Destacas en documentación clara, coordinación de equipos y capacitación a usuarios? "
        "Serás el eje operativo entre el equipo técnico, el Project Manager y la gerencia de negocio, "
        "asegurando trazabilidad requisito–entregable en cada sprint.",
        italic=True, size=11, color=C_GRIS)

    # ── ROL EN EL PROYECTO ────────────────────────────────────────────────────
    sec_title("🎯", "TU MISIÓN EN EL PROYECTO", color=C_PURPURA)
    pm = p(before=4, after=8)
    pm.paragraph_format.left_indent = Cm(0.4)
    shading(pm, RGBColor(0xF5, 0xF0, 0xFF))
    run(pm, "  Asegurar la gestión documental, el seguimiento del cronograma, la comunicación "
        "con stakeholders y la transferencia de conocimiento funcional, manteniendo trazabilidad "
        "requisito–entregable y materiales de capacitación actualizados a lo largo de todo el ciclo de vida del proyecto.",
        size=10.5, italic=True)

    # ── LO QUE HARÁS ──────────────────────────────────────────────────────────
    sec_title("🔧", "LO QUE HARÁS EN TU DÍA A DÍA")
    for t in [
        "Actualizar **semanalmente ClickUp:** estados, fechas, dependencias y bloqueos del equipo (287+ tareas activas).",
        "Coordinar reuniones: **convocatorias, agendas, actas, action items** y seguimiento hasta cierre confirmado.",
        "Redactar **manuales funcionales, quick-start** y material de capacitación para usuarios de operación en campo.",
        "Mantener la **matriz de trazabilidad** requisito–tarea–módulo funcional con cobertura total por sprint.",
        "Apoyar la elaboración de reportes de avance, gestión de riesgos y **comunicación quincenal con gerencia.**",
        "Preparar **escenarios de prueba UAT** y log de incidencias funcionales junto al equipo de QA.",
        "Diseñar y ejecutar **talleres de capacitación** a usuarios no técnicos en fases de despliegue.",
        "Compilar el **kit de cierre del proyecto:** índice documental, lecciones aprendidas y acta de transferencia a operaciones.",
        "Administrar el **onboarding documental** de nuevos recursos: RACI, organigrama del proyecto y guía de inducción ClickUp.",
    ]:
        bullet("◆", t, icon_color=C_AZUL_LI)
    sp()

    # ── REQUISITOS EXCLUYENTES ────────────────────────────────────────────────
    sec_title("✅", "REQUISITOS EXCLUYENTES", color=C_ROJO)
    pne = p(before=2, after=4); shading(pne, RGBColor(0xFF, 0xF0, 0xEF))
    run(pne, "  ⚠  La ausencia de cualquiera de estos puntos descalifica la postulación.",
        size=10, color=C_ROJO, italic=True)
    for e in [
        "**5+ años** de experiencia en análisis funcional, documentación técnica o soporte PMO en proyectos de software.",
        "Dominio operativo de **ClickUp, Jira o Azure DevOps** en proyectos con equipos de al menos 5–10 personas.",
        "Producción autónoma de **actas formales, manuales de usuario, procedimientos y matrices RACI.**",
        "Coordinación efectiva de **múltiples stakeholders** técnicos y de negocio (desarrollo, QA, gerencia, cliente).",
        "Experiencia en metodologías **Scrum/Kanban:** facilitar ceremonias, seguimiento de compromisos y bloqueos.",
        "Manejo avanzado de **Excel y Word** para documentación formal, reportes de avance y presentaciones gerenciales.",
        "**Español nativo o fluido** para redacción formal de alta calidad; inglés técnico de lectura.",
        "**Disponibilidad full time durante 6 meses** — inicio inmediato.",
    ]:
        bullet("▸", e, icon_color=C_ROJO)
    sp()

    # ── CONOCIMIENTOS IMPORTANTES ─────────────────────────────────────────────
    sec_title("⭐", "CONOCIMIENTOS IMPORTANTES")
    pni = p(before=2, after=4); shading(pni, C_FONDO_TAG)
    run(pni, "  Alta valoración — marcan la diferencia en la evaluación.", size=10, color=C_GRIS, italic=True)
    for i in [
        "Experiencia como **capacitador de usuarios no TI** en operaciones de campo o plantas industriales.",
        "Conocimiento básico de **arquitectura web** (frontend/backend/base de datos) para dialogar con el equipo técnico.",
        "Gestión de **riesgos y plan de contingencias** en proyectos de software con alta presión de tiempo.",
        "Elaboración de **dashboards de avance** en Excel/ClickUp para reportes ejecutivos quincenales.",
        "Redacción de **especificaciones funcionales y casos de uso** que sirvan de contrato con el equipo técnico.",
    ]:
        bullet("▸", i, icon_color=C_AZUL_LI)
    sp()

    # ── CONOCIMIENTOS DESEABLES ───────────────────────────────────────────────
    sec_title("💡", "CONOCIMIENTOS DESEABLES", color=C_GRIS)
    pnd = p(before=2, after=4); shading(pnd, RGBColor(0xF8, 0xF8, 0xF8))
    run(pnd, "  No excluyentes — altamente valorados.", size=10, color=C_GRIS, italic=True)
    for d in [
        "Certificaciones **PMP, Scrum Master (CSM/PSM) o ITIL Foundation.**",
        "Experiencia en proyectos enterprise en **sector industrial, minería, utilities o energía.**",
        "Documentación de calidad alineada a **ISO 9001** o normativas de trazabilidad sectorial.",
        "Gestión documental integrada con **sistemas de versionado** (Git, Confluence, SharePoint).",
        "Uso básico de **Figma o herramientas de wireframe** para apoyar la documentación de UI/UX.",
    ]:
        bullet("▸", d, icon_color=C_GRIS)
    sp()

    # ── QUÉ TE OFRECEMOS ──────────────────────────────────────────────────────
    sec_title("🏆", "QUÉ TE OFRECEMOS", color=C_VERDE)
    for icon, txt in [
        ("✔", "**Rol transversal con visibilidad directa** ante gerencia y PM — tus entregables documentales impactan decisiones reales."),
        ("✔", "Proyecto enterprise de 6 meses con **equipo multidisciplinario de élite** (ARQ, BE, FE, QA, IA, SYS)."),
        ("✔", "**Autonomía real** para proponer mejoras en procesos de documentación, seguimiento y comunicación."),
        ("✔", "**Planilla completa desde el primer día.** Contrato 6 meses, renovable según desempeño."),
        ("✔", "Compensación **S/. 3,500 – S/. 5,000 bruto/mes** (planilla completa, según experiencia demostrada)."),
    ]:
        pa = doc.add_paragraph()
        pa.paragraph_format.left_indent = Cm(0.6)
        pa.paragraph_format.first_line_indent = Cm(-0.6)
        pa.paragraph_format.space_before = Pt(2); pa.paragraph_format.space_after = Pt(4)
        run(pa, f"{icon}  ", size=11, color=C_VERDE)
        runs(pa, parse(txt))
    sp()

    # ── CÓMO POSTULAR ─────────────────────────────────────────────────────────
    sec_title("📩", "¿CÓMO POSTULAR?")
    penv = p(before=4, after=2)
    run(penv, "Envía los siguientes elementos al correo ")
    run(penv, "[CORREO DE CONTACTO]", bold=True, size=11, color=C_ROJO)
    run(penv, " con el asunto:")

    pas = p(before=2, after=8); pas.paragraph_format.left_indent = Cm(0.8)
    shading(pas, C_FONDO_TAG)
    run(pas, '  "Analista Funcional / Soporte PMO — [TU NOMBRE COMPLETO]"',
        bold=True, italic=True, size=11, color=C_AZUL_LI)

    for num, btxt, rest in [
        ("1.", "CV actualizado en PDF", " (máx. 3 páginas)."),
        ("2.", "Un ejemplo concreto de documentación", " que hayas elaborado (acta, manual, matriz RACI o similar — opcional pero valorado)."),
        ("3.", "Respuesta breve en el correo", " (máx. 150 palabras) a la pregunta de filtro:"),
    ]:
        pb2 = doc.add_paragraph()
        pb2.paragraph_format.left_indent = Cm(0.8)
        pb2.paragraph_format.space_before = Pt(2); pb2.paragraph_format.space_after = Pt(3)
        run(pb2, f"{num}  ", bold=True, size=11, color=C_AZUL_LI)
        run(pb2, btxt, bold=True, size=11); run(pb2, rest, size=11)

    pq = p(before=4, after=4)
    pq.paragraph_format.left_indent = Cm(1.0); pq.paragraph_format.right_indent = Cm(0.5)
    shading(pq, C_FONDO_TAG)
    run(pq, '  "Describe un proyecto de software donde lideraste la documentación funcional y el seguimiento PMO. '
        '¿Cuál fue el entregable más crítico que produjiste, cómo garantizaste su calidad y cuál fue el impacto '
        'concreto en el avance del proyecto o en la satisfacción del cliente?"',
        italic=True, size=10.5)

    ppr = p(before=6, after=10)
    run(ppr, "Los candidatos que respondan esta pregunta tienen prioridad. "
        "Los que no la respondan no serán considerados en la primera revisión.",
        bold=True, size=10.5, color=C_ROJO)

    # ── PROCESO DE SELECCIÓN ──────────────────────────────────────────────────
    sec_title("📋", "PROCESO DE SELECCIÓN"); sp()

    tbl = doc.add_table(rows=6, cols=3)

    def cell_set(cell, text, bold=False, bg=None,
                 align=WD_ALIGN_PARAGRAPH.LEFT, color=C_NEGRO):
        if bg:
            tc = cell._tc; tcPr = tc.get_or_add_tcPr()
            s = OxmlElement('w:shd')
            s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto')
            s.set(qn('w:fill'), h(bg)); tcPr.append(s)
        pa = cell.paragraphs[0]
        pa.alignment = align
        pa.paragraph_format.space_before = Pt(4); pa.paragraph_format.space_after = Pt(4)
        r = pa.add_run(text)
        r.bold = bold; r.font.name = 'Segoe UI'
        r.font.size = Pt(9.5); r.font.color.rgb = color

    for i, hdr in enumerate(["FASE", "DURACIÓN", "RESPONSABLE"]):
        cell_set(tbl.rows[0].cells[i], hdr, bold=True,
                 bg=C_AZUL_LI, color=C_WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

    alt = [RGBColor(0xF0, 0xF6, 0xFD), C_WHITE]
    for ri, (f, d, r_) in enumerate([
        ("Revisión de CV + respuesta a la pregunta de filtro", "48 h", "PM / ARQ"),
        ("Revisión de muestra documental (si se adjuntó)", "24 h", "PM"),
        ("Entrevista funcional — documentación, PMO y herramientas (40 min)", "1 sesión", "PM / ARQ"),
        ("Challenge asíncrono: redacción de acta y matriz RACI (máx. 2 h)", "2–3 días", "Candidato"),
        ("Oferta, negociación y firma", "48 h", "PM / RRHH"),
    ]):
        for ci, (val, aln) in enumerate([
            (f, WD_ALIGN_PARAGRAPH.LEFT),
            (d, WD_ALIGN_PARAGRAPH.CENTER),
            (r_, WD_ALIGN_PARAGRAPH.CENTER),
        ]):
            cell_set(tbl.rows[ri+1].cells[ci], val, bg=alt[ri % 2], align=aln)

    sp(); sp()

    # ── CHALLENGE ─────────────────────────────────────────────────────────────
    sec_title("🧩", "EJEMPLO DE CHALLENGE FUNCIONAL", color=C_PURPURA)
    pch = p(before=4, after=4)
    pch.paragraph_format.left_indent = Cm(0.8); pch.paragraph_format.right_indent = Cm(0.5)
    shading(pch, RGBColor(0xF5, 0xF0, 0xFF))
    run(pch, '  "Eres el Analista Funcional de un sprint de 2 semanas. Al final del sprint se realizó una demo '
        'con el cliente en la que se detectaron 3 incidencias funcionales no registradas previamente. '
        'El PM pide una acta de la sesión, un log de incidencias priorizado y un correo de seguimiento '
        'al cliente con compromisos y fechas. Entrega los tres documentos en formato Word. '
        'Tiempo máximo: 2 horas. No hay una respuesta única — se evalúa claridad, completitud y redacción profesional."',
        italic=True, size=10.5)
    sp(); sp()

    # ── HASHTAGS ──────────────────────────────────────────────────────────────
    border_bottom(p(before=4, after=4), color=C_BORDE, sz=4)
    run(p(before=6, after=4),
        "🔁 Si conoces a alguien con este perfil, comparte esta publicación. "
        "El mejor talento funcional también llega por referidos.",
        size=10.5, color=C_GRIS, italic=True)

    ptags = p(before=6, after=6)
    for tag in ["#AnalistaFuncional","#SoportePMO","#GestionDeProyectos","#ClickUp",
                "#DocumentacionTecnica","#Scrum","#PMO","#UAT","#Capacitacion",
                "#TechJobsLATAM","#RemoteWork","#IndustriaMinera","#Hiring","#TransformacionDigital"]:
        run(ptags, tag + "  ", size=9.5, color=C_AZUL_LI)

    # ── PIE ───────────────────────────────────────────────────────────────────
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Convocatoria Laboral 03 — Analista Funcional & Soporte PMO · Mayo 2026 · Publicación LinkedIn")
    r.font.name = 'Segoe UI'; r.font.size = Pt(8); r.font.color.rgb = C_GRIS

    # ── GUARDAR ───────────────────────────────────────────────────────────────
    out_dir = r"c:\Users\BEEMETRY\Desktop\26-05-2026"
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "Convocatoria_Laboral_03_PAF_LinkedIn.docx")
    doc.save(out)
    print(f"✅  Generado: {out}")


if __name__ == "__main__":
    build()
