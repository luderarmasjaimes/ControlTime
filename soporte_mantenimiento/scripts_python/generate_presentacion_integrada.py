from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import os

OUTPUT = r"C:\InformeCliente\MineriaCampo\Presentacion_Integrada_Aurixa_MineriaCampo.pptx"
OUTPUT_DOCS = r"C:\InformeCliente\docs\Presentacion_Integrada_Aurixa_MineriaCampo.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

COLORS = {
    "navy": RGBColor(17, 34, 64),
    "green": RGBColor(34, 94, 78),
    "sand": RGBColor(245, 239, 224),
    "gold": RGBColor(184, 134, 11),
    "stone": RGBColor(219, 211, 189),
    "slate": RGBColor(74, 85, 104),
    "white": RGBColor(255, 255, 255),
    "ink": RGBColor(28, 37, 44),
}

slides_data = [
    {
        "title": "AURIXA & ERP: Ecosistema Minero Inteligente",
        "subtitle": "Integración de Telemetría IoT, Operaciones de Campo (Getac ZX10) e Inteligencia Artificial",
        "tag": "AURIXA | Visión de Integración de Sistemas Mineros",
        "bullets": [
            "Ecosistema unificado que conecta las operaciones de campo con la gestión estratégica corporativa.",
            "Ingesta masiva de datos en tiempo real con procesamiento analítico de alto rendimiento.",
            "Toma de decisiones basada en diagnósticos automáticos por IA e historial de mantenimiento.",
            "Gobernanza centralizada de datos y facturación inmediata mediante sincronización con el ERP."
        ],
        "highlight": "Estrategia Tecnológica: Unificar la captura de datos en el borde con la gestión y facturación en el ERP corporativo.",
        "type": "title"
    },
    {
        "title": "La Realidad y el Desafío Operativo en Mina",
        "tag": "Diagnóstico del Problema",
        "bullets": [
            "Diagnósticos Lentos: Técnicos en campo sin acceso a manuales ni historial de fallas, elevando el MTTR.",
            "Brecha de Datos: Formularios manuales en papel causan retrasos y errores al registrar datos en el ERP.",
            "Dependencia de Expertos: Cada anomalía requiere trasladar ingenieros especialistas a zonas de difícil acceso.",
            "Silos de Información: Desconexión entre sistemas de telemetría, mantenimiento y facturación de la empresa."
        ],
        "highlight": "La falta de trazabilidad en campo ralentiza el ciclo de facturación y compromete la vida útil de los activos críticos.",
        "type": "content"
    },
    {
        "title": "Arquitectura de Integración de Tres Pilares",
        "tag": "Modelo Objetivo",
        "bullets": [
            "ERP Corporativo (Oficina/Negocio):\nPlanificación, contratos, certificados de calibración y facturación final.",
            "Plataforma AURIXA (Servidor/Nube):\nIngesta de 10k sensores/seg (TimescaleDB Hypertables), motor de fórmulas e IA central.",
            "App Móvil de Campo (Edge/Borde):\nEjecución 100% offline en Getac ZX10, guiado paso a paso, YOLO y OCR local."
        ],
        "image": r"C:\InformeCliente\IMAGENES\platform_architecture.png",
        "type": "image"
    },
    {
        "title": "Ciclo de Vida: De la Ingesta al Cierre Financiero",
        "tag": "Integración de Procesos",
        "bullets": [
            "01. Oficina/ERP: Se crea el proyecto en el ERP, se definen los piezómetros asignados y sus certificados de calibración.",
            "02. Sincronización Móvil: El técnico descarga en la App los proyectos, dispositivos y datos técnicos necesarios.",
            "03. Campo (Offline): Formularios estructurados por etapas (operatividad, saturación, coordenadas, cota).",
            "04. Carga y Facturación: Generación de PDF firmado, envío al ERP y disponibilidad de datos listos para valorización."
        ],
        "highlight": "Un flujo integrado de punta a punta asegura trazabilidad forense 100% auditable y elimina trabajos manuales de oficina.",
        "type": "content"
    },
    {
        "title": "El Técnico Aumentado: IA en el Getac ZX10",
        "tag": "Inteligencia Artificial Edge",
        "bullets": [
            "Reconocimiento YOLO:\nIdentificación de modelos de sensores, estado físico y fallas visuales mediante la cámara.",
            "OCR Automático:\nEscaneo de placas metálicas de sensores y certificados de calibración directamente en campo.",
            "Voz a Texto (Whisper):\nDictado de reportes técnicos manos libres bajo condiciones de ruido extremo.",
            "Asistente IA Local:\nPreguntas y respuestas sobre manuales y procedimientos técnicos sin conexión a internet."
        ],
        "image": r"C:\InformeCliente\IMAGENES\tablet_field_work.png",
        "type": "image"
    },
    {
        "title": "Mantenimiento Inteligente: Correctivo y Predictivo",
        "tag": "Operación Inteligente",
        "bullets": [
            "Mantenimiento Correctivo: Diagnóstico guiado por IA local y procedimientos de reparación en tiempo real.",
            "Tutoría de Instalación: Pasos obligatorios de validación de lecturas (digits/Hz) antes del sellado del pozo.",
            "Mantenimiento Predictivo: Detección temprana de anomalías analizando tendencias históricas en TimescaleDB.",
            "Integración ERP: Envío automático de alertas predictivas que generan órdenes de trabajo preventivas en el ERP."
        ],
        "highlight": "Pasar de mantenimiento reactivo a predictivo reduce el costo de reparaciones imprevistas en un 35% y protege la operación.",
        "type": "content"
    },
    {
        "title": "Informes Técnicos y Auditoría Forense",
        "tag": "Cumplimiento y Calidad",
        "bullets": [
            "Generación Automatizada: Informes PDF por dispositivo al terminar tareas, incluyendo lecturas y fotos de campo.",
            "Firma Digital en Campo: Captura de la firma de conformidad del cliente en la tablet Getac.",
            "Logs Inalterables: Registro de auditoría append-only para garantizar que las lecturas no sean manipuladas.",
            "Respaldo ERP: Almacenamiento directo del PDF y lecturas en el expediente del activo del ERP."
        ],
        "highlight": "La firma en campo y la trazabilidad inalterable aseguran aprobaciones rápidas de valorizaciones y auditorías exitosas.",
        "type": "content"
    },
    {
        "title": "Tutoría del Proceso e Ingesta de Telemetría",
        "tag": "Capacitación y Calidad",
        "bullets": [
            "Guiado de Saturation Check: Pasos específicos para asegurar la saturación correcta del piezómetro antes de instalar.",
            "Monitoreo de Estabilización: Visualización en tiempo real de frecuencias y temperaturas para verificar estabilidad.",
            "Validación de Coordenadas: Registro manual obligatorio de coordenadas de topografía antes del cierre de etapa.",
            "Centro de Conocimiento: Acceso instantáneo a manuales interactivos de sensores (Geokon, etc.) en campo."
        ],
        "highlight": "El sistema guía al técnico paso a paso, eliminando errores humanos de instalación y asegurando la calidad del dato inicial.",
        "type": "content"
    },
    {
        "title": "KPIs y Retorno de la Inversión (ROI)",
        "tag": "Impacto Financiero",
        "bullets": [
            "Reducción de MTTR: Disminución del 40% en tiempos de reparación gracias a diagnósticos asistidos por IA local.",
            "Liberación de Tiempo: Reducción del 30% en horas-hombre administrativas al automatizar reportes en campo.",
            "Integridad de Datos: Coherencia del 99.9% entre los reportes de campo y el sistema de base de datos del ERP.",
            "Aceleración de Cobro: Ciclos de facturación acortados drásticamente al eliminar demoras de aprobación física."
        ],
        "image": r"C:\InformeCliente\IMAGENES\mining_roi_dashboard.png",
        "type": "image"
    },
    {
        "title": "Roadmap de Transformación e Integración",
        "tag": "Hoja de Ruta Estratégica",
        "bullets": [
            "Fase 1 (2026): Integración Core ERP - Plataforma AURIXA - App Móvil para instalación y calibración inicial.",
            "Fase 2 (2027): Telemetría IoT en tiempo real integrada a TimescaleDB e IA de borde (YOLO/OCR/Whisper).",
            "Fase 3 (2028): Módulo de Asistencia Avanzada y Tutorías con Realidad Aumentada sobre Getac ZX10.",
            "Fase 4 (2029-30): Gemelos Digitales de activos críticos en mina y orquestación con agentes autónomos IA."
        ],
        "highlight": "Un roadmap de 5 años enfocado en mitigar riesgos de adopción y consolidar la ventaja tecnológica en LATAM.",
        "type": "content"
    }
]

def add_background(slide, accent=False, title_slide=False):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = COLORS["sand"] if not title_slide else COLORS["navy"]

    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.65))
    band.fill.solid()
    band.fill.fore_color.rgb = COLORS["green"] if accent else COLORS["navy"]
    band.line.fill.background()

    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.35), Inches(6.85), Inches(1.7), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS["gold"]
    bar.line.fill.background()

    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(12.4), Inches(0), Inches(0.28), prs.slide_height)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = COLORS["gold"] if title_slide else COLORS["stone"]
    stripe.line.fill.background()

    if title_slide:
        panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(1.45), Inches(12.0), Inches(4.5))
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(245, 239, 224)
        panel.fill.transparency = 0.06
        panel.line.fill.background()

def add_branding(slide, tag, title_slide=False):
    tag_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(5.4), Inches(0.3))
    tf = tag_box.text_frame
    p = tf.paragraphs[0]
    p.text = tag
    p.font.name = "Aptos"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]

    logo = slide.shapes.add_textbox(Inches(10.55), Inches(0.13), Inches(1.4), Inches(0.32))
    tf = logo.text_frame
    p = tf.paragraphs[0]
    p.text = "AURIXA"
    p.font.name = "Aptos Display"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLORS["gold"] if title_slide else COLORS["white"]
    p.alignment = PP_ALIGN.RIGHT

    footer = slide.shapes.add_textbox(Inches(8.7), Inches(6.68), Inches(3.5), Inches(0.25))
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = "Integración Operativa & ERP | 2026"
    p.font.name = "Aptos"
    p.font.size = Pt(9)
    p.font.color.rgb = COLORS["slate"] if not title_slide else COLORS["white"]
    p.alignment = PP_ALIGN.RIGHT

def add_title(slide, title, subtitle=None, title_slide=False):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0 if title_slide else 0.8), Inches(11.3), Inches(0.9))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Aptos Display"
    p.font.size = Pt(28 if title_slide else 24)
    p.font.bold = True
    p.font.color.rgb = COLORS["navy"]

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.85 if title_slide else 1.45), Inches(11.0), Inches(0.7))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.name = "Aptos"
        p.font.size = Pt(13 if title_slide else 12)
        p.font.color.rgb = COLORS["slate"]

def add_bullets(slide, bullets, width_in=7.7, title_slide=False):
    body = slide.shapes.add_textbox(Inches(0.95), Inches(2.75 if title_slide else 2.0), Inches(width_in), Inches(3.6))
    tf = body.text_frame
    tf.word_wrap = True
    for index, item in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(18 if len(bullets) <= 4 else 16)
        p.font.color.rgb = COLORS["ink"]
        p.space_after = Pt(10)

def add_highlight(slide, text, title_slide=False):
    if title_slide:
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(5.65), Inches(10.4), Inches(0.7))
    else:
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.95), Inches(1.9), Inches(2.9), Inches(3.9))
    box.fill.solid()
    box.fill.fore_color.rgb = COLORS["green"] if title_slide else COLORS["navy"]
    box.fill.transparency = 0.06 if title_slide else 0.02
    box.line.fill.background()

    inner = slide.shapes.add_textbox(
        Inches(1.18) if title_slide else Inches(9.2),
        Inches(5.86) if title_slide else Inches(2.15),
        Inches(9.9) if title_slide else Inches(2.4),
        Inches(0.35) if title_slide else Inches(3.4),
    )
    tf = inner.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(11 if title_slide else 14)
    p.font.bold = True if title_slide else False
    p.font.color.rgb = COLORS["white"]
    p.alignment = PP_ALIGN.LEFT

def add_slide_image(slide, image_path):
    if not os.path.exists(image_path):
        print(f"Warning: Image file not found at {image_path}")
        return
    # Add a thin gray frame behind the image
    frame = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(7.75), Inches(1.85), Inches(4.2), Inches(4.0))
    frame.fill.solid()
    frame.fill.fore_color.rgb = COLORS["stone"]
    frame.line.fill.background()

    # Add the image slightly smaller inside the frame
    slide.shapes.add_picture(image_path, Inches(7.8), Inches(1.9), Inches(4.1), Inches(3.9))

# Compile Presentation
for idx, data in enumerate(slides_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    is_title = (data["type"] == "title")
    is_image = (data["type"] == "image")

    # Set background and layout elements
    add_background(slide, accent=(idx % 2 == 1), title_slide=is_title)
    add_branding(slide, data.get("tag", "Programa Ejecutivo"), title_slide=is_title)
    add_title(slide, data["title"], data.get("subtitle"), title_slide=is_title)

    if is_title:
        add_bullets(slide, data["bullets"], width_in=10.8, title_slide=True)
        add_highlight(slide, data.get("highlight", ""), title_slide=True)
    elif is_image:
        add_bullets(slide, data["bullets"], width_in=6.5, title_slide=False)
        add_slide_image(slide, data["image"])
    else:
        add_bullets(slide, data["bullets"], width_in=7.7, title_slide=False)
        add_highlight(slide, data.get("highlight", ""), title_slide=False)

# Create parent folders if necessary
os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
os.makedirs(os.path.dirname(OUTPUT_DOCS), exist_ok=True)

# Save presentation
prs.save(OUTPUT)
prs.save(OUTPUT_DOCS)
print(f"Successfully generated presentation at:\n1. {OUTPUT}\n2. {OUTPUT_DOCS}")
