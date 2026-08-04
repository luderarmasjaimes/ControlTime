from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT = r"c:\InformeCliente\docs\Presentacion_Ejecutiva_Portafolio_IA_Minera_2026.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

COLORS = {
    "navy": RGBColor(17, 34, 64),
    "green": RGBColor(34, 94, 78),
    "sand": RGBColor(245, 239, 224),
    "gold": RGBColor(184, 134, 11),
    "stone": RGBColor(219, 211, 189),
    "slate": RGBColor(74, 85, 104),
    "white": RGBColor(255, 255, 255),
    "ink": RGBColor(28, 37, 44),
}

slides = [
    {
        "title": "Portafolio IA Minero 2026",
        "subtitle": "Plataforma nueva en VPS + continuidad AWS + productividad compartida de ingenieria",
        "tag": "AURIXA | Directorio y Gerencia TI",
        "bullets": [
            "No se recomienda una sola plataforma de IA; se recomienda un portafolio multivendor gobernado por TI.",
            "La prioridad es combinar copilotos, modelos API, document AI y ML propio sobre una misma capa de acceso.",
            "La estrategia propuesta protege costo, continuidad, seguridad y velocidad de ejecucion.",
            "La adopcion se plantea por trimestres para reducir riesgo y construir capacidad interna.",
        ],
        "highlight": "Decision recomendada: aprobar un programa IA compartido y gobernado, no una compra aislada de licencias.",
    },
    {
        "title": "Problema de negocio que se resuelve",
        "tag": "Contexto Ejecutivo",
        "bullets": [
            "Construir una nueva plataforma minera sobre VPS sin perder continuidad con la plataforma actual en AWS.",
            "Aumentar productividad de ingenieria sobre varios proyectos de automatizacion y software minero.",
            "Habilitar document AI, RAG, biometria, telemetria y asistentes operativos con una sola estrategia de gobierno.",
            "Preparar la base para evolucionar a modelos predictivos mineros con datos propios.",
        ],
        "highlight": "El retorno esperado no viene solo de escribir mejor codigo, sino de crear una plataforma compartida de capacidades IA.",
    },
    {
        "title": "Arquitectura recomendada por capas",
        "tag": "Modelo Objetivo",
        "bullets": [
            "Copiloto corporativo base: GitHub Copilot para adopcion amplia y controlada.",
            "Capa senior de ingenieria: Claude Code Team o Windsurf Teams para cambios complejos y agentes de desarrollo.",
            "Modelos premium: Claude Sonnet 4.6 y GPT-5.4 para razonamiento, documentos y agentes empresariales.",
            "Volumen y costo: Gemini 2.5 Flash; portabilidad: Mistral Medium 3.1 y Codestral; continuidad AWS: Amazon Bedrock.",
        ],
        "highlight": "La combinacion correcta no es un proveedor unico; es un portafolio con routing segun costo, riesgo y caso de uso.",
    },
    {
        "title": "Lectura de scoring por audiencia",
        "tag": "Comite de Decision",
        "bullets": [
            "Directorio prioriza gobierno enterprise, control de costo y escalamiento seguro.",
            "Gerencia TI prioriza productividad de desarrollo, integracion con el stack y portabilidad hybrid.",
            "Operaciones mineras prioriza document AI, asistentes y compatibilidad con flujos reales de uso.",
            "La matriz Excel ya incorpora estos tres perfiles de pesos para ordenar el shortlist segun cada audiencia.",
        ],
        "highlight": "La misma tecnologia no se evalua igual para directorio, TI y operaciones; por eso se definieron pesos distintos.",
    },
    {
        "title": "Presupuesto orientativo y escenarios",
        "tag": "Costo y Gobierno",
        "bullets": [
            "Licencias base y premium: desde una capa de entrada controlada hasta un modelo senior de alta exigencia.",
            "API mensual: de PoC y arranque hacia operacion estable con RAG, document AI y asistentes internos.",
            "Escenario conservador anual: USD 44,130.",
            "Escenario medio anual: alrededor de USD 83,000. Escenario expandido: USD 103,050.",
        ],
        "highlight": "El costo debe tratarse como programa escalable con FinOps IA, no como gasto fijo sin gobierno.",
    },
    {
        "title": "Comparativa de decision: conservadora vs agresiva",
        "tag": "Opciones de Aprobacion",
        "bullets": [
            "Ruta conservadora: GitHub Copilot + Azure OpenAI/Gemini 2.5 Flash + Amazon Bedrock, con foco en control, gobierno y TCO.",
            "Ruta agresiva: Claude Sonnet 4.6 + GPT-5.4 + Windsurf Teams/Mistral, con foco en velocidad, amplitud funcional y ventaja tecnica.",
            "La ruta conservadora reduce riesgo presupuestal y simplifica el arranque corporativo.",
            "La ruta agresiva acelera construccion, experimentacion y sofisticacion de agentes, pero exige mas disciplina de costo.",
        ],
        "highlight": "Recomendacion: aprobar base conservadora corporativa y habilitar una celula agresiva controlada para arquitectura e ingenieria senior.",
    },
    {
        "title": "Top 5 finalistas para comite",
        "tag": "Shortlist Ejecutivo",
        "bullets": [
            "Gemini 2.5 Flash: mejor equilibrio entre costo, multimodalidad y volumen operativo.",
            "Azure OpenAI: fuerte opcion enterprise para gobierno, seguridad y residencia operacional.",
            "Gemini 2.5 Pro: mejor perfil para analisis documental y razonamiento largo.",
            "Mistral Medium 3.1: mejor lectura para portabilidad hybrid y control de TCO.",
            "Claude Sonnet 4.6: mejor perfil premium para codigo, documentos y calidad tecnica.",
        ],
        "highlight": "El shortlist de comite concentra plataformas defendibles para aprobacion inicial, sin dispersar la decision en demasiados proveedores.",
    },
    {
        "title": "Roadmap trimestral de ejecucion",
        "tag": "Implementacion 2026",
        "bullets": [
            "T1: copilotos, gateway IA, seguridad y RAG inicial.",
            "T2: document AI, asistentes operativos, observabilidad e integracion AWS/VPS.",
            "T3: plataforma ML, prediccion de eventos, vision y MLOps.",
            "T4: escalamiento corporativo, gobierno final, optimizacion de TCO y continuidad reforzada.",
        ],
        "highlight": "La ruta correcta es productividad primero, industrializacion despues y modelos propios al consolidar datos y gobierno.",
    },
    {
        "title": "Riesgos y mitigaciones para directorio",
        "tag": "Gobierno de Riesgo",
        "bullets": [
            "Riesgo de dispersion de proveedores. Mitigacion: gateway unico, politicas de aprobacion y FinOps IA centralizado.",
            "Riesgo de sobrecosto por consumo API. Mitigacion: cuotas, routing por costo, cache y seguimiento quincenal.",
            "Riesgo de fuga de datos o uso no gobernado. Mitigacion: SSO, auditoria, entornos controlados y politicas por caso de uso.",
            "Riesgo de baja adopcion o baja calidad operativa. Mitigacion: pilotos acotados, playbooks y medicion trimestral de valor.",
        ],
        "highlight": "El riesgo principal no es usar IA, sino hacerlo sin gobierno, sin control de costo y sin disciplina de adopcion.",
    },
    {
        "title": "Decision solicitada",
        "tag": "Cierre Ejecutivo",
        "bullets": [
            "Aprobar el portafolio IA multivendor como estandar de trabajo para la plataforma minera 2026.",
            "Autorizar presupuesto inicial y gobierno trimestral por escenarios conservador, medio y expandido.",
            "Nombrar a Gerencia TI como dueña del gateway IA, seguridad, observabilidad y control de proveedores.",
            "Iniciar con pilotos controlados y comite quincenal de avance, costo y riesgo.",
        ],
        "footer": "Recomendacion: aprobar el programa IA por fases y evaluar resultados por trimestre con metrica de costo, productividad y adopcion.",
        "highlight": "La ventaja competitiva no vendra de un modelo aislado, sino de una capacidad corporativa de IA bien gobernada.",
    },
]


def add_background(slide, accent=False, title_slide=False):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = COLORS["sand"] if not title_slide else COLORS["navy"]

    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.65))
    band.fill.solid()
    band.fill.fore_color.rgb = COLORS["green"] if accent else COLORS["navy"]
    band.line.fill.background()

    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.35), Inches(6.85), Inches(1.7), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS["gold"]
    bar.line.fill.background()

    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(12.4), Inches(0), Inches(0.28), prs.slide_height)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = COLORS["gold"] if title_slide else COLORS["stone"]
    stripe.line.fill.background()

    if title_slide:
        panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(1.45), Inches(12.0), Inches(4.5))
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(245, 239, 224)
        panel.fill.transparency = 0.06
        panel.line.fill.background()


def add_branding(slide, tag, title_slide=False):
    tag_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(5.4), Inches(0.3))
    tf = tag_box.text_frame
    p = tf.paragraphs[0]
    p.text = tag
    p.font.name = "Aptos"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]

    logo = slide.shapes.add_textbox(Inches(10.55), Inches(0.13), Inches(1.4), Inches(0.32))
    tf = logo.text_frame
    p = tf.paragraphs[0]
    p.text = "AURIXA"
    p.font.name = "Aptos Display"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLORS["gold"] if title_slide else COLORS["white"]
    p.alignment = PP_ALIGN.RIGHT

    footer = slide.shapes.add_textbox(Inches(8.7), Inches(6.68), Inches(3.5), Inches(0.25))
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = "Programa IA Minera | 2026"
    p.font.name = "Aptos"
    p.font.size = Pt(9)
    p.font.color.rgb = COLORS["slate"] if not title_slide else COLORS["white"]
    p.alignment = PP_ALIGN.RIGHT


def add_title(slide, title, subtitle=None, title_slide=False):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0 if title_slide else 0.8), Inches(11.3), Inches(0.9))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Aptos Display"
    p.font.size = Pt(28 if title_slide else 24)
    p.font.bold = True
    p.font.color.rgb = COLORS["navy"]

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.85 if title_slide else 1.45), Inches(11.0), Inches(0.7))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.name = "Aptos"
        p.font.size = Pt(13 if title_slide else 12)
        p.font.color.rgb = COLORS["slate"]


def add_bullets(slide, bullets, footer=None, title_slide=False):
    body = slide.shapes.add_textbox(Inches(0.95), Inches(2.75 if title_slide else 2.0), Inches(7.7 if not title_slide else 10.8), Inches(3.6))
    tf = body.text_frame
    tf.word_wrap = True
    for index, item in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(20 if len(bullets) <= 4 else 18)
        p.font.color.rgb = COLORS["ink"]
        p.space_after = Pt(10)

    if footer:
        foot = slide.shapes.add_textbox(Inches(0.85), Inches(6.1), Inches(11.6), Inches(0.55))
        tf = foot.text_frame
        p = tf.paragraphs[0]
        p.text = footer
        p.font.name = "Aptos"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLORS["green"]
        p.alignment = PP_ALIGN.LEFT


def add_highlight(slide, text, title_slide=False):
    if title_slide:
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(5.65), Inches(10.4), Inches(0.7))
    else:
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.95), Inches(1.9), Inches(2.9), Inches(3.9))
    box.fill.solid()
    box.fill.fore_color.rgb = COLORS["green"] if title_slide else COLORS["navy"]
    box.fill.transparency = 0.06 if title_slide else 0.02
    box.line.fill.background()

    inner = slide.shapes.add_textbox(
        Inches(1.18) if title_slide else Inches(9.2),
        Inches(5.86) if title_slide else Inches(2.15),
        Inches(9.9) if title_slide else Inches(2.4),
        Inches(0.35) if title_slide else Inches(3.4),
    )
    tf = inner.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(11 if title_slide else 14)
    p.font.bold = True if title_slide else False
    p.font.color.rgb = COLORS["white"]
    p.alignment = PP_ALIGN.LEFT


for idx, data in enumerate(slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    is_title = idx == 0
    add_background(slide, accent=(idx % 2 == 1), title_slide=is_title)
    add_branding(slide, data.get("tag", "Programa Ejecutivo"), title_slide=is_title)
    add_title(slide, data["title"], data.get("subtitle"), title_slide=is_title)
    add_bullets(slide, data["bullets"], data.get("footer"), title_slide=is_title)
    add_highlight(slide, data.get("highlight", ""), title_slide=is_title)

prs.save(OUTPUT)
print(OUTPUT)