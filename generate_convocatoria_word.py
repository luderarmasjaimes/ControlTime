"""
Genera CONVOCATORIA_LINKEDIN.docx — Documento Word profesional listo para presentar.
Ejecutar: python generate_convocatoria_word.py
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── Paleta de colores corporativos ───────────────────────────────────────────
C_NAVY      = RGBColor(0x0F, 0x2A, 0x4A)   # Azul oscuro (encabezados primarios)
C_TEAL      = RGBColor(0x00, 0x7A, 0x8A)   # Teal (encabezados de sección)
C_GOLD      = RGBColor(0xC9, 0x9A, 0x06)   # Dorado (acentos / iconos)
C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT_BG  = RGBColor(0xF2, 0xF7, 0xFA)   # Fondo secciones alternadas
C_RED_EXCL  = RGBColor(0xC0, 0x39, 0x2B)   # Rojo — requisitos excluyentes
C_GREEN_OK  = RGBColor(0x1A, 0x7A, 0x3C)   # Verde — qué se ofrece
C_GRAY_TEXT = RGBColor(0x44, 0x44, 0x44)
C_TAG_BG    = RGBColor(0xE8, 0xF4, 0xF8)   # Fondo etiqueta hashtag


def hex_to_rgb_str(color: RGBColor) -> str:
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_cell_bg(cell, color: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_to_rgb_str(color))
    tcPr.append(shd)


def set_para_shading(para, color: RGBColor):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_to_rgb_str(color))
    pPr.append(shd)


def set_para_border(para, side='bottom', color: RGBColor = C_TEAL, sz=12):
    pPr = para._p.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is None:
        pBdr = OxmlElement('w:pBdr')
        pPr.append(pBdr)
    bdr = OxmlElement(f'w:{side}')
    bdr.set(qn('w:val'), 'single')
    bdr.set(qn('w:sz'), str(sz))
    bdr.set(qn('w:space'), '4')
    bdr.set(qn('w:color'), hex_to_rgb_str(color))
    pBdr.append(bdr)


def remove_para_spacing(para):
    pPr = para._p.get_or_add_pPr()
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), '0')
    spacing.set(qn('w:after'), '0')
    pPr.append(spacing)


def add_run_bold(para, text, color=None, size=None):
    run = para.add_run(text)
    run.bold = True
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return run


def add_run_normal(para, text, color=None, size=None, italic=False):
    run = para.add_run(text)
    run.bold = False
    run.italic = italic
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return run


def section_heading(doc, icon, title, bg_color=C_NAVY, text_color=C_WHITE, size=13):
    """Encabezado de sección con fondo de color."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = para.paragraph_format
    pf.space_before = Pt(14)
    pf.space_after = Pt(6)
    pf.left_indent = Cm(0.3)
    set_para_shading(para, bg_color)
    run = para.add_run(f"  {icon}  {title}  ")
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = text_color
    run.font.name = 'Calibri'
    return para


def bullet_item(doc, icon, text_parts, indent=Cm(0.8), size=10.5, after=Pt(4)):
    """Ítem de lista con icono y texto (text_parts: list de (texto, bold))."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = indent
    para.paragraph_format.first_line_indent = Cm(-0.5)
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = after
    # Icono
    r = para.add_run(f"{icon}  ")
    r.font.size = Pt(size)
    r.font.color.rgb = C_TEAL
    # Texto con porciones bold/normal
    for seg_text, is_bold in text_parts:
        r = para.add_run(seg_text)
        r.font.size = Pt(size)
        r.bold = is_bold
        r.font.color.rgb = C_GRAY_TEXT
    return para


def divider(doc, color=C_TEAL):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(4)
    set_para_border(para, 'bottom', color, sz=8)
    return para


def note_box(doc, text, bg=C_LIGHT_BG, size=9.5, italic=True):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.5)
    para.paragraph_format.right_indent = Cm(0.5)
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(6)
    set_para_shading(para, bg)
    r = para.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    r.font.color.rgb = C_GRAY_TEXT
    return para


def parse_bold_segments(text: str):
    """
    Parsea segmentos **bold** dentro de un string y devuelve
    lista de (texto, is_bold).
    """
    segments = []
    parts = text.split('**')
    for i, part in enumerate(parts):
        if part:
            segments.append((part, i % 2 == 1))
    return segments


# ─────────────────────────────────────────────────────────────────────────────

def build_document():
    doc = Document()

    # ── Márgenes del documento ───────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # Fuente por defecto
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style.font.color.rgb = C_GRAY_TEXT

    # ══════════════════════════════════════════════════════════════════════════
    # BLOQUE DE CABECERA — Título principal con fondo navy
    # ══════════════════════════════════════════════════════════════════════════
    header_tbl = doc.add_table(rows=1, cols=1)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = header_tbl.cell(0, 0)
    set_cell_bg(cell, C_NAVY)

    # Eliminar bordes de la tabla
    tbl_xml = header_tbl._tbl
    tblPr = tbl_xml.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl_xml.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        bdr = OxmlElement(f'w:{border_name}')
        bdr.set(qn('w:val'), 'none')
        tblBorders.append(bdr)
    tblPr.append(tblBorders)

    # Padding interno de la celda
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side in ('top', 'bottom', 'left', 'right'):
        m = OxmlElement(f'w:{side}')
        m.set(qn('w:w'), '180')
        m.set(qn('w:type'), 'dxa')
        tcMar.append(m)
    tcPr.append(tcMar)

    # Línea dorada superior (decorativa — párrafo vacío en otra tabla)
    # Se simula con un párrafo de fondo dorado encima
    gold_tbl = doc.add_table(rows=1, cols=1)
    gold_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    gold_cell = gold_tbl.cell(0, 0)
    set_cell_bg(gold_cell, C_GOLD)
    gp = gold_cell.paragraphs[0]
    remove_para_spacing(gp)
    gp.paragraph_format.space_before = Pt(3)
    gp.paragraph_format.space_after = Pt(3)
    # Mover gold_tbl ANTES del header_tbl en el XML
    doc._body._body.remove(gold_tbl._tbl)
    doc._body._body.remove(header_tbl._tbl)
    doc._body._body.insert(0, gold_tbl._tbl)
    doc._body._body.insert(1, header_tbl._tbl)

    # Textos dentro del header_tbl
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(10)
    p1.paragraph_format.space_after = Pt(2)
    r = p1.add_run("CONVOCATORIA LABORAL")
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.bold = False
    r.font.color.rgb = RGBColor(0xAA, 0xCC, 0xDD)
    r.font.all_caps = True

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(4)
    r2 = p2.add_run("Arquitecto de Datos & Backend Senior")
    r2.font.name = 'Calibri'
    r2.font.size = Pt(22)
    r2.bold = True
    r2.font.color.rgb = C_WHITE

    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(4)
    r3 = p3.add_run("Plataforma Industrial de Alto Impacto")
    r3.font.name = 'Calibri'
    r3.font.size = Pt(13)
    r3.bold = False
    r3.font.color.rgb = RGBColor(0xAA, 0xD8, 0xE8)

    p4 = cell.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(4)
    p4.paragraph_format.space_after = Pt(10)
    for seg, sep in [("Remoto / Híbrido", "  ·  "), ("Full Time", "  ·  "), ("6 meses renovable", "  ·  "), ("Incorporación Inmediata", "")]:
        r4 = p4.add_run(seg)
        r4.font.size = Pt(10)
        r4.font.color.rgb = C_GOLD
        r4.bold = True
        if sep:
            rs = p4.add_run(sep)
            rs.font.size = Pt(10)
            rs.font.color.rgb = C_WHITE

    # ══════════════════════════════════════════════════════════════════════════
    # INTRO
    # ══════════════════════════════════════════════════════════════════════════
    doc.add_paragraph()
    p_intro = doc.add_paragraph()
    p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro.paragraph_format.space_before = Pt(8)
    p_intro.paragraph_format.space_after = Pt(6)
    r = p_intro.add_run(
        "Organización del sector industrial convoca a un profesional "
    )
    r.font.size = Pt(11)
    rb = p_intro.add_run("Senior")
    rb.bold = True
    rb.font.size = Pt(11)
    r2 = p_intro.add_run(
        " para liderar la capa de datos, persistencia y servicios backend de una plataforma "
        "web enterprise de monitoreo operativo, reportabilidad técnica y gestión documental avanzada."
    )
    r2.font.size = Pt(11)

    p_hook = doc.add_paragraph()
    p_hook.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_hook.paragraph_format.space_before = Pt(0)
    p_hook.paragraph_format.space_after = Pt(10)
    r = p_hook.add_run(
        "¿Te apasionan la arquitectura de datos, la continuidad operativa y la trazabilidad "
        "en entornos con conectividad limitada? Esta es tu oportunidad de liderar un proyecto "
        "de transformación digital de alto impacto en la región."
    )
    r.font.size = Pt(11)
    r.italic = True
    r.font.color.rgb = C_TEAL

    # ══════════════════════════════════════════════════════════════════════════
    # TU MISIÓN
    # ══════════════════════════════════════════════════════════════════════════
    section_heading(doc, "🎯", "TU MISIÓN", bg_color=C_TEAL)
    p_mision = doc.add_paragraph()
    p_mision.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_mision.paragraph_format.left_indent = Cm(0.5)
    p_mision.paragraph_format.space_before = Pt(4)
    p_mision.paragraph_format.space_after = Pt(10)
    r = p_mision.add_run(
        "Garantizar que la arquitectura de datos y los servicios de núcleo soporten de forma "
        "confiable: informes técnicos estructurados, telemetría operativa de alta frecuencia, "
        "auditoría completa de acciones, operación offline con sincronización controlada e "
        "integración segura con los sistemas existentes de planta."
    )
    r.font.size = Pt(11)

    # ══════════════════════════════════════════════════════════════════════════
    # LO QUE HARÁS
    # ══════════════════════════════════════════════════════════════════════════
    section_heading(doc, "🔧", "LO QUE HARÁS EN TU DÍA A DÍA", bg_color=C_NAVY)

    responsabilidades = [
        "Diseñar e implementar modelos de datos transaccionales, documentales y de series temporales con soporte **multi-empresa**, auditoría y políticas de retención.",
        "Desarrollar y mantener servicios backend de **alto rendimiento**: APIs REST, canales de comunicación en tiempo real y autoservicio documental.",
        "Construir pipelines **ETL/ELT** e integraciones con fuentes legacy e IoT/industrial, garantizando idempotencia y monitoreo de rezagos.",
        "Diseñar estrategias de persistencia local, colas de cambios pendientes y mecanismos de reconciliación al restablecer la conectividad.",
        "Administrar entornos **Linux y contenedores**: backups probados, planes de recuperación ante desastres y migraciones de esquema versionadas.",
        "Aplicar controles de seguridad en capa de datos: segregación por tenant (**Row-Level Security**), cifrado en tránsito y en reposo, bitácoras inmutables y cumplimiento normativo.",
        "Definir estándares, lineamientos y buenas prácticas de arquitectura de datos; guiar y asesorar a los equipos de desarrollo, QA y PMO.",
        "Colaborar en contratos de integración con frontend, revisiones de carga y documentación operativa.",
        "Participar activamente en metodología ágil con **Git** y herramienta de gestión de proyectos (ClickUp o equivalente).",
    ]
    for item in responsabilidades:
        bullet_item(doc, "◆", parse_bold_segments(item))

    # ══════════════════════════════════════════════════════════════════════════
    # REQUISITOS EXCLUYENTES
    # ══════════════════════════════════════════════════════════════════════════
    section_heading(doc, "✅", "REQUISITOS EXCLUYENTES", bg_color=C_RED_EXCL)
    note_box(doc, "  ⚠  La ausencia de cualquiera de estos puntos descalifica la postulación.", bg=RGBColor(0xFD, 0xED, 0xEC))

    excluyentes = [
        "**8+ años** de experiencia en backend y arquitectura de datos; mínimo **5 años** con motor relacional en producción.",
        "**Dominio avanzado de PostgreSQL:** modelado relacional y documental (JSON/JSONB), partición de tablas, índices avanzados, replicación, backup/restore y tuning de consultas.",
        "Experiencia comprobable en servicios backend de **alto rendimiento**, APIs REST y comunicación en tiempo real (WebSockets, SSE o equivalente).",
        "Diseño de **trazabilidad forense, auditoría de acciones y gobierno de datos** alineado a estándares de Industria 4.0/5.0.",
        "Pipelines **ETL/ELT**, integración con sistemas externos (legacy e IoT) y sincronización incremental controlada.",
        "Seguridad aplicada a datos: **RBAC, cifrado, OWASP** aplicado a APIs y accesos, hardening de bases de datos.",
        "Operación fluida en **Linux, Docker y entornos productivos con CI/CD.**",
        "Experiencia en **proyectos industriales:** manufactura, energía, utilities, minería o plantas automatizadas.",
        "Trabajo en equipos multidisciplinarios con **Git y gestión ágil** de proyectos.",
        "**Español fluido; inglés técnico de lectura** (documentación, RFCs, whitepapers).",
    ]
    for item in excluyentes:
        bullet_item(doc, "▸", parse_bold_segments(item))

    # ══════════════════════════════════════════════════════════════════════════
    # REQUISITOS IMPORTANTES
    # ══════════════════════════════════════════════════════════════════════════
    section_heading(doc, "⭐", "REQUISITOS IMPORTANTES", bg_color=C_NAVY)
    note_box(doc, "  Alta valoración — marcan la diferencia en la evaluación.")

    importantes = [
        "Diseño de arquitecturas de datos end-to-end: ingestión, procesamiento, almacenamiento y consumo, con foco en **escalabilidad, resiliencia y seguridad.**",
        "Conocimiento aplicado de patrones de arquitectura: **Data Lake, Lakehouse, Data Warehouse,** procesamiento batch y streaming.",
        "Experiencia con bases de datos **relacionales** (PostgreSQL, SQL Server, MySQL) **y NoSQL** (MongoDB, Redis o equivalentes), con criterios claros de selección por caso de uso.",
        "Técnicas de **modelado de datos** (ERD, diagramas UML, Draw.io) y planificación de capacidad.",
        "**SQL avanzado** y Python para análisis, scripting y soporte a decisiones de arquitectura.",
        "Capacidad para definir lineamientos técnicos y **asesorar a perfiles junior y mid** (Data Engineers, BI, DevOps).",
    ]
    for item in importantes:
        bullet_item(doc, "▸", parse_bold_segments(item))

    # ══════════════════════════════════════════════════════════════════════════
    # CONOCIMIENTOS DESEABLES
    # ══════════════════════════════════════════════════════════════════════════
    section_heading(doc, "💡", "CONOCIMIENTOS DESEABLES", bg_color=C_TEAL)
    note_box(doc, "  No excluyentes, pero altamente valorados.")

    deseables = [
        "Motores de **series temporales** (TimescaleDB, InfluxDB) y telemetría de alta frecuencia.",
        "Plataformas de **mensajería/streaming** para ingesta ordenada de eventos (Kafka, Redpanda, NATS o equivalentes).",
        "**Multi-tenancy** con Row-Level Security y políticas de retención por regulación legal.",
        "Modelado de eventos con **validación biométrica** integrada a flujos de acceso.",
        "Experiencia en operaciones con **conectividad intermitente** y reconciliación de datos offline.",
        "Certificaciones en **DBA, seguridad de la información** o normas IEC 62443.",
        "Marcos **ISO** aplicables a calidad de software y evidencias de trazabilidad documental.",
        "Conocimientos conceptuales de infraestructura como código (IaC) y orquestación de contenedores.",
    ]
    for item in deseables:
        bullet_item(doc, "▸", parse_bold_segments(item))

    # ══════════════════════════════════════════════════════════════════════════
    # QUÉ TE OFRECEMOS
    # ══════════════════════════════════════════════════════════════════════════
    section_heading(doc, "🏆", "QUÉ TE OFRECEMOS", bg_color=C_GREEN_OK)

    ofrecemos = [
        ("Proyecto de **alto impacto real** en transformación digital industrial", " — tus decisiones de arquitectura se verán en producción desde el primer sprint."),
        ("**Autonomía técnica plena** en datos, persistencia y núcleo de servicios", ": aquí se trabaja con argumentos, no con jerarquías."),
        ("Equipo multidisciplinario consolidado", ", con metodología ágil operativa y procesos definidos."),
        ("**Planilla completa desde el primer día.**", " Contrato inicial de 6 meses, renovable según desempeño."),
        ("Compensación **acorde a perfil Senior,**", " a coordinar en el proceso de selección."),
    ]
    for bold_part, rest in ofrecemos:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.8)
        para.paragraph_format.first_line_indent = Cm(-0.5)
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(4)
        tick = para.add_run("✔  ")
        tick.font.color.rgb = C_GREEN_OK
        tick.font.size = Pt(11)
        for seg, is_bold in parse_bold_segments(bold_part):
            r = para.add_run(seg)
            r.bold = is_bold
            r.font.size = Pt(10.5)
        r2 = para.add_run(rest)
        r2.font.size = Pt(10.5)

    # ══════════════════════════════════════════════════════════════════════════
    # CÓMO POSTULAR
    # ══════════════════════════════════════════════════════════════════════════
    section_heading(doc, "📩", "¿CÓMO POSTULAR?", bg_color=C_NAVY)

    p_correo = doc.add_paragraph()
    p_correo.paragraph_format.left_indent = Cm(0.5)
    p_correo.paragraph_format.space_before = Pt(6)
    p_correo.paragraph_format.space_after = Pt(2)
    r = p_correo.add_run("Envía los siguientes elementos al correo ")
    r.font.size = Pt(10.5)
    r_correo = p_correo.add_run("[CORREO DE CONTACTO]")
    r_correo.bold = True
    r_correo.font.color.rgb = C_RED_EXCL
    r_correo.font.size = Pt(10.5)
    r2 = p_correo.add_run(" con el asunto:")
    r2.font.size = Pt(10.5)

    p_asunto = doc.add_paragraph()
    p_asunto.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_asunto.paragraph_format.left_indent = Cm(1.0)
    p_asunto.paragraph_format.space_before = Pt(2)
    p_asunto.paragraph_format.space_after = Pt(8)
    set_para_shading(p_asunto, C_LIGHT_BG)
    r3 = p_asunto.add_run('  "Senior Backend / Arquitecto de Datos — [TU NOMBRE COMPLETO]"')
    r3.bold = True
    r3.italic = True
    r3.font.size = Pt(10.5)
    r3.font.color.rgb = C_NAVY

    pasos = [
        ("1. ", "CV actualizado en PDF ", "(máximo 3 páginas)."),
        ("2. ", "Portafolio o repositorios de trabajo técnico ", "(GitHub, GitLab, Bitbucket — opcional pero valorado)."),
        ("3. ", "Respuesta breve en el cuerpo del correo ", "(máx. 150 palabras) a la pregunta de filtro."),
    ]
    for num, bold_txt, rest in pasos:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        rn = p.add_run(num)
        rn.bold = True
        rn.font.color.rgb = C_TEAL
        rn.font.size = Pt(10.5)
        rb = p.add_run(bold_txt)
        rb.bold = True
        rb.font.size = Pt(10.5)
        rr = p.add_run(rest)
        rr.font.size = Pt(10.5)

    note_box(
        doc,
        '   "¿Cuál fue el proyecto de datos de mayor complejidad técnica en el que participaste? '
        'Describe específicamente tu rol, las decisiones de arquitectura que tomaste y el impacto '
        'medible que generaste."',
        bg=C_TAG_BG,
        size=10,
        italic=True,
    )

    p_prio = doc.add_paragraph()
    p_prio.paragraph_format.left_indent = Cm(0.5)
    p_prio.paragraph_format.space_before = Pt(4)
    p_prio.paragraph_format.space_after = Pt(10)
    r_p = p_prio.add_run(
        "Los candidatos que respondan esta pregunta tienen prioridad en la evaluación. "
        "Los que no la respondan no serán considerados en la primera revisión."
    )
    r_p.font.size = Pt(10)
    r_p.bold = True
    r_p.font.color.rgb = C_RED_EXCL

    # ══════════════════════════════════════════════════════════════════════════
    # PROCESO DE SELECCIÓN
    # ══════════════════════════════════════════════════════════════════════════
    section_heading(doc, "📋", "PROCESO DE SELECCIÓN", bg_color=C_TEAL)
    doc.add_paragraph()

    tbl = doc.add_table(rows=6, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Estilos de borde
    def set_table_borders(table):
        tbl2 = table._tbl
        tblPr2 = tbl2.find(qn('w:tblPr'))
        if tblPr2 is None:
            tblPr2 = OxmlElement('w:tblPr')
            tbl2.insert(0, tblPr2)
        borders = OxmlElement('w:tblBorders')
        for name in ('top','left','bottom','right','insideH','insideV'):
            b = OxmlElement(f'w:{name}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '4')
            b.set(qn('w:color'), hex_to_rgb_str(RGBColor(0xCC, 0xDD, 0xEE)))
            borders.append(b)
        tblPr2.append(borders)

    set_table_borders(tbl)

    headers = ["FASE", "DURACIÓN ESTIMADA", "RESPONSABLE"]
    rows_data = [
        ("Revisión de CV + respuesta técnica obligatoria", "48 horas", "Líder técnico / ARQ"),
        ("Entrevista técnica — Arquitectura de datos y PostgreSQL (45 min)", "1 sesión", "ARQ"),
        ("Challenge técnico asíncrono (máximo 4 h del candidato)", "3–5 días", "Candidato"),
        ("Entrevista final — encaje técnico y cultural (45 min)", "1 sesión", "ARQ + PM"),
        ("Oferta, negociación y firma", "48 horas", "PM / RRHH"),
    ]

    # Cabecera
    hrow = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        set_cell_bg(cell, C_NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = C_WHITE

    # Filas
    for ri, (fase, dur, resp) in enumerate(rows_data):
        row = tbl.rows[ri + 1]
        bg = C_LIGHT_BG if ri % 2 == 0 else C_WHITE
        for ci, val in enumerate([fase, dur, resp]):
            cell = row.cells[ci]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.size = Pt(9.5)

    doc.add_paragraph()

    # ══════════════════════════════════════════════════════════════════════════
    # HASHTAGS
    # ══════════════════════════════════════════════════════════════════════════
    divider(doc, C_GOLD)
    p_tags = doc.add_paragraph()
    p_tags.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tags.paragraph_format.space_before = Pt(6)
    p_tags.paragraph_format.space_after = Pt(6)
    tags = [
        "#ArquitectoDeDatos", "#BackendSenior", "#PostgreSQL", "#DataEngineering",
        "#TimescaleDB", "#Docker", "#Linux", "#IndustriaMinera",
        "#TransformacionDigital", "#Industria40", "#TechJobsLATAM", "#RemoteWork",
        "#DataArchitecture", "#ETL",
    ]
    r_tags = p_tags.add_run("  ".join(tags))
    r_tags.font.size = Pt(8.5)
    r_tags.font.color.rgb = C_TEAL

    # ══════════════════════════════════════════════════════════════════════════
    # PIE DE PÁGINA
    # ══════════════════════════════════════════════════════════════════════════
    section_obj = doc.sections[0]
    footer = section_obj.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f = fp.add_run("Documento de uso interno — Versión LinkedIn Ready · Mayo 2026")
    r_f.font.size = Pt(8)
    r_f.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    # ── Guardar ───────────────────────────────────────────────────────────────
    out = r"c:\InformeCliente\CONVOCATORIA_LINKEDIN.docx"
    doc.save(out)
    print(f"✅  Documento generado: {out}")


if __name__ == "__main__":
    build_document()
